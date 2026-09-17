"""
Bloco A - recorte geográfico: região metropolitana (RM) da capital do estado.

Módulo compartilhado, sem saída própria. Regra decidida pelo autor em
16/09/2026: fica no recorte o estabelecimento cujo município é a capital do
seu estado ou integra a RM dessa capital. Fica fora quem está no interior ou
em outra RM do mesmo estado (ex.: Santos, na RM da Baixada Santista).

  - Município e UF saem do endereço do Google, não da capital da busca (a
    busca "board game em Boa Vista" também devolve casas do Rio de Janeiro).
  - A RM de cada capital é a categoria metropolitana do IBGE que contém o
    município da capital (Composição dos Recortes Metropolitanos, situação
    2025). Todas as subcategorias contam (núcleo, área de expansão,
    sub-regiões).
  - Brasília e Teresina não têm RM, e sim RIDE (Região Integrada de
    Desenvolvimento), que faz o papel da RM. A capital de referência vem da
    UF do endereço: uma casa em Timon-MA, que é da RIDE da Grande Teresina,
    seria conferida contra São Luís e sairia do recorte. Em 16/09/2026 não há
    nenhum caso assim na curadoria.
  - Rio Branco não integra recorte metropolitano no arquivo do IBGE: a "RM"
    se reduz ao próprio município.

Fonte: IBGE, Composicao_RM_2025.xlsx, baixado em 16/09/2026 de
geoftp.ibge.gov.br/organizacao_do_territorio/estrutura_territorial/
municipios_por_regioes_metropolitanas/Situacao_2020a2029/
"""
import pathlib
import re
import unicodedata
from collections import defaultdict
from functools import lru_cache

from openpyxl import load_workbook

ROOT = pathlib.Path(__file__).resolve().parents[2]
RM_PATH = ROOT / "_data" / "raw" / "[E] Bases publicas" / "ibge_rm" / "Composicao_RM_2025.xlsx"

# Capital de cada UF, com a UF por sigla e por nome (o Google usa os dois).
CAPITAL_POR_UF = {
    "AC": "Rio Branco", "AL": "Maceió", "AM": "Manaus", "AP": "Macapá",
    "BA": "Salvador", "CE": "Fortaleza", "DF": "Brasília", "ES": "Vitória",
    "GO": "Goiânia", "MA": "São Luís", "MG": "Belo Horizonte", "MS": "Campo Grande",
    "MT": "Cuiabá", "PA": "Belém", "PB": "João Pessoa", "PE": "Recife",
    "PI": "Teresina", "PR": "Curitiba", "RJ": "Rio de Janeiro", "RN": "Natal",
    "RO": "Porto Velho", "RR": "Boa Vista", "RS": "Porto Alegre", "SC": "Florianópolis",
    "SE": "Aracaju", "SP": "São Paulo", "TO": "Palmas",
}
UF_POR_NOME = {
    "acre": "AC", "alagoas": "AL", "amazonas": "AM", "amapa": "AP", "bahia": "BA",
    "ceara": "CE", "distritofederal": "DF", "espiritosanto": "ES", "goias": "GO",
    "maranhao": "MA", "minasgerais": "MG", "matogrossodosul": "MS", "matogrosso": "MT",
    "para": "PA", "paraiba": "PB", "pernambuco": "PE", "piaui": "PI", "parana": "PR",
    "riodejaneiro": "RJ", "riograndedonorte": "RN", "rondonia": "RO", "roraima": "RR",
    "riograndedosul": "RS", "santacatarina": "SC", "sergipe": "SE", "saopaulo": "SP",
    "tocantins": "TO",
}

# "..., Rio de Janeiro - RJ, 22271-041, Brazil" ou "..., Goiânia, GO, 74595-331, Brasil"
RE_MUNICIPIO = [
    re.compile(r",\s*([^,]+?)\s+-\s+([A-Z]{2})(?:,|\s*$)"),
    re.compile(r"(?:^|[,-])\s*([^,\-]+?),\s*([A-Z]{2}),\s*\d{5}-?\d{3}"),
]
# "Pituaçu, Salvador - Bahia, 41741-170, Brazil": UF por extenso.
RE_UF_POR_EXTENSO = re.compile(r"([^,]+?)\s+-\s+([^,\d]+?),\s*\d{5}-?\d{3},\s*Bra[sz]il$")
# Só a UF, quando o município não é legível.
RE_UF = re.compile(r"[,-]\s*([^,-]+?)\s*,(?:\s*[\d-]+\s*,)?\s*Bra[sz]il\s*$")

# Abreviações que o Google usa no nome do município.
ABREVIACOES = [(r"\bsra\b\.?", "senhora"), (r"\bsto\b\.?", "santo"), (r"\bsta\b\.?", "santa")]

SITUACOES_NO_RECORTE = {"municipio_da_capital", "dentro_da_rm"}


def chave(texto):
    """Nome normalizado para comparação: sem acento, sem pontuação, sem
    espaço e com abreviações expandidas."""
    t = unicodedata.normalize("NFKD", texto or "").encode("ascii", "ignore").decode("ascii").lower()
    for padrao, troca in ABREVIACOES:
        t = re.sub(padrao, troca, t)
    return re.sub(r"[^a-z0-9]", "", t)


def municipio_uf(endereco):
    """(município, UF) do endereço; ("", "") se ilegível."""
    for regex in RE_MUNICIPIO:
        achados = regex.findall(endereco or "")
        if achados:
            municipio, uf = achados[-1]
            return municipio.strip(), uf
    m = RE_UF_POR_EXTENSO.search(endereco or "")
    if m and chave(m.group(2)) in UF_POR_NOME:
        return m.group(1).strip(), UF_POR_NOME[chave(m.group(2))]
    return "", ""


def uf_do_endereco(endereco):
    _, uf = municipio_uf(endereco)
    if uf:
        return uf
    m = RE_UF.search(endereco or "")
    if m:
        uf = m.group(1).strip()
        return uf if uf in CAPITAL_POR_UF else UF_POR_NOME.get(chave(uf), "")
    return ""


@lru_cache(maxsize=1)
def rms():
    """Capital -> (nome do recorte no IBGE, conjunto de (município, UF))."""
    ws = load_workbook(RM_PATH, read_only=True).worksheets[0]
    linhas = ws.iter_rows(values_only=True)
    idx = {c: i for i, c in enumerate(next(linhas))}
    por_categoria = defaultdict(set)
    for r in linhas:
        if r[idx["COD_CATMETROPOL"]] is None:
            continue
        por_categoria[r[idx["NOME_CATMETROPOL"]]].add((chave(r[idx["NOME_MUN"]]), r[idx["SIGLA_UF"]]))
    saida = {}
    for uf, capital in CAPITAL_POR_UF.items():
        k = (chave(capital), uf)
        cats = [c for c, muns in por_categoria.items() if k in muns]
        assert len(cats) <= 1, (capital, cats)
        saida[capital] = (cats[0], frozenset(por_categoria[cats[0]])) if cats else ("sem recorte", frozenset({k}))
    return saida


def classifica(endereco):
    """Situação do endereço no recorte. Devolve dict com municipio, uf,
    capital_referencia, recorte_ibge e situacao, que é uma de:
    municipio_da_capital, dentro_da_rm, fora_da_rm, sem_municipio_no_endereco."""
    municipio, uf = municipio_uf(endereco)
    uf = uf or uf_do_endereco(endereco)
    capital = CAPITAL_POR_UF.get(uf, "")
    recorte, muns = rms().get(capital, ("", frozenset()))
    if not municipio or not capital:
        situacao = "sem_municipio_no_endereco"
    elif (chave(municipio), uf) == (chave(capital), uf):
        situacao = "municipio_da_capital"
    elif (chave(municipio), uf) in muns:
        situacao = "dentro_da_rm"
    else:
        situacao = "fora_da_rm"
    return {
        "municipio": municipio, "uf": uf, "capital_referencia": capital,
        "recorte_ibge": recorte, "situacao": situacao,
    }


def no_recorte(endereco):
    return classifica(endereco)["situacao"] in SITUACOES_NO_RECORTE
