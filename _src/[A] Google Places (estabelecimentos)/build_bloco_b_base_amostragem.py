"""
Bloco A - base de sorteio e planilha de registro da amostra que define os
limiares do recorte (mínimo de avaliações no Maps e de publicações no
Instagram).

bloco_b_candidatos.csv tem uma linha por combinação capital x termo de busca:
a mesma casa aparece várias vezes. Sortear linhas daria mais chance às casas
encontradas por mais termos. A base tem uma linha por place_id.

Decisões:
  - Só endereços no Brasil; município e UF extraídos do endereço, não da
    capital da busca (a busca "board game em Boa Vista" também devolve casas
    do Rio de Janeiro).
  - Estratos pela REGIC 2018 (IBGE), em dois grupos: Metrópoles (14 capitais)
    e Capitais Regionais (13). Divisão extraída pelo autor em Metrópoles.md.
    Casas fora das capitais ficam fora do recorte.
  - Sorteio: cada casa recebe ALEATÓRIO() na planilha; uma fórmula ordena as
    casas por esse número dentro do estrato, e a aba Amostragem busca a 1ª,
    a 2ª, a 3ª casa de cada estrato. ALEATÓRIO() recalcula a cada edição: o
    autor copia e cola como valores o resultado do sorteio antes de registrar
    a amostra, e esses valores colados são o registro do sorteio.
  - O volume de avaliações NÃO entra na base: a checagem das fotos é feita
    sem saber quantas avaliações a casa tem. Na aba Amostragem, as colunas
    de evidência vêm antes das de contagem pelo mesmo motivo.

Saída: _data/processed/[A] Estabelecimentos e cardapios/bloco_b_base_amostragem.xlsx
  Aba "Base de sorteio" - uma linha por estabelecimento, com estrato e ordem
  Aba "Amostragem"      - registro da amostra, com fórmulas de sorteio
  Aba "Resumo"          - progresso por estrato até a meta de casas do tipo

Uso:
    python3 "_src/[A] Google Places (estabelecimentos)/build_bloco_b_base_amostragem.py"
"""
import csv
import pathlib
import re
import unicodedata
from collections import Counter, defaultdict
from urllib.parse import quote

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = pathlib.Path(__file__).resolve().parents[2]
RAW = ROOT / "_data" / "raw" / "[A] Estabelecimentos e cardapios"
CANDIDATOS_PATH = RAW / "bloco_b_candidatos.csv"
REGIC_PATH = RAW / "Metrópoles.md"
RESUMO_CAPITAIS_PATH = (ROOT / "_data" / "processed" / "[E] Bases publicas"
                        / "[E] Resumo de População e Renda por Capital 2022-2026 (IBGE).csv")
OUT_PATH = ROOT / "_data" / "processed" / "[A] Estabelecimentos e cardapios" / "bloco_b_base_amostragem.xlsx"

META_POR_ESTRATO = 45
FORA = "fora do recorte"

# "..., Rio de Janeiro - RJ, 22271-041, Brazil" ou "..., Goiânia, GO, 74595-331, Brasil"
RE_MUNICIPIO = [
    re.compile(r",\s*([^,]+?)\s+-\s+([A-Z]{2})(?:,|\s*$)"),
    re.compile(r"(?:^|[,-])\s*([^,\-]+?),\s*([A-Z]{2}),\s*\d{5}-?\d{3}"),
]

FONTE = Font(name="Arial", size=10, color="000000")
NEGRITO = Font(name="Arial", size=10, color="000000", bold=True)
FINA = Side(style="thin", color="000000")


def normaliza(texto):
    sem_acento = unicodedata.normalize("NFKD", texto or "").encode("ascii", "ignore").decode("ascii")
    return sem_acento.lower().strip()


def municipio_uf(endereco):
    for regex in RE_MUNICIPIO:
        achados = regex.findall(endereco or "")
        if achados:
            municipio, uf = achados[-1]
            return municipio.strip(), uf
    return "", ""


def load_capitais():
    with open(RESUMO_CAPITAIS_PATH, newline="", encoding="utf-8") as f:
        return {(normaliza(r["capital"]), r["uf"]) for r in csv.DictReader(f)}


def load_regic():
    """Lê Metrópoles.md: '# ' é o estrato, '## ' o nível REGIC, '- ' a capital."""
    regic = {}
    estrato = nivel = None
    for linha in REGIC_PATH.read_text(encoding="utf-8").splitlines():
        linha = linha.strip()
        if linha.startswith("## "):
            nivel = linha[3:].strip()
        elif linha.startswith("# "):
            estrato = "Metrópole" if normaliza(linha).startswith("# metropole") else "Capital regional"
        elif linha.startswith("- "):
            regic[normaliza(linha[2:])] = (estrato, nivel)
    return regic


def load_casas():
    casas, capitais_busca, termos = {}, defaultdict(set), defaultdict(set)
    with open(CANDIDATOS_PATH, newline="", encoding="utf-8") as f:
        for r in csv.DictReader(f):
            pid = r["place_id"]
            if not pid:
                continue
            casas.setdefault(pid, r)
            capitais_busca[pid].add(r["capital_busca"])
            termos[pid].add(r["termo_busca"])
    return casas, capitais_busca, termos


def cabecalho(ws, titulo, colunas):
    ws.cell(row=1, column=1, value=titulo).font = NEGRITO
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(colunas))
    ws.cell(row=1, column=1).alignment = Alignment(horizontal="center")
    for j, (rotulo, largura) in enumerate(colunas, start=1):
        c = ws.cell(row=2, column=j, value=rotulo)
        c.font = NEGRITO
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = Border(top=FINA, bottom=FINA)
        ws.column_dimensions[get_column_letter(j)].width = largura
    ws.row_dimensions[2].height = 40


def celula(ws, linha, coluna, valor, direita=False, ultima=False, formato=None):
    c = ws.cell(row=linha, column=coluna, value=valor)
    c.font = FONTE
    c.alignment = Alignment(horizontal="right" if direita else "left")
    if formato:
        c.number_format = formato
    if ultima:
        c.border = Border(bottom=FINA)
    return c


def aba_base(ws, linhas):
    colunas = [
        ("Número aleatório", 12),        # A
        ("place_id", 30),                # B
        ("Nome", 40),                    # C
        ("Município", 22),               # D
        ("UF", 6),                       # E
        ("Nível REGIC 2018", 26),        # F
        ("Estrato", 16),                 # G
        ("Ordem no estrato", 10),        # H  fórmula
        ("Chave de sorteio", 22),        # I  fórmula
        ("Link do Google Maps", 40),     # J
        ("Endereço", 60),                # K
        ("Capitais da busca", 28),       # L
        ("Nº de termos que acharam", 12),  # M
    ]
    cabecalho(ws, "Base de sorteio do Bloco A: uma linha por estabelecimento candidato", colunas)
    fim = len(linhas) + 2
    for i, l in enumerate(linhas, start=3):
        u = i == fim
        celula(ws, i, 1, "=RAND()", direita=True, ultima=u, formato="0.000000")
        celula(ws, i, 2, l["place_id"], ultima=u)
        celula(ws, i, 3, l["nome"], ultima=u)
        celula(ws, i, 4, l["municipio"], ultima=u)
        celula(ws, i, 5, l["uf"], ultima=u)
        celula(ws, i, 6, l["nivel_regic"], ultima=u)
        celula(ws, i, 7, l["estrato"], ultima=u)
        # Posição da casa no sorteio do seu estrato: quantas casas do mesmo
        # estrato têm número aleatório menor ou igual ao dela. SUMPRODUCT
        # compara números; COUNTIFS com "<="&A converteria o número em texto
        # de 15 dígitos e poderia errar a ordem por arredondamento.
        celula(
            ws, i, 8,
            f'=IF(G{i}="{FORA}","",SUMPRODUCT(($G$3:$G${fim}=G{i})*($A$3:$A${fim}<=A{i})))',
            direita=True, ultima=u, formato="0",
        )
        celula(ws, i, 9, f'=IF(H{i}="","",G{i}&"|"&H{i})', ultima=u)
        link = celula(ws, i, 10, l["link_maps"], ultima=u)
        link.hyperlink = l["link_maps"]
        celula(ws, i, 11, l["endereco"], ultima=u)
        celula(ws, i, 12, l["capitais_busca"], ultima=u)
        celula(ws, i, 13, l["n_termos"], direita=True, ultima=u, formato="0")
    ws.freeze_panes = "D3"
    ws.auto_filter.ref = f"A2:M{fim}"
    return fim


def aba_amostragem(ws, tamanhos, fim_base):
    colunas = [
        ("Nº do item", 8),                                   # A
        ("Estrato", 16),                                     # B
        ("Ordem no estrato", 10),                            # C
        ("Número aleatório sorteado", 12),                   # D
        ("place_id", 30),                                    # E
        ("Nome do estabelecimento", 40),                     # F
        ("Município", 20),                                   # G
        ("UF", 6),                                           # H
        ("Link do Google Maps", 14),                         # I
        ("Tem fotos no Maps?", 11),                          # J
        ("Fotos comprovam espaço de jogo?", 14),             # K
        ("Tem Instagram?", 11),                              # L
        ("Instagram comprova espaço de jogo?", 14),          # M
        ("É do tipo?", 12),                                  # N
        ("Casas do tipo acumuladas no estrato", 12),         # O
        ("Nº de avaliações no Maps", 12),                    # P
        ("Usuário do Instagram", 22),                        # Q
        ("Nº de publicações no Instagram", 12),              # R
        ("Data da última publicação", 13),                   # S
        ("Observação", 40),                                  # T
    ]
    cabecalho(ws, "Amostragem do Bloco A: registro das casas sorteadas por estrato", colunas)

    chave = f"'Base de sorteio'!$I$3:$I${fim_base}"

    def busca(coluna_base, i):
        return (
            f"=IFERROR(INDEX('Base de sorteio'!${coluna_base}$3:${coluna_base}${fim_base},"
            f'MATCH($B{i}&"|"&$C{i},{chave},0)),"")'
        )

    linhas = [(e, k) for e in ("Metrópole", "Capital regional") for k in range(1, tamanhos[e] + 1)]
    fim = len(linhas) + 2
    for i, (estrato, k) in enumerate(linhas, start=3):
        u = i == fim
        celula(ws, i, 1, f"=ROW()-2", direita=True, ultima=u, formato="0")
        celula(ws, i, 2, estrato, ultima=u)
        celula(ws, i, 3, k, direita=True, ultima=u, formato="0")
        celula(ws, i, 4, busca("A", i), direita=True, ultima=u, formato="0")
        celula(ws, i, 5, busca("B", i), ultima=u)
        celula(ws, i, 6, busca("C", i), ultima=u)
        celula(ws, i, 7, busca("D", i), ultima=u)
        celula(ws, i, 8, busca("E", i), ultima=u)
        celula(ws, i, 9, f'=IF(E{i}="","",HYPERLINK(' + busca("J", i)[1:] + ',"abrir"))', ultima=u)
        for col in (10, 11, 12, 13, 16, 17, 18, 19, 20):
            celula(ws, i, col, None, direita=col in (16, 18), ultima=u, formato="0" if col in (16, 18) else None)
        ws.cell(row=i, column=19).number_format = "dd/mm/yyyy"
        # Do tipo: basta uma das duas fontes comprovar. "inconclusivo" só
        # quando nenhuma comprova e ao menos uma ficou em dúvida.
        celula(
            ws, i, 14,
            f'=IF(AND(K{i}="",M{i}=""),"",IF(OR(K{i}="sim",M{i}="sim"),"sim",'
            f'IF(OR(K{i}="inconclusivo",M{i}="inconclusivo"),"inconclusivo","não")))',
            ultima=u,
        )
        celula(ws, i, 15, f'=IF(N{i}="sim",COUNTIFS($B$3:$B{i},$B{i},$N$3:$N{i},"sim"),"")', direita=True, ultima=u, formato="0")

    sim_nao = DataValidation(type="list", formula1='"sim,não"', allow_blank=True)
    comprova = DataValidation(type="list", formula1='"sim,não,inconclusivo"', allow_blank=True)
    ws.add_data_validation(sim_nao)
    ws.add_data_validation(comprova)
    sim_nao.add(f"J3:J{fim}")
    sim_nao.add(f"L3:L{fim}")
    comprova.add(f"K3:K{fim}")
    comprova.add(f"M3:M{fim}")
    ws.freeze_panes = "G3"
    ws.auto_filter.ref = f"A2:T{fim}"
    return fim


def aba_resumo(ws, fim_amostra):
    colunas = [
        ("Estrato", 18),
        ("Candidatos na base", 12),
        ("Casas verificadas", 12),
        ("Casas do tipo", 12),
        ("Meta de casas do tipo", 12),
        ("Faltam", 10),
        ("Proporção de casas do tipo", 14),
    ]
    cabecalho(ws, "Resumo da amostragem por estrato", colunas)
    amostra = f"Amostragem!$B$3:$B${fim_amostra}"
    for i, estrato in enumerate(("Metrópole", "Capital regional", "Total"), start=3):
        u = i == 5
        celula(ws, i, 1, estrato, ultima=u)
        if estrato == "Total":
            for col in "BCDEF":
                j = ord(col) - 64
                celula(ws, i, j, f"=SUM({col}3:{col}4)", direita=True, ultima=u, formato="0")
        else:
            celula(ws, i, 2, f"=COUNTIF('Base de sorteio'!$G:$G,A{i})", direita=True, formato="0")
            celula(ws, i, 3, f'=COUNTIFS({amostra},A{i},Amostragem!$K$3:$K${fim_amostra},"<>")', direita=True, formato="0")
            celula(ws, i, 4, f'=COUNTIFS({amostra},A{i},Amostragem!$N$3:$N${fim_amostra},"sim")', direita=True, formato="0")
            celula(ws, i, 5, META_POR_ESTRATO, direita=True, formato="0")
            celula(ws, i, 6, f"=MAX(0,E{i}-D{i})", direita=True, formato="0")
        celula(ws, i, 7, f'=IF(C{i}=0,"",D{i}/C{i})', direita=True, ultima=u, formato="0.00%")


def main():
    capitais = load_capitais()
    regic = load_regic()
    casas, capitais_busca, termos = load_casas()

    faltando = {c for c, _ in capitais} - set(regic)
    if faltando:
        raise SystemExit(f"Capitais sem nível REGIC em Metrópoles.md: {sorted(faltando)}")

    linhas = []
    for pid in sorted(casas):
        r = casas[pid]
        endereco = (r["endereco"] or "").strip()
        if not normaliza(endereco).endswith(("brazil", "brasil")):
            continue
        municipio, uf = municipio_uf(endereco)
        chave = (normaliza(municipio), uf)
        estrato, nivel = regic[chave[0]] if chave in capitais else (FORA, "")
        linhas.append(
            {
                "place_id": pid,
                "nome": r["nome"],
                "municipio": municipio,
                "uf": uf,
                "nivel_regic": nivel,
                "estrato": estrato,
                "endereco": endereco,
                "capitais_busca": "; ".join(sorted(capitais_busca[pid])),
                "n_termos": len(termos[pid]),
                "link_maps": (
                    "https://www.google.com/maps/search/?api=1"
                    f"&query={quote(r['nome'] or '')}&query_place_id={pid}"
                ),
            }
        )

    linhas.sort(key=lambda x: (x["estrato"], x["uf"], x["municipio"], x["nome"] or ""))

    tamanhos = Counter(l["estrato"] for l in linhas)

    wb = Workbook()
    fim_base = aba_base(wb.active, linhas)
    wb.active.title = "Base de sorteio"
    fim_amostra = aba_amostragem(wb.create_sheet("Amostragem"), tamanhos, fim_base)
    aba_resumo(wb.create_sheet("Resumo"), fim_amostra)
    wb.active = 1
    OUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT_PATH)

    print(f"OK: {len(linhas)} estabelecimentos únicos no Brasil em {OUT_PATH.name}")
    for estrato, n in sorted(tamanhos.items()):
        print(f"    {estrato}: {n}")
    print(f"    {len(casas) - len(linhas)} fora do Brasil excluídos")


if __name__ == "__main__":
    main()
