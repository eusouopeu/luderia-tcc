"""
Bloco B1 - separa os comentários coletados à mão das luderias aprovadas em
colunas, um comentário por linha.

O autor anotou cada comentário numa célula, com os campos separados por "::":
  data relativa :: nota (notas por aspecto) :: preço :: contexto :: texto
O contexto junta barulho, tamanho do grupo e espera, separados por vírgula.
Nenhum campo é obrigatório; cada parte é reconhecida pelo conteúdo, não pela
posição. Campo ausente no comentário recebe "R" (revisar).

Padronizações:
  - data: letra da unidade antes do número, para a ordenação alfabética
    agrupar por unidade: "2 dias" -> "d2", "3S" -> "s3", "fim de um mes atrás"
    -> "m1"; a anotação aceita as duas ordens ("2d" ou "d2");
  - notas por aspecto: "comida e ambiente 4" vale 4 para os dois aspectos;
    aspecto não citado, "sem ambiente" ou "não avaliado" fica "R";
  - preço: "+200" -> "200+";
  - barulho, grupo e espera: rótulos únicos para grafias diferentes.
Partes que não se encaixam em nenhum campo vão para a coluna "sobras" e são
listadas no terminal.

Entrada: estabelecimentos/3_triagem/B1_aprovadas_metro_av100.xlsx
         (B1-b_triagem_manual.py)
Saída em dados/[B1] Estabelecimentos e cardapios/:
  estabelecimentos/3_triagem/B1_comentarios_aprovadas_metro_av100.xlsx
Fica fora do repositório (contém texto de avaliações).

Uso:
    python3 "_src/[B1] Estabelecimentos e cardapios/Google Places (Negócios)/build/B1-b_comentarios_aprovadas.py"
"""
# Caminhos: _src/extras/ (log comum) e outros/ do objeto (auxiliares); log em ../logs/
import sys
from pathlib import Path
_AQUI = Path(__file__).resolve()
sys.path[:0] = [str(_AQUI.parents[3] / "extras"), str(_AQUI.parents[1] / "outros")]
from log_execucao import registrar_execucao  # noqa: E402
if __name__ == "__main__":
    registrar_execucao(__file__)

import pathlib
import re

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

ROOT = pathlib.Path(__file__).resolve().parents[4]
PROCESSED = ROOT / "dados" / "[B1] Estabelecimentos e cardapios"
ENTRADA = PROCESSED / "estabelecimentos/3_triagem/B1_aprovadas_metro_av100.xlsx"
SAIDA = PROCESSED / "estabelecimentos/3_triagem/B1_comentarios_aprovadas_metro_av100.xlsx"

R = "R"
ASPECTOS = ("comida", "serviço", "ambiente")
CAMPOS = [
    "data relativa", "nota geral", "nota comida", "nota serviço", "nota ambiente",
    "preço por pessoa (R$)", "nível de barulho", "tamanho do grupo", "tempo de espera", "texto",
]

RE_DATA = re.compile(r"^(\d+)\s*(h|d|s|m|dias?|semanas?|mes|mês|meses)$", re.I)
RE_DATA_LETRA = re.compile(r"^(h|d|s|m)\s*(\d+)$", re.I)
RE_NOTA = re.compile(r"^([1-5])\s*(?:\((.*?)\)?)?$")
RE_PRECO = re.compile(r"^\+?\d{2,3}(?:\s*-\s*\d{2,3})?\+?$")
UNIDADE = {"dia": "d", "dias": "d", "semana": "s", "semanas": "s", "mes": "m", "mês": "m", "meses": "m"}


def sem_acento(texto):
    return texto.lower().translate(str.maketrans("áâãéêíóôõúç", "aaaeeiooouc"))


def data(parte):
    m = RE_DATA.match(parte)
    if m:
        numero, unidade = m.group(1), m.group(2).lower()
    elif m := RE_DATA_LETRA.match(parte):
        unidade, numero = m.group(1).lower(), m.group(2)
    else:
        return None
    return f"{UNIDADE.get(unidade, unidade)}{numero}"


def notas_aspecto(texto):
    """'ambiente 5, comida e serviço 1' -> {'ambiente': '5', 'comida': '1', 'serviço': '1'}."""
    notas = {}
    texto = texto.replace("comida 40", "comida 4")  # parêntese não fechado na anotação
    pendentes = []  # "comida, serviço e ambiente 5": aspectos sem número herdam o próximo
    for trecho in texto.split(","):
        t = sem_acento(trecho)
        if re.search(r"\b(sem|nao avaliado)\b", t):
            continue
        pendentes += [a for a in ASPECTOS if sem_acento(a) in t]
        numero = re.search(r"(\d)\s*$", trecho.strip())
        if numero:
            notas.update({a: numero.group(1) for a in pendentes})
            pendentes = []
    return notas


def barulho(trecho):
    t = sem_acento(trecho)
    if "silencioso" in t:
        return "silencioso, fácil de conversar"
    if "muito alto" in t:
        return "muito alto, difícil conversar"
    if "alto" in t:
        return "alto, mas possível conversar"
    if "moderado" in t:
        return "moderado"
    return None


def grupo(trecho):
    t = sem_acento(trecho)
    if "mais de 9" in t:
        return "mais de 9 pessoas"
    m = re.search(r"(\d+)\s*a\s*(\d+)", t)
    if m:
        return f"{m.group(1)} a {m.group(2)} pessoas"
    m = re.search(r"(\d+)\s*pessoas?", t)
    return f"{m.group(1)} pessoas" if m else None


def espera(trecho):
    t = sem_acento(trecho)
    if "mais de uma hora" in t or "mais de 1 hora" in t:
        return "mais de 1 hora"
    if "1 hora" in t or "uma hora" in t:
        return "1 hora"
    m = re.search(r"(\d+)\s*a\s*(\d+)", t)
    if m:
        return f"{m.group(1)} a {m.group(2)} minutos"
    if re.search(r"10 minutos", t):
        return "até 10 minutos"
    return None


def contexto(parte):
    """Divide o contexto em barulho, grupo e espera; devolve None se não for contexto."""
    t = sem_acento(parte)
    if not re.search(r"barulho|silencioso|nivel alto|pessoas|grupo|espera|\d+\s*a\s*\d+$", t) or len(parte) > 160:
        return None
    achados = {}
    for trecho in parte.split(","):
        ts = sem_acento(trecho)
        if "espera" in ts:
            achados["tempo de espera"] = espera(trecho)
        elif "grupo" in ts or "pessoa" in ts or re.search(r"\d+\s*a\s*\d+", ts):
            achados["tamanho do grupo"] = grupo(trecho)
        elif barulho(trecho):
            achados["nível de barulho"] = barulho(trecho)
        # "fácil de conversar" sozinho completa o barulho já lido
    return {k: v for k, v in achados.items() if v}


def separa(comentario):
    linha = {campo: None for campo in CAMPOS}
    sobras = []
    partes = [p.strip() for p in str(comentario).split("::") if p.strip()]
    for i, parte in enumerate(partes):
        ultima = i == len(partes) - 1
        if linha["data relativa"] is None and data(parte):
            linha["data relativa"] = data(parte)
            continue
        nota = RE_NOTA.match(parte)
        if linha["nota geral"] is None and nota:
            linha["nota geral"] = nota.group(1)
            dentro = nota.group(2) or ""
            if "mes atras" in sem_acento(dentro):
                linha["data relativa"] = linha["data relativa"] or "m1"
            for aspecto, valor in notas_aspecto(dentro).items():
                linha[f"nota {aspecto}"] = valor
            continue
        if linha["preço por pessoa (R$)"] is None and RE_PRECO.match(parte):
            valor = parte.replace(" ", "")
            linha["preço por pessoa (R$)"] = valor.lstrip("+") + "+" if valor.startswith("+") else valor
            continue
        ctx = contexto(parte) if not ultima or len(parte) < 80 else None
        if ctx:
            for campo, valor in ctx.items():
                linha[campo] = linha[campo] or valor
            continue
        if linha["texto"] is None and (ultima or len(parte) > 40):
            linha["texto"] = parte
            continue
        sobras.append(parte)
    linha = {k: (R if v is None else v) for k, v in linha.items()}
    linha["sobras"] = " | ".join(sobras)
    return linha


def main():
    aprovadas = pd.read_excel(ENTRADA)
    colunas_coment = [c for c in aprovadas.columns if str(c).startswith("Comentário ")]
    linhas = []
    for _, est in aprovadas.iterrows():
        for col in colunas_coment:
            if pd.isna(est[col]):
                continue
            linhas.append({
                "place_id": est["place_id"],
                "nome": est["nome"],
                "capital_endereco": est["capital_endereco"],
                "nº do comentário": int(col.split()[-1]),
                **separa(est[col]),
                "comentário original": est[col],
            })
    df = pd.DataFrame(linhas)
    df.to_excel(SAIDA, index=False, sheet_name="comentarios")

    wb = load_workbook(SAIDA)
    ws = wb.active
    larguras = {"nome": 30, "texto": 70, "comentário original": 50, "sobras": 25, "place_id": 14}
    for i, celula in enumerate(ws[1], start=1):
        celula.font = Font(bold=True)
        celula.alignment = Alignment(wrap_text=True, vertical="top")
        ws.column_dimensions[get_column_letter(i)].width = larguras.get(celula.value, 13)
    ws.freeze_panes = "E2"
    ws.auto_filter.ref = ws.dimensions
    wb.save(SAIDA)

    total_celulas = aprovadas[colunas_coment].notna().sum().sum()
    print(f"comentários: {len(df)} | células preenchidas na entrada: {total_celulas} | diferença: {total_celulas - len(df)}")
    print(f"estabelecimentos: {df['place_id'].nunique()} de {len(aprovadas)}")
    print("\n'R' por coluna:")
    print((df[CAMPOS] == R).sum().to_string())
    for campo in ["data relativa", "nota geral", "preço por pessoa (R$)", "nível de barulho", "tamanho do grupo", "tempo de espera"]:
        print(f"\n{campo}: {df[campo].value_counts().to_dict()}")
    print("\nsobras:")
    print(df.loc[df["sobras"] != "", ["nome", "nº do comentário", "sobras"]].to_string())
    print(f"\n{SAIDA.relative_to(ROOT)}")


if __name__ == "__main__":
    main()
