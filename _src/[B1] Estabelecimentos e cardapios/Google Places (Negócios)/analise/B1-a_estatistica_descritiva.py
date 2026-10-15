"""
Bloco B1 - estatística descritiva das luderias do estudo: média, mediana,
desvio padrão, mínimo, quartis, máximo, amplitude e coeficiente de variação de
cada variável, para o conjunto inteiro (54 casas em 15/10/2026), por metrópole
e por região do Brasil.

Entrada: estabelecimentos/B1_resumo_estabelecimentos.xlsx (B1-b_resumo.py),
abas estabelecimentos e cesta. As cestas são somadas aqui a partir das colunas
de preço da aba cesta (couvert + bebida + entrada ou prato, mesmo critério da
fórmula da planilha), para não depender dos valores em cache.

Regras:
- cada place_id conta uma vez, na metrópole do endereço (capital_endereco) e
  na região da UF dessa capital (IBGE);
- Instagram de LudoCafé e Dominaria fica vazio: o perfil coletado não é o da
  casa aprovada (editora e unidade de São Paulo);
- couvert entra só nas casas que cobram couvert (sem couvert = vazio, não 0);
- shoppings de porte só nas casas em que o porte foi buscado;
- interações por post = (curtidas + comentários) ÷ posts do período; taxa de
  engajamento = interações por post ÷ seguidores;
- desvio padrão amostral (STDEV); quartis pelo método inclusivo (PERCENTILE),
  o mesmo do pandas.

Saídas (dados/[B1] Estabelecimentos e cardapios/estabelecimentos/4_estatistica_descritiva/):
  B1_estatistica_descritiva_geral.xlsx      abas geral, frequencias, dados, dicionario
  B1_estatistica_descritiva_por_grupo.xlsx  abas metropole_painel, metropole_detalhe,
                                            regiao_painel, regiao_detalhe, dados
  Todas as estatísticas são fórmulas sobre a aba dados; recalcular no
  LibreOffice ou no Excel antes de ler os valores por script.
  _src/site-apresentacao/src/conteudo/b1-estatisticas.json  (páginas do site)

Uso:
    python3 "_src/[B1] Estabelecimentos e cardapios/Google Places (Negócios)/analise/B1-a_estatistica_descritiva.py"
"""
# Caminhos: _src/extras/ (log comum); log em ../logs/
import sys
from pathlib import Path
_AQUI = Path(__file__).resolve()
sys.path[:0] = [str(_AQUI.parents[3] / "extras")]
from log_execucao import registrar_execucao  # noqa: E402
if __name__ == "__main__":
    registrar_execucao(__file__)

import json
from datetime import date

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.formula import ArrayFormula

ROOT = _AQUI.parents[4]
B1 = ROOT / "dados" / "[B1] Estabelecimentos e cardapios"
RESUMO = B1 / "estabelecimentos" / "B1_resumo_estabelecimentos.xlsx"
PASTA = B1 / "estabelecimentos" / "4_estatistica_descritiva"
SAIDA_GERAL = PASTA / "B1_estatistica_descritiva_geral.xlsx"
SAIDA_GRUPO = PASTA / "B1_estatistica_descritiva_por_grupo.xlsx"
SAIDA_SITE = ROOT / "_src" / "site-apresentacao" / "src" / "conteudo" / "b1-estatisticas.json"
DATA_REF = date(2026, 10, 15)
FONTE = "Arial"

# Capital da metrópole -> região do Brasil (IBGE, pela UF da capital).
REGIAO = {
    "Belém": "Norte", "Manaus": "Norte",
    "Fortaleza": "Nordeste", "Recife": "Nordeste", "Salvador": "Nordeste",
    "Brasília": "Centro-Oeste", "Goiânia": "Centro-Oeste",
    "Belo Horizonte": "Sudeste", "Rio de Janeiro": "Sudeste", "São Paulo": "Sudeste", "Vitória": "Sudeste",
    "Curitiba": "Sul", "Florianópolis": "Sul", "Porto Alegre": "Sul",
}
ORDEM_REGIAO = ["Norte", "Nordeste", "Centro-Oeste", "Sudeste", "Sul"]
PERFIL_ERRADO = {"LudoCafé", "Dominaria Cards & Games"}
DIAS = ["segunda", "terca", "quarta", "quinta", "sexta", "sabado", "domingo"]
NOME_DIA = {"segunda": "segunda", "terca": "terça", "quarta": "quarta", "quinta": "quinta",
            "sexta": "sexta", "sabado": "sábado", "domingo": "domingo"}

# coluna da aba dados -> (grupo, rótulo, unidade, formato)
REAIS, INT, DEC1, DEC2, PCT = '"R$" #,##0.00', "#,##0", "#,##0.0", "#,##0.00", "0.0%"
VARS = {
    "avaliacao_media": ("Google Maps", "Nota média no Google", "1 a 5", DEC2),
    "volume_avaliacoes": ("Google Maps", "Nº de avaliações no Google", "avaliações", INT),
    "dias_abertos": ("Horários", "Dias abertos por semana", "dias", DEC1),
    "horas_semana": ("Horários", "Horas abertas por semana", "horas", DEC1),
    "abre_cedo_h": ("Horários", "Abertura mais cedo da semana", "hora do dia", DEC1),
    "fecha_tarde_h": ("Horários", "Fechamento mais tarde da semana (25 = 1h da manhã)", "hora do dia", DEC1),
    "metro_500m": ("Entorno", "Estações de metrô até 500 m", "estações", DEC2),
    "metro_1000m": ("Entorno", "Estações de metrô até 1000 m", "estações", DEC2),
    "dist_metro_m": ("Entorno", "Distância à estação de metrô mais próxima", "metros (linha reta)", INT),
    "shopping_porte_500m": ("Entorno", "Shoppings de porte até 500 m (onde o porte foi buscado)", "shoppings", DEC2),
    "shopping_porte_1000m": ("Entorno", "Shoppings de porte até 1000 m (onde o porte foi buscado)", "shoppings", DEC2),
    "dist_shopping_m": ("Entorno", "Distância ao shopping mais próximo (qualquer porte)", "metros (linha reta)", INT),
    "seguidores": ("Instagram", "Seguidores", "seguidores", INT),
    "n_posts_total": ("Instagram", "Posts desde a criação do perfil", "posts", INT),
    "idade_perfil_anos": ("Instagram", "Idade do perfil (desde o primeiro post até 15/10/2026)", "anos", DEC1),
    "n_posts_periodo": ("Instagram", "Posts de 01/09 a 01/10/2026", "posts", DEC1),
    "curtidas_periodo": ("Instagram", "Curtidas no período", "curtidas", INT),
    "comentarios_periodo": ("Instagram", "Comentários no período", "comentários", INT),
    "interacoes_por_post": ("Instagram", "Interações por post no período (curtidas + comentários)", "interações", DEC1),
    "engajamento_pct": ("Instagram", "Taxa de engajamento (interações por post ÷ seguidores)", "%", "0.00%"),
    "itens_com_preco": ("Cardápio", "Itens com preço no cardápio", "itens", INT),
    "couvert_mediana": ("Cardápio", "Couvert ou taxa de jogo (só casas que cobram)", "R$", REAIS),
    "bebida_mediana": ("Cardápio", "Bebida, mediana do cardápio", "R$", REAIS),
    "entrada_ou_prato_mediana": ("Cardápio", "Entrada ou prato, mediana do cardápio", "R$", REAIS),
    "cesta_baixa": ("Cardápio", "Cesta por pessoa, nível baixo (P25)", "R$", REAIS),
    "cesta_mediana": ("Cardápio", "Cesta por pessoa, nível mediano", "R$", REAIS),
    "cesta_alta": ("Cardápio", "Cesta por pessoa, nível alto (P75)", "R$", REAIS),
    "hamburguer_mediana": ("Cardápio", "Hambúrguer, mediana", "R$", REAIS),
    "batata_mediana": ("Cardápio", "Batata, mediana", "R$", REAIS),
    "cerveja_mediana": ("Cardápio", "Cerveja, mediana", "R$", REAIS),
    "refrigerante_mediana": ("Cardápio", "Refrigerante, mediana", "R$", REAIS),
    "drinks_mediana": ("Cardápio", "Drinks, mediana", "R$", REAIS),
    "titulos_acervo": ("Acervo", "Títulos no acervo publicado", "jogos", INT),
    "avaliacoes_b2": ("B2", "Avaliações no recorte do B2 (até '3 meses atrás')", "avaliações", INT),
    "avaliacoes_b2_com_texto": ("B2", "Dessas, com texto", "avaliações", INT),
}
# Proporções (0/1 por casa): a média é a proporção de casas.
FLAGS = {
    "tem_metro_1000m": "Com estação de metrô até 1000 m",
    "tem_cardapio": "Com cardápio transcrito",
    "cobra_couvert": "Cobra couvert ou taxa de jogo (entre as com cardápio)",
    **{f"abre_{d}": f"Abre {NOME_DIA[d]}" for d in DIAS},
}
CORES = {"Google Maps": "DDEBF7", "Horários": "E2EFDA", "Entorno": "FFF2CC", "Instagram": "FCE4D6",
         "Cardápio": "EDE2F6", "Acervo": "DEEBF7", "B2": "F2F2F2", "Proporção": "E7E6E6"}


def hora(txt):
    if not isinstance(txt, str) or ":" not in txt:
        return None
    h, m = txt.split(":")[:2]
    return int(h) + int(m) / 60


def montar_dados():
    e = pd.read_excel(RESUMO, sheet_name="estabelecimentos")
    c = pd.read_excel(RESUMO, sheet_name="cesta")
    for nivel, suf in [("cesta_baixa", "p25"), ("cesta_mediana", "mediana"), ("cesta_alta", "p75")]:
        ok = c[f"bebida_{suf}"].notna() & c[f"entrada_ou_prato_{suf}"].notna()
        c[nivel] = (c[f"couvert_{suf}"] + c[f"bebida_{suf}"] + c[f"entrada_ou_prato_{suf}"]).where(ok).round(2)
    c["couvert_mediana"] = c["couvert_mediana"].where(c["couvert_no_cardapio"] == "sim")
    precos = ["couvert_mediana", "bebida_mediana", "entrada_ou_prato_mediana", "cesta_baixa", "cesta_mediana",
              "cesta_alta", "hamburguer_mediana", "batata_mediana", "cerveja_mediana", "refrigerante_mediana",
              "drinks_mediana"]
    d = e.drop(columns=["cesta_baixa", "cesta_mediana", "cesta_alta"]).merge(
        c[["place_id", *precos]], on="place_id", how="left")
    assert len(d) == len(e) and d.place_id.is_unique

    d["metropole"] = d["capital_endereco"]
    d["regiao"] = d["metropole"].map(REGIAO)
    assert d.regiao.notna().all(), d[d.regiao.isna()].metropole.unique()
    d["abre_cedo_h"] = d["abre_cedo"].map(hora)
    d["fecha_tarde_h"] = d["fecha_tarde"].map(hora)
    pp = pd.to_datetime(d["primeiro_post"], errors="coerce")
    d["idade_perfil_anos"] = ((pd.Timestamp(DATA_REF) - pp).dt.days / 365.25).round(2)
    ig = ["seguidores", "n_posts_total", "idade_perfil_anos", "n_posts_periodo", "curtidas_periodo",
          "comentarios_periodo"]
    d.loc[d.nome.isin(PERFIL_ERRADO), ig] = None
    for col in ig:
        d[col] = pd.to_numeric(d[col], errors="coerce")
    d["interacoes_por_post"] = ((d.curtidas_periodo + d.comentarios_periodo) / d.n_posts_periodo).where(
        d.n_posts_periodo > 0).round(2)
    d["engajamento_pct"] = (d.interacoes_por_post / d.seguidores).round(5)
    d["tem_metro_1000m"] = (d.metro_1000m > 0).astype(int)
    d["tem_cardapio"] = d.itens_com_preco.notna().astype(int)
    d["cobra_couvert"] = d.couvert_no_cardapio.map({"sim": 1, "não": 0})
    for dia in DIAS:
        d[f"abre_{dia}"] = d[dia].notna().astype(int)
    cols = ["place_id", "nome", "metropole", "regiao", *VARS, *FLAGS]
    return d[cols].sort_values(["regiao", "metropole", "nome"]).reset_index(drop=True)


# ---------- estatísticas em pandas (para o site e para conferir a planilha) ----------
def resumo(s):
    s = pd.to_numeric(s, errors="coerce").dropna()
    if not len(s):
        return dict(n=0)
    r = dict(n=int(len(s)), media=s.mean(), mediana=s.median(), dp=s.std(ddof=1) if len(s) > 1 else None,
             min=s.min(), p25=s.quantile(.25), p75=s.quantile(.75), max=s.max())
    r["amplitude"] = r["max"] - r["min"]
    r["cv"] = r["dp"] / r["media"] if r["dp"] is not None and r["media"] else None
    return {k: (None if v is None or pd.isna(v) else round(float(v), 4)) for k, v in r.items()}


# ---------- planilhas ----------
def estilo(ws, n_cols, congelar="C2"):
    for c in range(1, n_cols + 1):
        cel = ws.cell(row=1, column=c)
        cel.font = Font(name=FONTE, bold=True)
        cel.fill = PatternFill("solid", fgColor="DDEBF7")
        cel.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = congelar
    ws.auto_filter.ref = ws.dimensions


def fonte(ws):
    for row in ws.iter_rows(min_row=2):
        for c in row:
            if c.font is None or not c.font.bold:
                c.font = Font(name=FONTE)


def larguras(ws, padrao, especiais):
    for i, cel in enumerate(ws[1], start=1):
        ws.column_dimensions[get_column_letter(i)].width = especiais.get(cel.value, padrao)


def aba_dados(wb, d):
    ws = wb.create_sheet("dados")
    ws.append(list(d.columns))
    for r in d.itertuples(index=False):
        ws.append([None if (isinstance(v, float) and pd.isna(v)) else v for v in r])
    for i, c in enumerate(d.columns, start=1):
        fmt = VARS[c][3] if c in VARS else None
        if fmt:
            for r in range(2, len(d) + 2):
                ws.cell(row=r, column=i).number_format = fmt
    estilo(ws, len(d.columns))
    larguras(ws, 12, {"nome": 34, "metropole": 16, "regiao": 13})
    fonte(ws)
    letra = {c: get_column_letter(i + 1) for i, c in enumerate(d.columns)}
    n = len(d) + 1
    return lambda c: f"dados!${letra[c]}$2:${letra[c]}${n}"


def aba_geral(wb, R):
    ws = wb.create_sheet("geral", 0)
    cab = ["grupo", "variável", "descrição", "unidade", "n (casas com dado)", "sem dado", "média", "mediana",
           "desvio padrão", "mínimo", "1º quartil (P25)", "3º quartil (P75)", "máximo", "amplitude",
           "coef. de variação"]
    ws.append(cab)
    for i, (c, (g, rot, uni, fmt)) in enumerate(VARS.items(), start=2):
        x = R(c)
        ws.append([g, c, rot, uni, f"=COUNT({x})", f"=ROWS({x})-E{i}",
                   f"=IFERROR(AVERAGE({x}),\"\")", f"=IFERROR(MEDIAN({x}),\"\")", f"=IFERROR(STDEV({x}),\"\")",
                   f"=IFERROR(MIN({x}),\"\")", f"=IFERROR(PERCENTILE({x},0.25),\"\")",
                   f"=IFERROR(PERCENTILE({x},0.75),\"\")", f"=IFERROR(MAX({x}),\"\")",
                   f"=IFERROR(M{i}-J{i},\"\")", f"=IFERROR(I{i}/G{i},\"\")"])
        for col in "GHIJKLMN":
            ws[f"{col}{i}"].number_format = fmt
        ws[f"O{i}"].number_format = "0.00"
        ws[f"A{i}"].fill = PatternFill("solid", fgColor=CORES[g])
    estilo(ws, len(cab), "C2")
    larguras(ws, 12, {"descrição": 46, "variável": 24, "unidade": 14, "grupo": 12})
    fonte(ws)
    for row in ws.iter_rows(min_row=2, max_col=4):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")


def aba_frequencias(wb, R, d):
    ws = wb.create_sheet("frequencias", 1)
    ws.append(["tabela", "categoria", "casas", "% do total"])
    total = len(d) + 1  # linha do total de casas na planilha: COUNTA de place_id
    linhas = []

    def bloco(titulo, cats, crit_col):
        ini = ws.max_row + 1
        for cat in cats:
            r = ws.max_row + 1
            ws.append([titulo, cat, f"=COUNTIF({R(crit_col)},B{r})", f"=C{r}/COUNTA({R('place_id')})"])
        fim = ws.max_row
        r = fim + 1
        ws.append([titulo, "Total", f"=SUM(C{ini}:C{fim})", f"=SUM(D{ini}:D{fim})"])
        ws.append([titulo, "Conferência (total − casas; deve ser 0)", f"=C{r}-COUNTA({R('place_id')})", None])
        linhas.extend([r, r + 1])

    bloco("casas por região", ORDEM_REGIAO, "regiao")
    bloco("casas por metrópole", sorted(d.metropole.unique()), "metropole")
    ws.append([])
    ws.append(["proporções (não exclusivas: cada linha tem seu próprio total)", "característica", "casas com a característica", "% das casas com dado"])
    linhas.append(ws.max_row)
    for c, rot in FLAGS.items():
        r = ws.max_row + 1
        ws.append(["proporção", rot, f"=SUM({R(c)})", f"=IFERROR(SUM({R(c)})/COUNT({R(c)}),\"\")"])
    for r in range(2, ws.max_row + 1):
        ws[f"D{r}"].number_format = PCT
    estilo(ws, 4, "A2")
    larguras(ws, 16, {"tabela": 30, "categoria": 46})
    fonte(ws)
    for r in linhas:
        for c in "ABCD":
            ws[f"{c}{r}"].font = Font(name=FONTE, bold=True)


def aba_dicionario(wb):
    ws = wb.create_sheet("dicionario")
    ws.append(["coluna", "grupo", "descrição", "unidade"])
    for c, (g, rot, uni, _) in VARS.items():
        ws.append([c, g, rot, uni])
    for c, rot in FLAGS.items():
        ws.append([c, "Proporção", rot + " (1 = sim, 0 = não)", "0 ou 1"])
    ws.append(["regiao", "Identificação", "Região do Brasil da capital da metrópole (IBGE)", "—"])
    ws.append(["—", "Regras", "Cada casa conta uma vez, na metrópole do endereço. Instagram de LudoCafé e Dominaria "
               "vazio (perfil coletado não é o da casa). Couvert só nas casas que cobram. Shoppings de porte só onde "
               "o porte foi buscado. Distâncias em linha reta (coletor de entorno). Desvio padrão amostral; quartis "
               "pelo método inclusivo.", "—"])
    estilo(ws, 4, "A2")
    larguras(ws, 18, {"descrição": 90, "coluna": 26})
    fonte(ws)
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")


# Indicadores do painel por grupo: (cabeçalho, tipo, coluna, formato)
PAINEL = [
    ("nota média (mediana)", "mediana", "avaliacao_media", DEC2),
    ("avaliações no Google (mediana)", "mediana", "volume_avaliacoes", INT),
    ("horas abertas por semana (média)", "media", "horas_semana", DEC1),
    ("fechamento mais tarde (mediana, h)", "mediana", "fecha_tarde_h", DEC1),
    ("abre domingo (%)", "media", "abre_domingo", PCT),
    ("abre segunda (%)", "media", "abre_segunda", PCT),
    ("metrô até 1000 m (%)", "media", "tem_metro_1000m", PCT),
    ("seguidores (mediana)", "mediana", "seguidores", INT),
    ("posts no período (média)", "media", "n_posts_periodo", DEC1),
    ("interações por post (mediana)", "mediana", "interacoes_por_post", DEC1),
    ("com cardápio (casas)", "soma", "tem_cardapio", INT),
    ("cesta mediana (mediana, R$)", "mediana", "cesta_mediana", REAIS),
    ("cesta mediana (mín., R$)", "min", "cesta_mediana", REAIS),
    ("cesta mediana (máx., R$)", "max", "cesta_mediana", REAIS),
    ("avaliações no B2 (soma)", "soma", "avaliacoes_b2", INT),
]


def cond(R, grp, cel, c):
    return f"({R(grp)}={cel})*({R(c)}<>\"\")"


def formula(tipo, R, grp, cel, c):
    x = R(c)
    if tipo == "media":
        return f"=IFERROR(AVERAGEIFS({x},{R(grp)},{cel}),\"\")"
    if tipo == "soma":
        return f"=SUMIF({R(grp)},{cel},{x})"
    if tipo == "n":
        return f"=COUNTIFS({R(grp)},{cel},{x},\"<>\")"
    if tipo == "min":
        return f"=IF(COUNTIFS({R(grp)},{cel},{x},\"<>\")=0,\"\",_xlfn.MINIFS({x},{R(grp)},{cel}))"
    if tipo == "max":
        return f"=IF(COUNTIFS({R(grp)},{cel},{x},\"<>\")=0,\"\",_xlfn.MAXIFS({x},{R(grp)},{cel}))"
    fn = {"mediana": "MEDIAN", "dp": "STDEV"}[tipo]
    return f"=IFERROR({fn}(IF({cond(R, grp, cel, c)},{x})),\"\")"


def escrever_formula(ws, ref, f):
    ws[ref] = ArrayFormula(ref, f) if ("MEDIAN(IF" in f or "STDEV(IF" in f) else f


def aba_painel(wb, R, grp, grupos, nome, rotulo):
    ws = wb.create_sheet(nome)
    cab = [rotulo, *(["região"] if grp == "metropole" else []), "casas", "% das casas", *[p[0] for p in PAINEL]]
    ws.append(cab)
    desl = 2 if grp == "metropole" else 1
    for i, g in enumerate(grupos, start=2):
        ws.append([g, *([REGIAO[g]] if grp == "metropole" else [])])
        ws.cell(row=i, column=desl + 1, value=f"=COUNTIF({R(grp)},A{i})")
        cn = get_column_letter(desl + 1)
        ws.cell(row=i, column=desl + 2, value=f"={cn}{i}/COUNTA({R('place_id')})").number_format = PCT
        for j, (_, tipo, c, fmt) in enumerate(PAINEL, start=desl + 3):
            ref = f"{get_column_letter(j)}{i}"
            escrever_formula(ws, ref, formula(tipo, R, grp, f"$A{i}", c))
            ws[ref].number_format = fmt
    t = len(grupos) + 2
    cn, cp = get_column_letter(desl + 1), get_column_letter(desl + 2)
    ws.cell(row=t, column=1, value="Total")
    ws.cell(row=t, column=desl + 1, value=f"=SUM({cn}2:{cn}{t - 1})")
    ws.cell(row=t, column=desl + 2, value=f"=SUM({cp}2:{cp}{t - 1})").number_format = PCT
    ws.cell(row=t + 1, column=1, value="Conferência (total − casas da aba dados; deve ser 0)")
    ws.cell(row=t + 1, column=desl + 1, value=f"={cn}{t}-COUNTA({R('place_id')})")
    estilo(ws, len(cab), get_column_letter(desl + 1) + "2")
    larguras(ws, 14, {rotulo: 18, "região": 13})
    fonte(ws)
    for c in range(1, len(cab) + 1):
        ws.cell(row=t, column=c).font = Font(name=FONTE, bold=True)
    ws.cell(row=t + 1, column=1).font = Font(name=FONTE, italic=True)
    ws.cell(row=t + 1, column=1).alignment = Alignment(wrap_text=True)
    ws.row_dimensions[1].height = 48


def aba_detalhe(wb, R, grp, grupos, nome, rotulo):
    ws = wb.create_sheet(nome)
    cab = [rotulo, "grupo", "variável", "descrição", "n", "média", "mediana", "desvio padrão", "mínimo", "máximo",
           "amplitude"]
    ws.append(cab)
    r = 2
    for g in grupos:
        for c, (gr, rot, _, fmt) in VARS.items():
            ws.append([g, gr, c, rot])
            for col, tipo in zip("EFGHIJ", ["n", "media", "mediana", "dp", "min", "max"]):
                escrever_formula(ws, f"{col}{r}", formula(tipo, R, grp, f"$A{r}", c))
                ws[f"{col}{r}"].number_format = fmt if col != "E" else INT
            ws[f"K{r}"] = f"=IFERROR(J{r}-I{r},\"\")"
            ws[f"K{r}"].number_format = fmt
            ws[f"B{r}"].fill = PatternFill("solid", fgColor=CORES[gr])
            r += 1
    estilo(ws, len(cab), "E2")
    larguras(ws, 12, {rotulo: 16, "variável": 24, "descrição": 44})
    fonte(ws)


# ---------- site ----------
def site(d):
    num = lambda s: pd.to_numeric(s, errors="coerce")
    geral = {c: {"grupo": g, "rotulo": rot, "unidade": uni, **resumo(d[c])} for c, (g, rot, uni, _) in VARS.items()}

    def painel(grp, grupos):
        out = []
        for g in grupos:
            s = d[d[grp] == g]
            item = {"grupo": g, "casas": int(len(s)), "pct": round(len(s) / len(d), 4)}
            if grp == "metropole":
                item["regiao"] = REGIAO[g]
            for c in ["avaliacao_media", "volume_avaliacoes", "seguidores", "interacoes_por_post", "cesta_mediana",
                      "fecha_tarde_h", "horas_semana", "n_posts_periodo", "dist_metro_m", "bebida_mediana",
                      "entrada_ou_prato_mediana", "couvert_mediana"]:
                item[c] = resumo(s[c])
            for c in ["tem_metro_1000m", "abre_domingo", "abre_segunda"]:
                item[c] = round(float(s[c].mean()), 4)
            item["com_cardapio"] = int(s.tem_cardapio.sum())
            item["avaliacoes_b2"] = int(num(s.avaliacoes_b2).fillna(0).sum())
            out.append(item)
        return out

    metro = painel("metropole", d.metropole.value_counts().index.tolist())
    reg = painel("regiao", ORDEM_REGIAO)
    assert sum(m["casas"] for m in metro) == len(d) == sum(r["casas"] for r in reg)

    pontos = lambda cols: [
        {"nome": r.nome, "metropole": r.metropole, "regiao": r.regiao, **{c: (None if pd.isna(getattr(r, c)) else float(getattr(r, c))) for c in cols}}
        for r in d.itertuples() if all(pd.notna(getattr(r, c)) for c in cols)]

    def faixas(s, passo, ini=None):
        s = num(s).dropna()
        a = int((s.min() // passo) * passo) if ini is None else ini
        out = []
        while a <= s.max():
            out.append({"faixa": f"{a}–{a + passo}", "casas": int(((s >= a) & (s < a + passo)).sum())})
            a += passo
        assert sum(o["casas"] for o in out) == len(s)
        return out

    def contagem_hora(col):
        s = num(d[col]).dropna() // 1  # hora cheia: 23h30 conta em 23h
        return [{"hora": f"{int(h) % 24}h", "casas": int(n)} for h, n in s.value_counts().sort_index().items()]

    dados = {
        "_leiaMe": "Gerado por _src/[B1] Estabelecimentos e cardapios/Google Places (Negócios)/analise/B1-a_estatistica_descritiva.py. Não editar à mão: rode o script de novo.",
        "atualizado": DATA_REF.isoformat(),
        "total_casas": int(len(d)),
        "n_metropoles": int(d.metropole.nunique()),
        "geral": geral,
        "dias": [{"dia": NOME_DIA[x], "pct": round(float(d[f"abre_{x}"].mean()), 4),
                  "casas": int(d[f"abre_{x}"].sum())} for x in DIAS],
        "proporcoes": {c: {"rotulo": FLAGS[c], "casas": int(num(d[c]).sum()), "n": int(num(d[c]).notna().sum())}
                       for c in FLAGS},
        "hist_cesta": faixas(d.cesta_mediana, 20),
        "hist_horas": faixas(d.horas_semana, 10),
        "hist_nota": [{"faixa": f"{a:.1f}", "casas": int((num(d.avaliacao_media).round(1) == round(a, 1)).sum())}
                      for a in [x / 10 for x in range(int(num(d.avaliacao_media).min() * 10), 51)]],
        "abertura": contagem_hora("abre_cedo_h"),
        "fechamento": contagem_hora("fecha_tarde_h"),
        "dispersao_instagram": pontos(["seguidores", "interacoes_por_post"]),
        "dispersao_google": pontos(["volume_avaliacoes", "avaliacao_media"]),
        "cestas": sorted(({"nome": r.nome, "metropole": r.metropole, "regiao": r.regiao,
                           "baixa": float(r.cesta_baixa), "mediana": float(r.cesta_mediana), "alta": float(r.cesta_alta)}
                          for r in d.itertuples() if pd.notna(r.cesta_mediana)), key=lambda x: x["mediana"]),
        "metropoles": metro,
        "regioes": reg,
    }
    SAIDA_SITE.write_text(json.dumps(dados, ensure_ascii=False, indent=1), encoding="utf-8")


def main():
    d = montar_dados()
    PASTA.mkdir(parents=True, exist_ok=True)

    wb = Workbook()
    wb.remove(wb.active)
    R = aba_dados(wb, d)
    aba_geral(wb, R)
    aba_frequencias(wb, R, d)
    aba_dicionario(wb)
    wb.save(SAIDA_GERAL)

    wb = Workbook()
    wb.remove(wb.active)
    R = aba_dados(wb, d)
    metros = sorted(d.metropole.unique())
    aba_painel(wb, R, "regiao", ORDEM_REGIAO, "regiao_painel", "região")
    aba_painel(wb, R, "metropole", metros, "metropole_painel", "metrópole")
    aba_detalhe(wb, R, "regiao", ORDEM_REGIAO, "regiao_detalhe", "região")
    aba_detalhe(wb, R, "metropole", metros, "metropole_detalhe", "metrópole")
    wb.move_sheet("dados", offset=len(wb.sheetnames))
    wb.save(SAIDA_GRUPO)

    site(d)
    print(f"{len(d)} casas, {d.metropole.nunique()} metrópoles, {len(VARS)} variáveis + {len(FLAGS)} proporções")
    print(d.regiao.value_counts().reindex(ORDEM_REGIAO).to_string())
    print("soma por região:", d.regiao.value_counts().sum(), "| soma por metrópole:", d.metropole.value_counts().sum())
    for p in (SAIDA_GERAL, SAIDA_GRUPO, SAIDA_SITE):
        print("->", p.relative_to(ROOT))


if __name__ == "__main__":
    main()
