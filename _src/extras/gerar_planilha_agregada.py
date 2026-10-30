"""
Planilhas com os dados agregados da pesquisa (B1, B3 e B4) e gráficos, uma por bloco.

Lê só arquivos processados de `dados/` e grava em `dados/Dados agregados/`:
`B1_dados_agregados_completo.xlsx`, `B1_dados_agregados_resumido.xlsx` (sem
eventos, dias e temas do Instagram e sem o custo do acervo),
`B3_dados_agregados.xlsx` e `B4_dados_agregados.xlsx`, uma aba por tabela,
com a figura ao lado, e os gráficos em PNG (300 dpi) em `graficos/`. Monta
tudo numa pasta só e depois separa, renumerando as tabelas de cada arquivo.

Fora por pedido do autor (30/10/2026): os comentários do B2 (texto, temas e
contagens de avaliações) e, no entorno, tudo o que vem do OpenStreetMap
(`shopping_porte_*` e `B1_entorno.csv`). Do entorno ficam só as leituras do
coletor do Google Maps (estações de metrô e distância ao shopping).

Tabelas do B1 usam fórmulas sobre a aba "Dados B1" (54 casas, valores do
`B1_estatistica_descritiva_geral.xlsx`); as de Instagram, B3 e B4 trazem os
valores das análises já feitas, com fonte no rodapé. Paleta e desenho dos
gráficos seguem o desenho do documento de referência do autor (rótulos em
negrito na ponta da barra; legenda acima do gráfico), com cores só das paletas
slate e purple do Tailwind CSS.

Uso: python3 _src/extras/gerar_planilha_agregada.py
"""

from __future__ import annotations

import sys
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl.styles import Alignment, Border
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.formula import ArrayFormula

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / "_src" / "extras"))
from log_execucao import registrar_execucao  # noqa: E402
import planilha_estilo as est  # noqa: E402
from planilha_estilo import *  # noqa: E402,F403

D = ROOT / "dados"
B1 = D / "[B1] Estabelecimentos e cardapios"
B3 = D / "[B3] Contagem direta de ocupacao"
B4 = D / "[B4] Bases publicas"
SAIDA = D / "Dados agregados"
GRAF = SAIDA / "graficos"

# ---------------------------------------------------------------- B1: base
registrar_execucao(__file__)
est.GRAF = GRAF

ARQ_DESC = B1 / "estabelecimentos" / "4_estatistica_descritiva" / "B1_estatistica_descritiva_geral.xlsx"
base = pd.read_excel(ARQ_DESC, sheet_name="dados")
COLS = [
    ("nome", "casa", "txt"), ("metropole", "metrópole", "txt"), ("regiao", "região", "txt"),
    ("avaliacao_media", "nota no Google", "dec"), ("volume_avaliacoes", "avaliações no Google", "int"),
    ("dias_abertos", "dias abertos", "int"), ("horas_semana", "horas por semana", "int"),
    ("abre_cedo_h", "abre mais cedo (h)", "dec"), ("fecha_tarde_h", "fecha mais tarde (h)", "dec"),
    ("abre_segunda", "abre segunda", "int"), ("abre_terca", "abre terça", "int"), ("abre_quarta", "abre quarta", "int"),
    ("abre_quinta", "abre quinta", "int"), ("abre_sexta", "abre sexta", "int"), ("abre_sabado", "abre sábado", "int"),
    ("abre_domingo", "abre domingo", "int"),
    ("metro_500m", "metrô até 500 m", "int"), ("metro_1000m", "metrô até 1000 m", "int"),
    ("dist_metro_m", "distância ao metrô (m)", "int"), ("dist_shopping_m", "distância ao shopping (m)", "int"),
    ("seguidores", "seguidores", "int"), ("n_posts_total", "posts desde o início", "int"),
    ("idade_perfil_anos", "idade do perfil (anos)", "dec"), ("n_posts_periodo", "posts 01/09 a 01/10/2026", "int"),
    ("interacoes_por_post", "interações por post", "dec"), ("engajamento_pct", "engajamento", "pct"),
    ("itens_com_preco", "itens com preço", "int"), ("tem_cardapio", "cardápio transcrito", "int"),
    ("cobra_couvert", "couvert encontrado", "int"),
    ("couvert_mediana", "couvert (R$)", "dec"), ("bebida_mediana", "bebida (R$)", "dec"),
    ("entrada_ou_prato_mediana", "entrada ou prato (R$)", "dec"), ("cesta_baixa", "cesta baixa (R$)", "dec"),
    ("cesta_mediana", "cesta mediana (R$)", "dec"), ("cesta_alta", "cesta alta (R$)", "dec"),
    ("hamburguer_mediana", "hambúrguer (R$)", "dec"), ("batata_mediana", "batata (R$)", "dec"),
    ("cerveja_mediana", "cerveja (R$)", "dec"), ("refrigerante_mediana", "refrigerante (R$)", "dec"),
    ("drinks_mediana", "drinks (R$)", "dec"), ("titulos_acervo", "títulos no acervo", "int"),
]
base = base.sort_values(["regiao", "metropole", "nome"]).reset_index(drop=True)
N = len(base)
LETRA = {c: get_column_letter(i) for i, (c, _, _) in enumerate(COLS, 1)}
ABA_BASE = "Dados B1"


def R(col):
    return f"'{ABA_BASE}'!${LETRA[col]}$3:${LETRA[col]}${N + 2}"


ORDEM_REG = ["Norte", "Nordeste", "Centro-Oeste", "Sudeste", "Sul"]


def mediana_se(col, crit_col, crit_ref):
    return ArrayFormula("A1", f'=IFERROR(MEDIAN(IF({R(crit_col)}={crit_ref},IF({R(col)}<>"",{R(col)}))),"")')


FONTE_B1 = "B1, Google Places API (09/2026 e 15/10/2026), coletor de entorno do Google Maps (15/10/2026), Instagram (01/09 a 01/10/2026) e cardápios"

# Tab 1: casas por região e metrópole
cont = base.groupby(["regiao", "metropole"]).size().reset_index(name="n")
cont["ord"] = cont.regiao.map(ORDEM_REG.index)
cont = cont.sort_values(["ord", "n", "metropole"], ascending=[True, False, True])
linhas_t1 = []
for k, r in enumerate(cont.itertuples(), 3):
    linhas_t1.append([r.regiao, r.metropole, f'=COUNTIF({R("metropole")},B{k})', f"=C{k}/C${3 + len(cont)}"])
tot = 3 + len(cont)
linhas_t1.append(["Total", "", f"=SUM(C3:C{tot - 1})", f"=SUM(D3:D{tot - 1})"])
linhas_t1.append(["Conferência (total − casas da base; deve ser 0)", "", f'=C{tot}-COUNTA({R("nome")})', ""])
c1 = cont.sort_values("n", ascending=False)
g1 = barras_h("01_casas_por_metropole", c1.metropole.tolist(), {"casas": c1.n.tolist()}, lambda v: br(v, 0),
              destaque="Salvador")
nova_aba("Tab 1", "Tabela 1 – Casas do B1 por região e metrópole: São Paulo concentra 19 das 54; Salvador tem uma",
         ["Região", "Metrópole", "Casas", "% do total"], linhas_t1, ["txt", "txt", "int", "pct"], FONTE_B1,
         [24, 20, 10, 12], g1, ["Nota: cada casa conta uma vez, na metrópole do endereço."])

# Tab 2: estatística descritiva geral
VARS_DESC = [
    ("Google Maps", "avaliacao_media", "Nota média no Google (1 a 5)"),
    ("Google Maps", "volume_avaliacoes", "Avaliações no Google"),
    ("Horários", "dias_abertos", "Dias abertos por semana"),
    ("Horários", "horas_semana", "Horas abertas por semana"),
    ("Horários", "abre_cedo_h", "Abertura mais cedo (hora do dia)"),
    ("Horários", "fecha_tarde_h", "Fechamento mais tarde (25 = 1h)"),
    ("Entorno (Maps)", "metro_500m", "Estações de metrô até 500 m"),
    ("Entorno (Maps)", "metro_1000m", "Estações de metrô até 1000 m"),
    ("Entorno (Maps)", "dist_metro_m", "Distância ao metrô mais próximo (m)"),
    ("Entorno (Maps)", "dist_shopping_m", "Distância ao shopping mais próximo (m)"),
    ("Instagram", "seguidores", "Seguidores"),
    ("Instagram", "n_posts_total", "Posts desde a criação do perfil"),
    ("Instagram", "idade_perfil_anos", "Idade do perfil (anos)"),
    ("Instagram", "n_posts_periodo", "Posts de 01/09 a 01/10/2026"),
    ("Instagram", "interacoes_por_post", "Interações por post"),
    ("Instagram", "engajamento_pct", "Engajamento (interações por post ÷ seguidores)"),
    ("Cardápio", "itens_com_preco", "Itens com preço no cardápio"),
    ("Cardápio", "couvert_mediana", "Couvert ou taxa de jogo (R$, casas que cobram)"),
    ("Cardápio", "bebida_mediana", "Bebida, mediana do cardápio (R$)"),
    ("Cardápio", "entrada_ou_prato_mediana", "Entrada ou prato, mediana (R$)"),
    ("Cardápio", "cesta_baixa", "Cesta por pessoa, nível baixo (R$)"),
    ("Cardápio", "cesta_mediana", "Cesta por pessoa, nível mediano (R$)"),
    ("Cardápio", "cesta_alta", "Cesta por pessoa, nível alto (R$)"),
    ("Acervo", "titulos_acervo", "Títulos no acervo publicado"),
]
linhas_t2 = []
for g, v, d in VARS_DESC:
    rg = R(v)
    linhas_t2.append([g, d, f"=COUNT({rg})", f"=ROWS({rg})-COUNT({rg})", f'=IFERROR(AVERAGE({rg}),"")',
                      f'=IFERROR(MEDIAN({rg}),"")', f'=IFERROR(STDEV({rg}),"")', f'=IFERROR(MIN({rg}),"")',
                      f'=IFERROR(PERCENTILE({rg},0.25),"")', f'=IFERROR(PERCENTILE({rg},0.75),"")',
                      f'=IFERROR(MAX({rg}),"")'])
nova_aba("Tab 2", "Tabela 2 – Estatística descritiva das 54 casas do B1",
         ["Grupo", "Variável", "n (com dado)", "Sem dado", "Média", "Mediana", "Desvio padrão", "Mínimo",
          "1º quartil", "3º quartil", "Máximo"],
         linhas_t2, ["txt", "txt", "int", "int"] + ["dec"] * 7, FONTE_B1, [16, 42] + [11] * 9,
         notas=["Nota: desvio padrão amostral; quartis pelo método inclusivo. Engajamento em fração (0,01 = 1%).",
                "Entorno sem OpenStreetMap: shoppings de porte fora. Distância ao metrô vazia onde a cidade não tem metrô lido (Curitiba, Florianópolis, Belém, Goiânia, Vitória)."])

# Tab 3: dias de funcionamento
DIAS = [("abre_segunda", "Segunda"), ("abre_terca", "Terça"), ("abre_quarta", "Quarta"), ("abre_quinta", "Quinta"),
        ("abre_sexta", "Sexta"), ("abre_sabado", "Sábado"), ("abre_domingo", "Domingo")]
linhas_t3 = [[d, f"=COUNTIF({R(c)},1)", f"=COUNT({R(c)})", f"=B{k}/C{k}"] for k, (c, d) in enumerate(DIAS, 3)]
pct_dias = [base[c].sum() / base[c].count() for c, _ in DIAS]
g3 = barras_v("03_dias_funcionamento", [d for _, d in DIAS], {"% das casas que abrem": pct_dias},
              lambda v: br(v * 100, 1, suf="%"), eixo_pct=True, ymax=1.12)
nova_aba("Tab 3", "Tabela 3 – Dias de funcionamento: todas abrem de quinta a sábado; só 28% abrem na segunda",
         ["Dia", "Casas que abrem", "Casas com horário", "% das casas"], linhas_t3, ["txt", "int", "int", "pct"],
         "B1, Google Places API (regularOpeningHours, 15/10/2026)", [18, 14, 14, 12], g3,
         ["Nota: contagem não exclusiva (cada casa abre em vários dias); cada linha tem o próprio total."])

# Tab 4 e 5: painéis por região e por metrópole
PAINEL_CAB = ["casas", "nota (mediana)", "horas abertas por semana (média)", "abre domingo (%)",
              "abre segunda (%)", "metrô até 1000 m (% das casas)", "seguidores (mediana)",
              "casas com cesta", "cesta mediana (mediana, R$)"]
PAINEL_TIPOS = ["int", "dec", "dec", "pct", "pct", "pct", "int", "int", "dec"]


def painel(chave, ref):
    return [
        f'=COUNTIF({R(chave)},{ref})',
        mediana_se("avaliacao_media", chave, ref),
        f'=IFERROR(AVERAGEIFS({R("horas_semana")},{R(chave)},{ref}),"")',
        f'=IFERROR(AVERAGEIFS({R("abre_domingo")},{R(chave)},{ref}),"")',
        f'=IFERROR(AVERAGEIFS({R("abre_segunda")},{R(chave)},{ref}),"")',
        f'=IFERROR(COUNTIFS({R(chave)},{ref},{R("metro_1000m")},">0")/COUNTIF({R(chave)},{ref}),"")',
        mediana_se("seguidores", chave, ref),
        f'=COUNTIFS({R(chave)},{ref},{R("cesta_mediana")},"<>")',
        mediana_se("cesta_mediana", chave, ref),
    ]


linhas_t4 = [[r] + painel("regiao", f"$A{k}") for k, r in enumerate(ORDEM_REG, 3)]
linhas_t4.append(["Total", f"=SUM(B3:B7)", f'=MEDIAN({R("avaliacao_media")})', f'=AVERAGE({R("horas_semana")})',
                  f'=AVERAGE({R("abre_domingo")})', f'=AVERAGE({R("abre_segunda")})',
                  f'=COUNTIF({R("metro_1000m")},">0")/COUNTA({R("nome")})', f'=MEDIAN({R("seguidores")})',
                  f'=COUNT({R("cesta_mediana")})', f'=MEDIAN({R("cesta_mediana")})'])
por_reg = base.groupby("regiao")
g4 = barras_h("04_regiao_funcionamento_metro", ORDEM_REG, {
    "abre domingo": [por_reg.abre_domingo.mean()[r] for r in ORDEM_REG],
    "abre segunda": [por_reg.abre_segunda.mean()[r] for r in ORDEM_REG],
    "metrô até 1000 m": [(por_reg.metro_1000m.apply(lambda s: (s > 0).sum()) / por_reg.size())[r] for r in ORDEM_REG],
}, lambda v: br(v * 100, 1, suf="%"), eixo_pct=True, xmax=1.18)
nova_aba("Tab 4", "Tabela 4 – Painel por região: o Sudeste reúne 31 das 54 casas; no Sul nenhuma tem metrô a até 1 km",
         ["Região"] + PAINEL_CAB, linhas_t4, ["txt"] + PAINEL_TIPOS, FONTE_B1, [16] + [13] * 9, g4,
         ["Nota: medianas calculadas só entre as casas com dado; cesta só nas casas com couvert, bebida e prato encontrados."])

metros = cont.sort_values(["ord", "metropole"])
linhas_t5 = [[r.metropole, r.regiao] + painel("metropole", f"$A{k}") for k, r in enumerate(metros.itertuples(), 3)]
fim5 = 2 + len(linhas_t5)
linhas_t5.append(["Total", "", f"=SUM(C3:C{fim5})", f'=MEDIAN({R("avaliacao_media")})', f'=AVERAGE({R("horas_semana")})',
                  f'=AVERAGE({R("abre_domingo")})', f'=AVERAGE({R("abre_segunda")})',
                  f'=COUNTIF({R("metro_1000m")},">0")/COUNTA({R("nome")})', f'=MEDIAN({R("seguidores")})',
                  f'=COUNT({R("cesta_mediana")})', f'=MEDIAN({R("cesta_mediana")})'])
cm = base.groupby("metropole").cesta_mediana.median().dropna().sort_values(ascending=False)
g5 = barras_h("05_cesta_por_metropole", cm.index.tolist(), {"cesta mediana": cm.tolist()},
              lambda v: br(v, 2, pre="R$ "), destaque="Salvador")
nova_aba("Tab 5", "Tabela 5 – Painel por metrópole: a cesta por pessoa varia de metrópole para metrópole",
         ["Metrópole", "Região"] + PAINEL_CAB, linhas_t5, ["txt", "txt"] + PAINEL_TIPOS, FONTE_B1,
         [18, 14] + [13] * 9, g5,
         ["Nota: o gráfico mostra a mediana da cesta mediana das casas com cesta em cada metrópole; metrópoles sem cesta ficam fora."])

# Tab 6: entorno (só Google Maps)
FAIXAS = [("Até 500 m", '"<=500"', None), ("De 501 a 1000 m", '">500"', '"<=1000"'),
          ("De 1001 a 2000 m", '">1000"', '"<=2000"'), ("Mais de 2000 m", '">2000"', None)]
linhas_t6 = []
for k, (rot, a, b) in enumerate(FAIXAS, 3):
    cel = []
    for col in ("dist_metro_m", "dist_shopping_m"):
        cel.append(f"=COUNTIFS({R(col)},{a}" + (f",{R(col)},{b})" if b else ")"))
    linhas_t6.append([rot, cel[0], f"=B{k}/B$9", cel[1], f"=D{k}/D$9"])
linhas_t6.append(["Sem leitura (cidade sem metrô ou leitura a mais de 50 km)",
                  f'=COUNTA({R("nome")})-COUNT({R("dist_metro_m")})', "=B7/B$9",
                  f'=COUNTA({R("nome")})-COUNT({R("dist_shopping_m")})', "=D7/D$9"])
linhas_t6.append(["Conferência (total − casas; deve ser 0)", f'=B9-COUNTA({R("nome")})', "",
                  f'=D9-COUNTA({R("nome")})', ""])
linhas_t6.append(["Total", "=SUM(B3:B7)", "=SUM(C3:C7)", "=SUM(D3:D7)", "=SUM(E3:E7)"])


def faixa(s):
    s = s.dropna()
    return [(s <= 500).sum(), ((s > 500) & (s <= 1000)).sum(), ((s > 1000) & (s <= 2000)).sum(), (s > 2000).sum()]


fm, fs = faixa(base.dist_metro_m), faixa(base.dist_shopping_m)
nm, ns = N - base.dist_metro_m.count(), N - base.dist_shopping_m.count()
g6 = empilhado_100("06_entorno_distancias", ["Estação de metrô", "Shopping"],
                   {rot: [fm[i] if i < 4 else nm, fs[i] if i < 4 else ns]
                    for i, rot in enumerate([f[0] for f in FAIXAS] + ["Sem leitura"])})
nova_aba("Tab 6", "Tabela 6 – Entorno pelo Google Maps: a maioria das casas tem shopping a até 1 km; metrô perto é minoria",
         ["Distância (linha reta)", "Metrô: casas", "Metrô: %", "Shopping: casas", "Shopping: %"], linhas_t6,
         ["txt", "int", "pct", "int", "pct"], "B1, coletor de entorno do Google Maps (15/10/2026)",
         [44, 13, 11, 14, 12], g6,
         ["Nota: sem dados do OpenStreetMap (porte do shopping fora). Shopping de qualquer porte lido pelo Maps."])

# Tab 7: preços do cardápio
CATS = [("couvert_mediana", "Couvert ou taxa de jogo (casas que cobram)"), ("bebida_mediana", "Bebida"),
        ("entrada_ou_prato_mediana", "Entrada ou prato"), ("hamburguer_mediana", "Hambúrguer"),
        ("batata_mediana", "Batata"), ("cerveja_mediana", "Cerveja"), ("refrigerante_mediana", "Refrigerante"),
        ("drinks_mediana", "Drinks"), ("cesta_baixa", "Cesta por pessoa, nível baixo"),
        ("cesta_mediana", "Cesta por pessoa, nível mediano"), ("cesta_alta", "Cesta por pessoa, nível alto")]
linhas_t7 = [[d, f"=COUNT({R(c)})", f'=IFERROR(PERCENTILE({R(c)},0.25),"")', f'=IFERROR(MEDIAN({R(c)}),"")',
              f'=IFERROR(PERCENTILE({R(c)},0.75),"")'] for c, d in CATS]
g7 = barras_h("07_precos_cardapio", [d for _, d in CATS], {"mediana entre as casas": [base[c].median() for c, _ in CATS]},
              lambda v: br(v, 2, pre="R$ "), larg_rotulo=26)
nova_aba("Tab 7", "Tabela 7 – Preços nos cardápios: a cesta mediana por pessoa fica em R$ 54 (metade das casas entre R$ 38,50 e R$ 70)",
         ["Categoria", "Casas com preço", "1º quartil (R$)", "Mediana (R$)", "3º quartil (R$)"], linhas_t7,
         ["txt", "int", "dec", "dec", "dec"], "B1, cardápios transcritos (B1_cardapios_itens_precos.xlsx)",
         [42, 14, 14, 14, 14], g7,
         ["Nota: cada casa entra com a mediana do próprio cardápio na categoria; quartis entre as casas.",
          "Cesta = couvert + bebida + entrada ou prato, só nas casas em que os três foram encontrados."])

# Tab 8: couvert
linhas_t8 = [
    ["Couvert ou taxa de jogo encontrado no cardápio", f'=COUNTIFS({R("tem_cardapio")},1,{R("cobra_couvert")},1)', "=B3/B$6"],
    ["Couvert não encontrado no cardápio", f'=COUNTIF({R("tem_cardapio")},1)-B3', "=B4/B$6"],
    ["Sem cardápio transcrito", f'=COUNTA({R("nome")})-COUNTIF({R("tem_cardapio")},1)', "=B5/B$6"],
    ["Total", "=SUM(B3:B5)", "=SUM(C3:C5)"],
    ["Conferência (total − casas; deve ser 0)", f'=B6-COUNTA({R("nome")})', ""],
]
enc = int(((base.tem_cardapio == 1) & (base.cobra_couvert == 1)).sum())
nenc = int((base.tem_cardapio == 1).sum()) - enc
sem = N - enc - nenc
g8 = empilhado_100("08_couvert", ["Casas do B1"], {"Encontrado": [enc], "Não encontrado": [nenc], "Sem cardápio": [sem]})
nova_aba("Tab 8", "Tabela 8 – Couvert: encontrado no cardápio de 8 das 36 casas com cardápio",
         ["Situação", "Casas", "% do total"], linhas_t8, ["txt", "int", "pct"],
         "B1, cardápios transcritos (B1_cardapios_itens_precos.xlsx)", [46, 10, 12], g8,
         ["Nota: \"não encontrado\" não quer dizer que a casa não cobra; o couvert pode estar fora do cardápio coletado."])

# ---------------------------------------------------------------- B1: Instagram
LEG = B1 / "redes_sociais" / "analise_legendas"
FONTE_IG = "B1, legendas do Instagram das casas (coletor do Instagram, 2026), análise de legendas"


def n_de(t):
    a, b = str(t).split(" de ")
    return int(a), int(b)


mc = pd.read_csv(LEG / "precos" / "B1_instagram_modelos_cobranca.csv", sep=";")
linhas_t9, gj, gl = [], [], []
for k, r in enumerate(mc.itertuples(index=False), 3):
    a, b = n_de(r[3])
    c, d = n_de(r[4])
    linhas_t9.append([r[0], r[1], r[2], a, b, f"=D{k}/E{k}", c, d, f"=G{k}/H{k}"])
    gj.append(a / b)
    gl.append(c / d)
ordem = np.argsort(gj)[::-1]
g9 = barras_h("09_instagram_modelos_cobranca", [mc.iloc[i, 0] for i in ordem],
              {"Casas de jogos": [gj[i] for i in ordem], "Lojas com espaço de jogo": [gl[i] for i in ordem]},
              lambda v: br(v * 100, 1, suf="%"), eixo_pct=True, xmax=1.2, larg_rotulo=26)
nova_aba("Tab 9", "Tabela 9 – Modelos de cobrança citados no Instagram: couvert lúdico e inscrição em evento são os mais comuns",
         ["Modelo", "Posts que citam", "Estabelecimentos que citam", "Casas de jogos que citam", "Casas de jogos (total)",
          "Casas de jogos (%)", "Lojas que citam", "Lojas (total)", "Lojas (%)"], linhas_t9,
         ["txt", "int", "int", "int", "int", "pct", "int", "int", "pct"], FONTE_IG, [36] + [12] * 8, g9,
         ["Nota: contagem não exclusiva (um post pode citar vários modelos)."])

ev = pd.read_csv(LEG / "eventos" / "B1_instagram_eventos_por_categoria.csv", sep=";")
linhas_t10 = []
for k, r in enumerate(ev.itertuples(index=False), 3):
    a, b = n_de(r[3])
    c, d = n_de(r[5])
    linhas_t10.append([r[0], r[1], r[2], a, b, r[4], c, d, r[6]])
evg = ev[~ev.iloc[:, 0].str.contains("todos|sorteio")].sort_values(ev.columns[4], ascending=False)
g10 = barras_h("10_instagram_eventos_categoria", evg.iloc[:, 0].tolist(),
               {"Casas de jogos": evg.iloc[:, 4].tolist(), "Lojas com espaço de jogo": evg.iloc[:, 6].tolist()},
               lambda v: br(v, 2))
nova_aba("Tab 10", "Tabela 10 – Eventos no Instagram: encontros/RPG, competições e noites temáticas puxam a agenda",
         ["Categoria", "Posts", "Estabelecimentos", "Casas de jogos com o evento", "Casas de jogos (total)",
          "Casas de jogos: posts por casa por mês", "Lojas com o evento", "Lojas (total)", "Lojas: posts por casa por mês"],
         linhas_t10, ["txt", "int", "int", "int", "int", "dec", "int", "int", "dec"], FONTE_IG, [30] + [13] * 8, g10,
         ["Nota: categorias não exclusivas; o gráfico mostra posts por casa por mês, sem \"todos os eventos\" e sem sorteio (que não é evento)."])

dw = pd.read_csv(LEG / "eventos" / "B1_instagram_eventos_dias_da_semana.csv", sep=";")
linhas_t11 = [[r[0], r[2], r[3] / 100, r[4], r[5], r[6] / 100, r[7]] for r in dw.itertuples(index=False)]
g11 = barras_v("11_instagram_eventos_dia", [d.capitalize() for d in dw.iloc[:, 0]],
               {"Casas de jogos": (dw.iloc[:, 3] / 100).tolist(), "Lojas com espaço de jogo": (dw.iloc[:, 6] / 100).tolist()},
               lambda v: br(v * 100, 1, suf="%"), eixo_pct=True)
nova_aba("Tab 11", "Tabela 11 – Dia citado nos eventos: casas de jogos anunciam mais domingo e quinta; lojas, sábado",
         ["Dia", "Casas de jogos: eventos que citam", "Casas de jogos: % dos eventos", "Casas de jogos: estabelecimentos",
          "Lojas: eventos que citam", "Lojas: % dos eventos", "Lojas: estabelecimentos"], linhas_t11,
         ["txt", "int", "pct", "int", "int", "pct", "int"], FONTE_IG, [14] + [14] * 6, g11,
         ["Nota: % sobre os eventos prováveis de cada tipo de casa; um evento pode citar mais de um dia ou nenhum."])

tf = pd.read_csv(LEG / "temas" / "B1_instagram_temas_frequencia.csv", sep=";")
tf = tf.iloc[:, [0, 1, 2, 3, 4, 5]]  # só legendas; colunas do B2 ficam fora
linhas_t12 = [[r[0], r[1], r[2], r[3] / 100, r[4] / 100, r[5]] for r in tf.itertuples(index=False)]
tg = tf.sort_values(tf.columns[4], ascending=False)
g12 = barras_h("12_instagram_temas", tg.iloc[:, 1].tolist(), {"% ponderado por estabelecimento": (tg.iloc[:, 4] / 100).tolist()},
               lambda v: br(v * 100, 1, suf="%"), eixo_pct=True, larg_rotulo=30)
nova_aba("Tab 12", "Tabela 12 – Temas das legendas do Instagram: programação, varejo e preço dominam",
         ["Código", "Tema", "Posts com o tema", "% dos posts", "% ponderado por estabelecimento", "Estabelecimentos com o tema"],
         linhas_t12, ["txt", "txt", "int", "pct", "pct", "int"], FONTE_IG, [10, 36, 12, 12, 14, 14], g12,
         ["Nota: temas não exclusivos; o % ponderado dá peso igual a cada estabelecimento. Comparação com o B2 retirada desta planilha."])

# ---------------------------------------------------------------- B1: acervo
AC = B1 / "acervo_jogos"
bj = pd.read_csv(AC / "B1_banco_jogos.csv")
cj = pd.read_excel(AC / "B1_cesta_jogos.xlsx", sheet_name="cesta")
top = bj.sort_values(["n_casas_acervo", "jogo"], ascending=[False, True]).head(20).merge(
    cj[["jogo", "preco_provavel_rs"]], on="jogo", how="left")
n_acervos = bj.acervo_no_local.dropna().str.split("; ").explode().nunique()
linhas_t13 = [[r.jogo, r.n_casas_acervo, r.jogadores_min, r.jogadores_max, r.duracao_min, r.preco_provavel_rs,
               r.aluguel_mediana_rs] for r in top.itertuples()]
g13 = barras_h("13_jogos_mais_presentes", top.jogo.tolist(), {"casas com o jogo": top.n_casas_acervo.tolist()},
               lambda v: br(v, 0))
nova_aba("Tab 13", f"Tabela 13 – Os 20 jogos mais presentes nos acervos publicados ({n_acervos} acervos): Catan, Saboteur e Potion Explosion lideram",
         ["Jogo", "Casas com o jogo", "Jogadores (mín.)", "Jogadores (máx.)", "Duração (min)", "Preço provável (R$)",
          "Aluguel (mediana, R$)"], linhas_t13, ["txt", "int", "int", "int", "int", "dec", "dec"],
         "B1, acervos publicados (acervodejogos, Ludopedia, Rejoy e listas das casas) e preços de lojas (B1_cesta_jogos.xlsx)",
         [30, 12, 12, 12, 12, 14, 14], g13, ["Nota: preço vazio = jogo sem preço nas lojas consultadas."])

ai = pd.read_excel(AC / "B1_cesta_jogos.xlsx", sheet_name="acervo_inicial")
rot_tam = {"p25": "Pequeno (1º quartil dos acervos)", "mediana": "Mediano (mediana dos acervos)",
           "p75": "Grande (3º quartil dos acervos)"}
linhas_t14 = [[rot_tam.get(r.tamanho, r.tamanho), r.titulos, r.nucleo_com_preco, f"=B{k}-C{k}", r.custo_otimista_rs,
               r.custo_provavel_rs, r.custo_pessimista_rs] for k, r in enumerate(ai.itertuples(), 3)]
g14 = barras_h("14_custo_acervo_inicial", [rot_tam.get(t, t) for t in ai.tamanho],
               {"Otimista": ai.custo_otimista_rs.tolist(), "Provável": ai.custo_provavel_rs.tolist(),
                "Pessimista": ai.custo_pessimista_rs.tolist()},
               lambda v: br(v / 1000, 1, pre="R$ ", suf=" mil"), cores=[ROXO_800, ROXO_500, SLATE_600], larg_rotulo=20)
nova_aba("Tab 14", "Tabela 14 – Custo do acervo inicial: de R$ 40 mil (acervo pequeno) a R$ 119 mil (acervo grande)",
         ["Tamanho do acervo", "Títulos", "Núcleo com preço", "Títulos restantes", "Custo otimista (R$)",
          "Custo provável (R$)", "Custo pessimista (R$)"], linhas_t14, ["txt", "int", "int", "int", "dec", "dec", "dec"],
         "B1, acervos publicados e preços de lojas (B1_cesta_jogos.xlsx, aba acervo_inicial)", [32, 10, 12, 12, 15, 15, 15],
         g14, ["Nota: títulos restantes (fora do núcleo com preço) pelo preço médio do núcleo."])

# ---------------------------------------------------------------- B3
rv = pd.read_csv(B3 / "contagem_resumo_visitas.csv", sep=";", encoding="utf-8-sig")
rv = rv[rv.observador != "media"].reset_index(drop=True)
UNID = {"BV": "Bela Vista", "PA": "Paralela"}
rot_vis = [f"{pd.to_datetime(r.visita_inicio):%d/%m} {UNID.get(r.unidade, r.unidade)}" for r in rv.itertuples()]
VARS_B3 = [("horas_observadas", "Horas observadas", "dec"), ("grupos_porta", "Grupos que entraram (porta)", "int"),
           ("pessoas_entraram", "Pessoas que entraram", "int"), ("pessoas_por_hora", "Pessoas por hora", "dec"),
           ("tamanho_grupo_porta", "Pessoas por grupo (porta)", "dec"),
           ("permanencia_porta_media_min", "Permanência média pela porta (min)", "dec"),
           ("grupos_mesa", "Grupos acompanhados na mesa", "int"), ("pessoas_mesa", "Pessoas nas mesas", "int"),
           ("tamanho_grupo_mesa", "Pessoas por grupo (mesa)", "dec"),
           ("permanencia_mesa_media_min", "Permanência média na mesa (min)", "dec"),
           ("permanencia_mesa_mediana_min", "Permanência mediana na mesa (min)", "dec")]
linhas_t15 = [[d] + [rv.loc[i, c] for i in range(len(rv))] for c, d, _ in VARS_B3]
# pessoas por hora e por grupo recalculadas por fórmula a partir das contagens
for i in range(len(rv)):
    L = get_column_letter(2 + i)
    linhas_t15[3][1 + i] = f"=IFERROR({L}5/{L}3,\"\")"
    linhas_t15[4][1 + i] = f"=IFERROR({L}5/{L}4,\"\")"
    linhas_t15[8][1 + i] = f"=IFERROR({L}10/{L}9,\"\")"
g15 = barras_v("15_b3_permanencia", rot_vis,
               {"pela porta (média)": rv.permanencia_porta_media_min.tolist(),
                "na mesa (média)": rv.permanencia_mesa_media_min.tolist(),
                "na mesa (mediana)": rv.permanencia_mesa_mediana_min.tolist()},
               lambda v: br(v, 0, suf=" min"))
nova_aba("Tab 15", "Tabela 15 – Contagem direta na São Jogue: grupos ficam de 1h20 a 2h na mesa",
         ["Medida"] + rot_vis, linhas_t15, ["txt"] + ["dec"] * len(rv),
         "B3, contagem direta nas unidades da São Jogue (app de observação, 10/10, 16/10 e 17/10/2026)",
         [38] + [16] * len(rv), g15,
         ["Nota: visitas de cerca de 3 h, do lado de fora; 10/10 só com o POV1."])
ws15 = wb["Tab 15"]
for i, (_, _, t) in enumerate(VARS_B3, 3):
    for j in range(2, 2 + len(rv)):
        ws15.cell(i, j).number_format = FMT[t]

# ---------------------------------------------------------------- B4
CAP_B1 = sorted(base.metropole.unique())
rp = pd.read_csv(B4 / "mercado_demanda" / "[B4] Resumo de População e Renda por Capital 2022-2026 (IBGE).csv")
rp = rp[rp.capital.isin(CAP_B1)].sort_values("renda_trabalho_habitual_pnadc_202602", ascending=False)
linhas_t16 = [[r.capital, r.uf, r.pop_censo_2022, r.pop_estimada_2026, r.pop_20_39_censo_2022,
               f"=E{k}/C{k}", r.renda_media_14mais_censo_2022, r.renda_mediana_14mais_censo_2022,
               r.renda_trabalho_habitual_pnadc_202602] for k, r in enumerate(rp.itertuples(), 3)]
g16 = barras_h("16_renda_trabalho_capitais", rp.capital.tolist(),
               {"renda do trabalho": rp.renda_trabalho_habitual_pnadc_202602.tolist()},
               lambda v: br(v, 0, pre="R$ "), destaque="Salvador")
nova_aba("Tab 16", "Tabela 16 – População e renda das capitais do B1: Salvador e Belém têm a menor renda do trabalho do grupo",
         ["Capital", "UF", "População 2022", "População estimada 2026", "População de 20 a 39 anos (2022)",
          "% de 20 a 39 anos", "Renda média, 14 anos ou mais (2022, R$)", "Renda mediana, 14 anos ou mais (2022, R$)",
          "Renda habitual do trabalho (2º tri. 2026, R$)"], linhas_t16,
         ["txt", "txt", "int", "int", "int", "pct", "dec", "dec", "dec"],
         "B4, IBGE: Censo 2022, estimativas de população 2026 e PNAD Contínua (2º trimestre de 2026)",
         [18, 6, 14, 14, 14, 11, 15, 15, 16], g16)

al = pd.read_csv(B4 / "financeiro" / "[B4] Aluguel Comercial por Capital 2026-07 (FipeZap).csv")
al = al[al.capital.isin(CAP_B1)].copy()
al["valor"] = al.comercial_locacao_preco_m2.fillna(al.comercial_estimado_preco_m2_mediana)
al["origem"] = np.where(al.comercial_locacao_preco_m2.notna(), "FipeZap comercial",
                        "estimado (residencial × razão comercial/residencial)")
al = al.sort_values("valor", ascending=False)
linhas_t17 = [[r.capital, r.uf, r.comercial_locacao_preco_m2, r.comercial_estimado_preco_m2_mediana,
               f"=IF(C{k}<>\"\",C{k},D{k})", r.comercial_locacao_var_12m_pct / 100 if pd.notna(r.comercial_locacao_var_12m_pct) else None,
               r.origem] for k, r in enumerate(al.itertuples(), 3)]
g17 = barras_h("17_aluguel_comercial", al.capital.tolist(), {"R$/m²": al.valor.tolist()},
               lambda v: br(v, 2, pre="R$ "), destaque="Salvador")
nova_aba("Tab 17", "Tabela 17 – Aluguel comercial por m²: Salvador está entre as capitais mais caras do grupo",
         ["Capital", "UF", "Locação comercial (R$/m²)", "Estimativa, mediana (R$/m²)", "Valor usado (R$/m²)",
          "Variação em 12 meses", "Origem do valor"], linhas_t17, ["txt", "txt", "dec", "dec", "dec", "pct", "txt"],
         "B4, FipeZap (índice de locação comercial e residencial, julho de 2026)", [18, 6, 14, 14, 13, 12, 40], g17,
         ["Nota: onde o FipeZap não cobre o comercial, estimativa pelo residencial vezes a razão comercial/residencial das capitais cobertas."])

sa = pd.read_csv(B4 / "multiuso" / "[B4] Salário Admissão por Capital 2025-2026 (CAGED).csv")
sa = sa[sa.recorte == "CNAE 56 - alimentação"]
ssa = sa[sa.capital == "Salvador"].set_index("ocupacao")
sbr = sa[sa.capital == "27 capitais"].set_index("ocupacao")
ocup = ssa.salario_mediana.sort_values(ascending=False).index.tolist()
linhas_t18 = [[o, ssa.loc[o, "admissoes"], ssa.loc[o, "salario_mediana"], sbr.loc[o, "admissoes"],
               sbr.loc[o, "salario_mediana"], f"=C{k}/E{k}", "sim" if ssa.loc[o, "amostra_pequena"] else "não"]
              for k, o in enumerate(ocup, 3)]
g18 = barras_h("18_salario_admissao", ocup, {"Salvador": [ssa.loc[o, "salario_mediana"] for o in ocup],
                                             "27 capitais": [sbr.loc[o, "salario_mediana"] for o in ocup]},
               lambda v: br(v, 0, pre="R$ "), larg_rotulo=24)
nova_aba("Tab 18", "Tabela 18 – Salário de admissão em alimentação: Salvador paga perto do salário mínimo e abaixo das capitais",
         ["Ocupação", "Salvador: admissões", "Salvador: salário mediano (R$)", "27 capitais: admissões",
          "27 capitais: salário mediano (R$)", "Salvador ÷ 27 capitais", "Amostra pequena em Salvador"], linhas_t18,
         ["txt", "int", "dec", "int", "dec", "pct", "txt"],
         "B4, Novo CAGED (admissões de 08/2025 a 07/2026, CNAE 56 – alimentação)", [34, 12, 14, 12, 14, 12, 12], g18)

sn = pd.read_csv(B4 / "juridico_tributario" / "[B4] Faixas e Alíquotas por Anexo 2026 (Simples Nacional).csv")
sn = sn[sn.anexo.isin(["I", "III"])]
linhas_t19 = [[r.anexo, r.faixa, r.rbt12_min, r.rbt12_max, r.aliquota_nominal_pct / 100, r.parcela_deduzir,
               f"=IFERROR((D{k}*E{k}-F{k})/D{k},\"\")"] for k, r in enumerate(sn.itertuples(), 3)]
x19 = [f"Faixa {f}" for f in sn[sn.anexo == "I"].faixa]
efet = lambda a: ((sn[sn.anexo == a].rbt12_max * sn[sn.anexo == a].aliquota_nominal_pct / 100 - sn[sn.anexo == a].parcela_deduzir)
                  / sn[sn.anexo == a].rbt12_max).tolist()
g19 = barras_v("19_simples_aliquota_efetiva", x19, {"Anexo I (A&B)": efet("I"), "Anexo III (couvert e eventos)": efet("III")},
               lambda v: br(v * 100, 1, suf="%"), eixo_pct=True)
nova_aba("Tab 19", "Tabela 19 – Simples Nacional, anexos I e III: alíquota efetiva no teto de cada faixa",
         ["Anexo", "Faixa", "Receita bruta em 12 meses, de (R$)", "Até (R$)", "Alíquota nominal", "Parcela a deduzir (R$)",
          "Alíquota efetiva no teto da faixa"], linhas_t19, ["txt", "int", "dec", "dec", "pct", "dec", "pct"],
         "B4, Lei Complementar 123/2006, anexos I e III (Simples Nacional, 2026)", [8, 7, 18, 16, 12, 16, 16], g19,
         ["Nota: alíquota efetiva = (receita × alíquota nominal − parcela a deduzir) ÷ receita."])

bc = pd.read_csv(B4 / "financeiro" / "[B4] Selic IPCA e Salário Mínimo 2017-2026 (BCB).csv", parse_dates=["data"])
fo = pd.read_csv(B4 / "financeiro" / "[B4] Expectativas Focus 2026-2030 (BCB).csv")
anual = {}
for serie, rot in (("selic_meta", "selic"), ("ipca_acumulado_12_meses", "ipca")):
    s = bc[bc.serie == serie].sort_values("data")
    s = s.groupby(s.data.dt.year).last()
    for ano, r in s.iterrows():
        anual.setdefault(ano, {})[rot] = r.valor
        anual[ano][f"{rot}_data"] = r.data
linhas_t20 = []
for ano in sorted(anual):
    a = anual[ano]
    linhas_t20.append([str(ano), "realizado", a.get("selic", np.nan) / 100, a.get("ipca", np.nan) / 100,
                       f"Selic {a['selic_data']:%m/%Y}; IPCA {a['ipca_data']:%m/%Y}"])
for ano in range(int(max(anual)), 2031):
    f_s = fo[(fo.Indicador == "Selic") & (fo.DataReferencia == ano)]
    f_i = fo[(fo.Indicador == "IPCA") & (fo.DataReferencia == ano)]
    if len(f_s) and len(f_i):
        linhas_t20.append([str(ano), "Focus (mediana)", f_s.Mediana.iloc[0] / 100, f_i.Mediana.iloc[0] / 100, "11/09/2026"])
real = [l for l in linhas_t20 if l[1] == "realizado"]
proj = [l for l in linhas_t20 if l[1] != "realizado"]
anos = [int(l[0]) for l in real] + [int(l[0]) + 0.0 for l in proj if int(l[0]) > int(real[-1][0])]
sel = [l[2] * 100 for l in real] + [l[2] * 100 for l in proj if int(l[0]) > int(real[-1][0])]
ipc = [l[3] * 100 for l in real] + [l[3] * 100 for l in proj if int(l[0]) > int(real[-1][0])]
g20 = linhas("20_selic_ipca", np.array(anos), {"Selic meta": sel, "IPCA em 12 meses": ipc},
             lambda v: br(v, 2, suf="%"), tracejar_depois=int(real[-1][0]))
nova_aba("Tab 20", "Tabela 20 – Selic e IPCA: juros altos em 2026, com queda esperada até 2030",
         ["Ano", "Tipo", "Selic meta", "IPCA em 12 meses", "Referência"], linhas_t20, ["txt", "txt", "pct", "pct", "txt"],
         "B4, Banco Central do Brasil: SGS (séries 432 e 13522) e Relatório Focus de 11/09/2026", [8, 16, 12, 14, 30], g20,
         ["Nota: realizado = último valor do ano (2026 até a data indicada); Focus = mediana das expectativas para o fim do ano."])

ip = pd.read_csv(B4 / "multiuso" / "[B4] IPCA Alimentação Fora do Domicílio 2023-2026 (IBGE).csv")
col_item = ip.columns[8]
ip = ip[ip.variavel.str.contains("12 meses") & ip.localidade.isin(["Brasil", "Salvador - BA"])
        & ip[col_item].isin(["Índice geral", "1201.Alimentação fora do domicílio"])]
pv = ip.pivot_table(index="periodo", columns=["localidade", col_item], values="valor").dropna().sort_index()
pv = pv[pv.index % 100 == pv.index.max() % 100]  # mesmo mês de cada ano
linhas_t21 = [[f"{str(p)[4:]}/{str(p)[:4]}", pv.loc[p, ("Brasil", "Índice geral")] / 100,
               pv.loc[p, ("Brasil", "1201.Alimentação fora do domicílio")] / 100,
               pv.loc[p, ("Salvador - BA", "Índice geral")] / 100,
               pv.loc[p, ("Salvador - BA", "1201.Alimentação fora do domicílio")] / 100] for p in pv.index]
g21 = barras_v("21_ipca_alimentacao_fora", [l[0] for l in linhas_t21],
               {"Brasil: geral": [l[1] for l in linhas_t21], "Brasil: alimentação fora": [l[2] for l in linhas_t21],
                "Salvador: geral": [l[3] for l in linhas_t21], "Salvador: alimentação fora": [l[4] for l in linhas_t21]},
               lambda v: br(v * 100, 1, suf="%"), cores=[ROXO_800, ROXO_500, ROXO_300, SLATE_600], eixo_pct=True)
nova_aba("Tab 21", "Tabela 21 – IPCA em 12 meses: alimentação fora do domicílio sobe mais que o índice geral",
         ["Mês", "Brasil: índice geral", "Brasil: alimentação fora do domicílio", "Salvador: índice geral",
          "Salvador: alimentação fora do domicílio"], linhas_t21, ["txt", "pct", "pct", "pct", "pct"],
         "B4, IBGE, IPCA (tabela 7060, variação acumulada em 12 meses)", [12, 14, 16, 14, 16], g21)

# ---------------------------------------------------------------- base e índice
ws = wb.create_sheet(ABA_BASE)
ws.sheet_view.showGridLines = False
ws.cell(1, 1, f"Base – as {N} casas do B1 (valores; as Tabelas 1 a 8 calculam por fórmula sobre esta aba)")
ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(COLS))
for j, (_, rot, _) in enumerate(COLS, 1):
    ws.cell(2, j, rot)
for i, r in enumerate(base.itertuples(index=False), 3):
    for j, (c, _, _) in enumerate(COLS, 1):
        v = getattr(r, c)
        ws.cell(i, j, None if (isinstance(v, float) and np.isnan(v)) else (v.item() if hasattr(v, "item") else v))
ult = N + 2
for rr in range(1, ult + 1):
    for j, (_, _, t) in enumerate(COLS, 1):
        cel = ws.cell(rr, j)
        cel.fill = BRANCO
        if rr == 1:
            cel.font, cel.alignment = ARIAL_N, Alignment(horizontal="left")
        elif rr == 2:
            cel.font = ARIAL_N
            cel.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cel.border = Border(top=FINA, bottom=FINA, left=SEM, right=SEM)
        else:
            cel.font = ARIAL
            cel.alignment = Alignment(horizontal="left" if t == "txt" else "right")
            if t != "txt":
                cel.number_format = FMT[t]
            cel.border = Border(top=SEM, left=SEM, right=SEM, bottom=FINA if rr == ult else SEM)
ws.row_dimensions[2].height = 42
ws.freeze_panes = "B3"
for j in range(1, len(COLS) + 1):
    ws.column_dimensions[get_column_letter(j)].width = 30 if j == 1 else 13
ws.cell(ult + 1, 1, f"Fonte: {FONTE_B1}; arquivo B1_estatistica_descritiva_geral.xlsx (aba dados).").font = ARIAL
ws.cell(ult + 2, 1, "Nota: 1 = sim, 0 = não nas colunas de dia e de cardápio; \"couvert encontrado\" 0 = não encontrado no cardápio, não \"não cobra\".").font = ARIAL

SAIDA.mkdir(parents=True, exist_ok=True)
for velho in ("dados_agregados.xlsx",):  # pasta única de antes da divisão por bloco (30/10/2026)
    (SAIDA / velho).unlink(missing_ok=True)
total = SAIDA / "_todas_as_tabelas.tmp.xlsx"
wb.save(total)
TABS = [a for a, _, _ in INDICE]
B1_TABS, B3_TABS, B4_TABS = TABS[:14], TABS[14:15], TABS[15:]
B1_FORA = {"Tab 10", "Tab 11", "Tab 12", "Tab 14"}  # eventos, dias e temas do Instagram e custo do acervo
BASE_B1 = [(ABA_BASE, "Base das 54 casas do B1", FONTE_B1)]
NOTA_GRAF = "Gráficos em PNG de 300 dpi em dados/Dados agregados/graficos/. Gerado por _src/extras/gerar_planilha_agregada.py."
NOTA_FORA = "Fora, por pedido do autor (30/10/2026): comentários do B2 (planilha própria) e dados do OpenStreetMap no entorno."
feitos = dividir(total, INDICE, [
    (SAIDA / "B1_dados_agregados_completo.xlsx", "Dados agregados do B1 (estabelecimentos) – TCC Luderia, versão completa",
     B1_TABS, BASE_B1, [NOTA_FORA, NOTA_GRAF]),
    (SAIDA / "B1_dados_agregados_resumido.xlsx", "Dados agregados do B1 (estabelecimentos) – TCC Luderia, versão resumida",
     [a for a in B1_TABS if a not in B1_FORA], BASE_B1,
     [NOTA_FORA, "Versão resumida: sem eventos por categoria, dias citados nos eventos e temas das legendas do Instagram, "
      "e sem o custo do acervo inicial (tabelas 10, 11, 12 e 14 da versão completa).", NOTA_GRAF]),
    (SAIDA / "B3_dados_agregados.xlsx", "Dados agregados do B3 (contagem direta) – TCC Luderia", B3_TABS, [], [NOTA_GRAF]),
    (SAIDA / "B4_dados_agregados.xlsx", "Dados agregados do B4 (bases públicas) – TCC Luderia", B4_TABS, [], [NOTA_GRAF]),
])
total.unlink()
for caminho, n in feitos:
    print(f"{caminho.relative_to(ROOT)}: {n} tabelas")
print(f"{len(list(GRAF.glob('*.png')))} gráficos em {GRAF.relative_to(ROOT)}")
print(f"Conferência: {N} casas na base; soma por metrópole = {int(cont.n.sum())} (diferença {N - int(cont.n.sum())})")
