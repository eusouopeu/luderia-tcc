"""
Cap. 9 - modelo financeiro da luderia (_instrucoes/8_financeiro.md).

Gera `textos/capitulos/9_plano_financeiro_modelo.xlsx`, dirigido por premissas:
abas premissas, pessoal, simples, receitas, custos, investimento, dre,
fluxo_caixa, indicadores, sensibilidade e checagens. Toda fórmula de cálculo lê
a aba premissas (coluna "Valor usado", escolhida pelo cenário ativo em
premissas!D2: 1 pessimista, 2 provável, 3 otimista) pelos nomes definidos
P_<ID> (ex.: P_R04). Nenhum número fica digitado dentro de fórmula.

O mesmo modelo é calculado em Python para o tornado (§4.2), o ponto de ruptura
(§4.3) e a tabela dos três cenários, gravados como valores na aba
sensibilidade. A planilha é recalculada no LibreOffice nos três cenários e o
VPL de cada um tem de bater com o do Python (diferença < R$ 1,00); senão o
script para.

Modelo em preços constantes (out/2026), com TMA real. Mensal nos 60 meses
(rampa de maturação no ano 1), anual na DRE e no fluxo de caixa. Sem valor
residual: no ano 5 voltam só a caução e o capital de giro.

Entradas: B1 (cesta de jogos), B3 (contagem_resumo_visitas.csv,
contagem_precos.csv, contagem_notas_intervalo.csv), B4 (FipeZap, CAGED, BCB,
Simples Nacional, IBGE PAS e população). As premissas N6 ficam em PREMISSAS,
com a justificativa da faixa.

Saídas: o .xlsx e a figura textos/capitulos/figuras/9_tornado.png; imprime os
números citados no texto do capítulo.
Uso: python3 _src/extras/gerar_modelo_financeiro.py
"""
import sys
from pathlib import Path
_AQUI = Path(__file__).resolve()
sys.path[:0] = [str(_AQUI.parent)]
from log_execucao import registrar_execucao  # noqa: E402
if __name__ == "__main__":
    registrar_execucao(__file__)

import shutil
import subprocess
import tempfile

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt  # noqa: E402
import matplotlib.ticker  # noqa: E402,F401
import openpyxl  # noqa: E402
import pandas as pd  # noqa: E402
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side  # noqa: E402
from openpyxl.utils import get_column_letter  # noqa: E402
from openpyxl.workbook.defined_name import DefinedName  # noqa: E402

ROOT = Path(__file__).resolve().parents[2]
DADOS = ROOT / "dados"
B1 = DADOS / "[B1] Estabelecimentos e cardapios"
B3 = DADOS / "[B3] Contagem direta de ocupacao"
B4 = DADOS / "[B4] Bases publicas"
SAIDA = ROOT / "textos" / "capitulos" / "9_plano_financeiro_modelo.xlsx"
FIGURA = ROOT / "textos" / "capitulos" / "figuras" / "9_tornado.png"
SOFFICE = "/Applications/LibreOffice.app/Contents/MacOS/soffice"

MESES = 60
ANOS = 5
CENARIOS = ("pessimista", "provavel", "otimista")
JOGOS_NA_NOTA = {"Quartz BOVINO", "Dobble"}  # jogos comprados na loja, fora do consumo


# ---------------------------------------------------------------- dados-fonte
def ler_fontes() -> dict:
    f = {}
    # B3: médias das visitas (observador válido) e notas por dia
    r = pd.read_csv(B3 / "contagem_resumo_visitas.csv", sep=";", encoding="utf-8-sig")
    r = r[r["observador"] != "media"]
    f["visitas"] = r
    f["notas"] = pd.read_csv(B3 / "contagem_notas_intervalo.csv", sep=";", encoding="utf-8-sig")
    p = pd.read_csv(B3 / "_dados-brutos" / "app" / "contagem_precos.csv", sep=";", encoding="utf-8-sig")
    p = p[p["fonte_registro"] == "transcricao"].copy()
    p["total"] = p["qtd"] * p["valor"]
    eh_taxa = p["item"].str.startswith("Taxa de Entrada")
    eh_ab = ~eh_taxa & ~p["item"].isin(JOGOS_NA_NOTA) & ~p["item"].str.startswith("Acréscimo")
    notas = p.groupby("origem").apply(lambda g: pd.Series({
        "pessoas": g.loc[eh_taxa[g.index], "qtd"].sum(),
        "taxa": g.loc[eh_taxa[g.index], "total"].sum(),
        "ab": g.loc[eh_ab[g.index], "total"].sum(),
    }), include_groups=False)
    f["nfce"] = notas
    f["ab_por_pessoa"] = notas["ab"].sum() / notas["pessoas"].sum()
    f["valor_por_nota"] = (notas["taxa"] + notas["ab"]).mean()
    # B1: cesta de jogos (acervo inicial) e taxas de jogo dos concorrentes
    ai = pd.read_excel(B1 / "acervo_jogos" / "B1_cesta_jogos.xlsx", sheet_name="acervo_inicial")
    f["acervo"] = dict(zip(ai["tamanho"], ai["custo_provavel_rs"]))
    f["acervo_titulos"] = dict(zip(ai["tamanho"], ai["titulos"]))
    ps = pd.read_excel(B1 / "cardapios" / "B1_cardapios_itens_precos.xlsx", sheet_name="precos_subcategoria")
    taxa = ps["couvert ou taxa de jogo | couvert ou taxa de jogo"].dropna()
    taxa = taxa[taxa > 0]
    f["taxa_b1_mediana"], f["taxa_b1_n"] = taxa.median(), len(taxa)
    # B4
    fz = pd.read_csv(B4 / "financeiro" / "[B4] Aluguel Comercial por Capital 2026-07 (FipeZap).csv")
    fz = fz[fz["capital"] == "Salvador"].iloc[0]
    f["aluguel_m2"], f["aluguel_var12m"] = fz["comercial_locacao_preco_m2"], fz["comercial_locacao_var_12m_pct"]
    bcb = pd.read_csv(B4 / "financeiro" / "[B4] Selic IPCA e Salário Mínimo Último Valor 2026 (BCB).csv")
    f["selic"] = bcb.loc[bcb["serie"] == "selic_meta", "valor"].iloc[0]
    focus = pd.read_csv(B4 / "financeiro" / "[B4] Expectativas Focus 2026-2030 (BCB).csv")
    f["ipca_2027"] = focus[(focus["Indicador"] == "IPCA") & (focus["DataReferencia"] == 2027)]["Mediana"].iloc[0]
    f["juro_real"] = ((1 + f["selic"] / 100) / (1 + f["ipca_2027"] / 100) - 1) * 100
    pas_c = pd.read_csv(B4 / "financeiro" / "[B4] Serviços às Famílias Custos 2019-2023 (IBGE PAS).csv")
    pas_r = pd.read_csv(B4 / "financeiro" / "[B4] Serviços às Famílias Receita 2019-2023 (IBGE PAS).csv")
    sel = lambda d, var: d[(d["periodo"] == 2023) & (d["variavel"] == var)
                           & d["Serviços prestados às famílias"].str.startswith("1.2")]["valor"].iloc[0]
    f["pas_mercadorias_pct"] = sel(pas_c, "Mercadorias, materiais de consumo e de reposição") / sel(pas_r, "Receita operacional líquida") * 100
    cg = pd.read_csv(B4 / "multiuso" / "[B4] Salário Admissão por Capital 2025-2026 (CAGED).csv")
    f["caged"] = cg[cg["capital"] == "Salvador"]
    sn = pd.read_csv(B4 / "juridico_tributario" / "[B4] Faixas e Alíquotas por Anexo 2026 (Simples Nacional).csv")
    f["simples"] = {a: sn[sn["anexo"] == a].reset_index(drop=True) for a in ("I", "III")}
    pop = pd.read_csv(B4 / "mercado_demanda" / "[B4] Resumo de População e Renda por Capital 2022-2026 (IBGE).csv")
    f["pop_salvador"] = pop.loc[pop["capital"] == "Salvador", "pop_estimada_2026"].iloc[0]
    return f


def salario(f, ocupacao: str, recorte: str = "CNAE 56 - alimentação") -> float:
    c = f["caged"]
    return float(c[(c["ocupacao"] == ocupacao) & (c["recorte"] == recorte)]["salario_mediana"].iloc[0])


def quadro_pessoal(f) -> list[tuple]:
    """(cargo, CBO, quantidade, salário, fonte). Quadro de referência: dois turnos, 68 h por semana."""
    cn56 = "CAGED, admissões ago/2025-jul/2026, Salvador, CNAE 56, mediana"
    todas = "CAGED, admissões ago/2025-jul/2026, Salvador, todas as atividades, mediana (CNAE 56 com amostra pequena)"
    return [
        ("Gerente", "141510", 1, salario(f, "Gerente de restaurante"), cn56),
        ("Cozinheiro", "513205", 2, salario(f, "Cozinheiro geral"), cn56),
        ("Auxiliar de cozinha", "513505", 1, salario(f, "Auxiliar nos serviços de alimentação"), cn56),
        ("Garçom", "513405", 3, salario(f, "Garçom"), cn56),
        ("Atendente de bar e balcão", "513435", 1, salario(f, "Atendente de lanchonete"), cn56),
        ("Operador de caixa", "421125", 1, salario(f, "Operador de caixa"), cn56),
        ("Instrutor de jogos (recreador)", "371410", 2, salario(f, "Recreador", "todas as atividades"), todas),
        ("Auxiliar de limpeza", "514320", 1, salario(f, "Faxineiro"), cn56),
    ]


def premissas(f) -> list[dict]:
    """Registro de premissas (§3). Percentuais como fração. Ordem = ordem da aba."""
    vis = f["visitas"]
    pg, pm, oc = vis["tamanho_grupo_mesa"], vis["permanencia_mesa_media_min"], vis["ocupacao_mesas_media_pct"]
    a = f["acervo"]
    R = lambda *x: dict(zip(("id", "nome", "unidade", "pess", "prov", "otim", "fonte", "nivel", "just"), x))
    return [
        R("R01", "Mesas do salão", "mesas", 8, 10, 11,
          "B3: São Jogue Paralela (10) e Bela Vista (11), informação do autor", "N5",
          "Formato de referência igual a uma unidade da São Jogue; 8 se o ponto for menor."),
        R("R02", "Pessoas por grupo de mesa", "pessoas", 2.8, round(pg.mean(), 1), 4.4,
          "B3, 4 visitas (out/2026)", "N5",
          f"Média das visitas {pg.mean():.1f}, de {pg.min():.1f} a {pg.max():.1f}."),
        R("R03", "Permanência média", "min", 137, round(pm.mean()), 83,
          "B3, 4 visitas (out/2026)", "N5",
          f"Médias por visita de {pm.min():.0f} a {pm.max():.0f} min; permanência maior reduz o giro."),
        R("R04", "Taxa de ocupação, fim de semana (sex a dom)", "% das mesas", 0.25, 0.35, 0.45,
          "B3, noites de sexta e sábado (out/2026)", "N5 + N6",
          f"Noite de sexta e sábado: {oc.min():.1f}% a {oc.max():.1f}% (média {oc.mean():.1f}%). O dia inteiro fica abaixo do pico da noite."),
        R("R05", "Taxa de ocupação, dia útil (seg a qui)", "% das mesas", 0.10, 0.18, 0.25,
          "Estimativa do autor", "N6",
          "Cerca de metade do fim de semana. As visitas medem o pico, por orientação; sem dado de dia útil."),
        R("R06", "Taxa de entrada, fim de semana", "R$ por pessoa", 25, 28, 30,
          "São Jogue: NFC-e de out/2026 e informação do autor", "N3 + N5",
          "R$ 28,00 por 3 h de sexta a domingo na São Jogue; faixa de ± R$ 3,00."),
        R("R07", "Taxa de entrada, dia útil", "R$ por pessoa", 20, 22, 25,
          "São Jogue: informação do autor", "N3",
          "R$ 22,00 de segunda a quinta na São Jogue."),
        R("R08", "Gasto com alimentos e bebidas", "R$ por pessoa", 18, round(f["ab_por_pessoa"]), 32,
          "B3, 4 NFC-e das mesas dos observadores; cardápio da São Jogue (B1)", "N3 + N5",
          f"Média de R$ {f['ab_por_pessoa']:.2f} por pessoa nas notas (sem taxa, jogos e serviço); hambúrguer a R$ 45,00 e drinque a R$ 28,00 no cardápio."),
        R("R09", "Aluguéis de jogos", "aluguéis por mês", 10, 30, 60,
          "Estimativa do autor", "N6",
          "Serviço oferecido pela São Jogue; volume não observado."),
        R("R10", "Preço médio do aluguel de jogo", "R$", 15, 30, 40,
          "Tabela da São Jogue (4 dias: R$ 15,00 a R$ 60,00)", "N3",
          "Categoria prata como provável."),
        R("R11", "Ocupação no mês 1", "% da madura", 0.30, 0.40, 0.50,
          "Estimativa do autor (_instrucoes, §5)", "N6",
          "Sem ocupação madura no mês 1."),
        R("R12", "Meses até a maturidade", "meses", 12, 9, 6,
          "Estimativa do autor (_instrucoes, §5)", "N6",
          "Rampa de 6 a 12 meses."),
        R("R13", "Horas abertas por semana, fim de semana", "h", 28, 28, 28,
          "Horário da São Jogue (B1 e autor)", "N3",
          "Sexta e sábado das 12h às 22h; domingo das 13h às 21h."),
        R("R14", "Horas abertas por semana, dia útil", "h", 40, 40, 40,
          "Horário da São Jogue (B1 e autor)", "N3",
          "Segunda a quinta das 12h às 22h."),
        R("C01", "CMV de alimentos e bebidas", "% do A&B", 0.35, 0.30, 0.26,
          "IBGE PAS 2023, serviços de alimentação (B4)", "N2",
          f"Mercadorias e materiais = {f['pas_mercadorias_pct']:.1f}% da receita líquida do setor."),
        R("C02", "Taxa de cartão", "% da receita", 0.035, 0.030, 0.025,
          "Estimativa do autor", "N6",
          "Mistura de débito, crédito e Pix."),
        R("C03", "Área do ponto", "m²", 220, 180, 150,
          "Estimativa do autor", "N6",
          "Salão de 10 mesas, cozinha, caixa e uma sala reservável."),
        R("C04", "Aluguel comercial", "R$ por m² ao mês", round(f["aluguel_m2"] * (1 + f["aluguel_var12m"] / 100), 2), round(f["aluguel_m2"], 2), 45,
          "FipeZap, locação comercial, Salvador, jul/2026 (B4)", "N2",
          f"Pessimista com a alta de {f['aluguel_var12m']:.1f}% em 12 meses; otimista para ponto fora de shopping."),
        R("C05", "Condomínio e IPTU", "% do aluguel", 0.30, 0.20, 0.10,
          "Estimativa do autor", "N6", "Maior em shopping."),
        R("C06", "Fator do quadro de pessoal", "× quadro de 12", 14 / 12, 1.0, 10 / 12,
          "Quadro de referência (aba pessoal)", "N6",
          "14, 12 ou 10 pessoas em dois turnos; a São Jogue tinha de 3 a 7 instrutores nas visitas."),
        R("C07", "Encargos e provisões sobre a folha", "% do salário", 0.45, 0.37, 0.30,
          "CLT (FGTS, 13º, férias) e estimativa de benefícios", "N1 + N6",
          "No Simples, a contribuição patronal fica no DAS dos anexos I e III."),
        R("C08", "Pró-labore", "R$ por mês", 4000, 3000, 2000, "Estimativa do autor", "N6", "Um sócio na gestão."),
        R("C09", "Utilidades (energia, água, gás, internet)", "R$ por mês", 6000, 4500, 3500, "Estimativa do autor", "N6", "Cozinha própria e ar-condicionado."),
        R("C10", "Marketing", "R$ por mês", 2000, 1500, 1000, "Estimativa do autor", "N6", "Redes sociais e eventos."),
        R("C11", "Contabilidade", "R$ por mês", 1500, 1200, 900, "Estimativa do autor", "N6", "Escritório para empresa do Simples com folha."),
        R("C12", "Sistemas (caixa, reservas, nota fiscal)", "R$ por mês", 600, 450, 300, "Estimativa do autor", "N6", "Assinaturas."),
        R("C13", "Reposição do acervo", "% do acervo ao ano", 0.15, 0.10, 0.05, "Estimativa do autor", "N6", "Peças perdidas e jogos novos."),
        R("I01", "Reforma e instalações", "R$ por m²", 1400, 1000, 600, "Estimativa do autor", "N6", "Sem cotação; ponto alugado sem obra estrutural."),
        R("I02", "Mobiliário (mesas, cadeiras, estantes)", "R$", 60000, 45000, 30000, "Estimativa do autor; descrição das unidades da São Jogue", "N6", "10 mesas, 48 assentos e estantes de pinus com metal."),
        R("I03", "Cozinha e bar", "R$", 80000, 60000, 40000, "Estimativa do autor", "N6", "Cozinha própria, como a da São Jogue."),
        R("I04", "Caixa e informática", "R$", 8000, 6000, 4000, "Estimativa do autor", "N6", "1 computador, 1 impressora de nota e 2 maquininhas."),
        R("I05", "Acervo inicial de jogos", "R$", round(a["p75"], 2), round(a["mediana"], 2), round(a["p25"], 2),
          "B1, cesta de jogos dos concorrentes, preço provável nas lojas", "N3",
          f"Acervo do 3º quartil ({f['acervo_titulos']['p75']} títulos), da mediana ({f['acervo_titulos']['mediana']}) e do 1º quartil ({f['acervo_titulos']['p25']}); a São Jogue lista 370."),
        R("I06", "Gastos pré-operacionais", "R$", 30000, 20000, 12000, "Estimativa do autor", "N6", "Projeto, licenças, contratação e lançamento."),
        R("I07", "Caução do aluguel", "meses de aluguel", 3, 3, 2, "Estimativa do autor", "N6", "Prática de locação comercial."),
        R("I08", "Capital de giro", "meses de custo fixo", 3, 2, 1.5, "Estimativa do autor", "N6", "Cobre a rampa de maturação."),
        R("M01", "Taxa mínima de atratividade (real)", "% ao ano", 0.18, 0.15, 0.12,
          "BCB: Selic meta e Focus (B4); prêmio de risco do autor", "N2 + N6",
          f"Juro real de {f['juro_real']:.1f}% (Selic {f['selic']:.1f}% e IPCA esperado de {f['ipca_2027']:.1f}% para 2027) mais 3 a 9 p.p. de risco de pequeno negócio."),
        R("M02", "Semanas por mês", "semanas", 52 / 12, 52 / 12, 52 / 12, "Calendário", "N1", "52 semanas ÷ 12 meses."),
    ]


# --------------------------------------------------------------- modelo (Py)
def aliquota(f, anexo: str, rbt: float) -> float:
    t = f["simples"][anexo]
    if rbt <= 0:
        return 0.0
    linha = t[t["rbt12_min"] <= rbt].iloc[-1]
    return max(0.0, (rbt * linha["aliquota_nominal_pct"] / 100 - linha["parcela_deduzir"]) / rbt)


def calcula(f, v: dict, folha: float) -> dict:
    giro_fds = v["R13"] * 60 / v["R03"]
    giro_du = v["R14"] * 60 / v["R03"]
    m = []
    for mes in range(1, MESES + 1):
        rampa = min(1.0, v["R11"] + (1 - v["R11"]) * (mes - 1) / (v["R12"] - 1))
        p_fds = v["R01"] * v["R02"] * giro_fds * v["R04"] * rampa * v["M02"]
        p_du = v["R01"] * v["R02"] * giro_du * v["R05"] * rampa * v["M02"]
        taxa_fds, taxa_du = p_fds * v["R06"], p_du * v["R07"]
        ab = (p_fds + p_du) * v["R08"]
        alug = v["R09"] * v["R10"] * rampa
        m.append(dict(ano=(mes - 1) // 12 + 1, pessoas=p_fds + p_du, taxa=taxa_fds + taxa_du,
                      ab=ab, alug=alug, bruta=taxa_fds + taxa_du + ab + alug))
    m = pd.DataFrame(m)
    rbt = m.groupby("ano")["bruta"].sum()
    al_i = {a: aliquota(f, "I", rbt[a]) for a in rbt.index}
    al_iii = {a: aliquota(f, "III", rbt[a]) for a in rbt.index}
    m["simples"] = m["ab"] * m["ano"].map(al_i) + (m["taxa"] + m["alug"]) * m["ano"].map(al_iii)
    m["cmv"] = m["ab"] * v["C01"]
    m["cartao"] = m["bruta"] * v["C02"]
    aluguel = v["C03"] * v["C04"]
    fixos = (aluguel * (1 + v["C05"]) + folha * v["C06"] * (1 + v["C07"]) + v["C08"] + v["C09"]
             + v["C10"] + v["C11"] + v["C12"] + v["I05"] * v["C13"] / 12)
    m["fixos"] = fixos
    m["ebitda"] = m["bruta"] - m["simples"] - m["cmv"] - m["cartao"] - m["fixos"]
    inv = dict(reforma=v["I01"] * v["C03"], mobiliario=v["I02"], cozinha=v["I03"], ti=v["I04"],
               acervo=v["I05"], preop=v["I06"], caucao=v["I07"] * aluguel, giro=v["I08"] * fixos)
    total_inv = sum(inv.values())
    ano = m.groupby("ano")[["bruta", "simples", "cmv", "cartao", "fixos", "ebitda"]].sum()
    fc = [-total_inv] + list(ano["ebitda"])
    fc[-1] += inv["caucao"] + inv["giro"]
    desc = [x / (1 + v["M01"]) ** t for t, x in enumerate(fc)]
    vpl = sum(desc)
    acum, payback = 0.0, float("nan")
    for t, d in enumerate(desc):
        if t and acum < 0 <= acum + d:
            payback = t - 1 + (-acum) / d
            break
        acum += d
    mc_pct = 1 - (ano.loc[ANOS, ["simples", "cmv", "cartao"]].sum() / ano.loc[ANOS, "bruta"])
    return dict(vpl=vpl, tir=tir(fc), payback=payback, investimento=total_inv, inv=inv, fc=fc,
                receita_madura=m["bruta"].iloc[-1], pessoas_madura=m["pessoas"].iloc[-1],
                fixos=fixos, ebitda5=ano.loc[ANOS, "ebitda"], mc_pct=mc_pct,
                receita_equilibrio=fixos / mc_pct, alug_madura=m["alug"].iloc[-1], ano=ano)


def tir(fc: list) -> float:
    def vpl(r):
        return sum(x / (1 + r) ** t for t, x in enumerate(fc))
    lo, hi = -0.99, 10.0
    if vpl(lo) * vpl(hi) > 0:
        return float("nan")
    for _ in range(200):
        mid = (lo + hi) / 2
        if vpl(lo) * vpl(mid) <= 0:
            hi = mid
        else:
            lo = mid
    return (lo + hi) / 2


def valores(prem: list[dict], cenario: str) -> dict:
    return {p["id"]: p[cenario[:4] if cenario != "provavel" else "prov"] for p in prem}


def ruptura(f, prem, folha, pid: str):
    """Valor da premissa que zera o VPL (demais no provável), por bisseção."""
    base = valores(prem, "provavel")
    def vpl(x):
        v = dict(base, **{pid: x})
        return calcula(f, v, folha)["vpl"]
    p0 = base[pid]
    lo, hi = p0 * 0.05, p0 * 6
    if pid in ("R04", "R05"):
        lo, hi = 0.001, 1.0
    if vpl(lo) * vpl(hi) > 0:
        return float("nan")
    for _ in range(80):
        mid = (lo + hi) / 2
        if vpl(lo) * vpl(mid) <= 0:
            hi = mid
        else:
            lo = mid
    return (lo + hi) / 2


# --------------------------------------------------------------- planilha
FONTE = Font(name="Arial", size=10, color="000000")
NEGRITO = Font(name="Arial", size=10, color="000000", bold=True)
BRANCO = PatternFill("solid", fgColor="FFFFFF")
FINA = Side(style="thin", color="000000")
MOEDA, PCT, DEC1, DEC2, INT = "#,##0.00", "0.00%", "#,##0.0", "#,##0.00", "#,##0"
COL_MES = [get_column_letter(3 + i) for i in range(MESES)]  # C..BJ


def titulo(ws, texto, ncol):
    ws.cell(1, 1, texto)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ncol)


def estilo(ws, linha_cab, ultima, ncol, fmt=None):
    """Arial 10, fundo branco, cabeçalho centralizado em negrito com borda em cima e embaixo,
    borda inferior na última linha; texto à esquerda e número à direita."""
    for row in ws.iter_rows(min_row=1, max_row=max(ws.max_row, ultima), max_col=ncol):
        for c in row:
            c.font, c.fill = FONTE, BRANCO
            if c.row == 1:
                c.font, c.alignment = NEGRITO, Alignment(horizontal="center")
            elif c.row == linha_cab:
                c.font = NEGRITO
                c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
                c.border = Border(top=FINA, bottom=FINA)
            elif linha_cab < c.row <= ultima:
                txt = isinstance(c.value, str) and not str(c.value).startswith("=")
                c.alignment = Alignment(horizontal="left" if txt else "right", vertical="top", wrap_text=txt)
                if c.row == ultima:
                    c.border = Border(bottom=FINA)
                if fmt and not txt and c.column > 1:
                    c.number_format = fmt


def nota(ws, linha, texto, ncol):
    ws.cell(linha, 1, texto).font = FONTE
    ws.cell(linha, 1).alignment = Alignment(wrap_text=True, vertical="top")
    ws.merge_cells(start_row=linha, start_column=1, end_row=linha, end_column=ncol)
    ws.row_dimensions[linha].height = 15 * max(1, len(texto) // 110 + 1)


def planilha(f, prem, quadro, resultados, tornado, rupt):
    wb = openpyxl.Workbook()
    # ---- premissas
    ws = wb.active
    ws.title = "premissas"
    titulo(ws, "Registro de premissas do modelo financeiro", 11)
    ws["A2"], ws["D2"] = "Cenário ativo (1 = pessimista, 2 = provável, 3 = otimista)", 2
    cab = ["ID", "Premissa", "Unidade", "Pessimista", "Provável", "Otimista", "Valor usado",
           "Fonte", "Nível", "Justificativa da faixa", "Impacto no VPL"]
    for j, h in enumerate(cab, 1):
        ws.cell(4, j, h)
    impacto = {t["id"]: ("Alto" if i < 5 else "Médio" if i < 10 else "Baixo") for i, t in enumerate(tornado)}
    linha_de = {}
    for i, p in enumerate(prem):
        r = 5 + i
        linha_de[p["id"]] = r
        for j, k in enumerate(("id", "nome", "unidade", "pess", "prov", "otim"), 1):
            ws.cell(r, j, p[k])
        ws.cell(r, 7, f"=CHOOSE($D$2,D{r},E{r},F{r})")
        ws.cell(r, 8, p["fonte"]); ws.cell(r, 9, p["nivel"]); ws.cell(r, 10, p["just"])
        ws.cell(r, 11, impacto.get(p["id"], "Baixo"))
        wb.defined_names[f"P_{p['id']}"] = DefinedName(f"P_{p['id']}", attr_text=f"premissas!$G${r}")
    ult = 4 + len(prem)
    estilo(ws, 4, ult, 11)
    for p in prem:
        r = linha_de[p["id"]]
        fmt = PCT if p["unidade"].startswith("%") else (MOEDA if p["unidade"].startswith("R$") else DEC2)
        for c in range(4, 8):
            ws.cell(r, c).number_format = fmt
    ws["A2"].font = NEGRITO
    ws["D2"].font = NEGRITO
    nota(ws, ult + 2, "Fonte: elaboração própria com dados dos Blocos B1, B3 e B4 (out/2026). Níveis de evidência N1 (mais confiável) a N6 (estimativa do autor), conforme _instrucoes/8_financeiro.md. Impacto no VPL pelo tornado (aba sensibilidade): 5 maiores = Alto, 5 seguintes = Médio.", 11)
    for col, w in zip("ABCDEFGHIJK", (6, 34, 16, 12, 12, 12, 12, 36, 9, 52, 10)):
        ws.column_dimensions[col].width = w

    # ---- pessoal
    ws = wb.create_sheet("pessoal")
    titulo(ws, "Quadro de pessoal de referência e salários de admissão em Salvador", 6)
    for j, h in enumerate(["Cargo", "CBO", "Quantidade", "Salário mensal (R$)", "Total mensal (R$)", "Fonte"], 1):
        ws.cell(2, j, h)
    for i, (cargo, cbo, q, s, fonte) in enumerate(quadro):
        r = 3 + i
        ws.cell(r, 1, cargo); ws.cell(r, 2, cbo); ws.cell(r, 3, q); ws.cell(r, 4, s)
        ws.cell(r, 5, f"=C{r}*D{r}"); ws.cell(r, 6, fonte)
    rt = 3 + len(quadro)
    ws.cell(rt, 1, "Total"); ws.cell(rt, 3, f"=SUM(C3:C{rt - 1})"); ws.cell(rt, 5, f"=SUM(E3:E{rt - 1})")
    estilo(ws, 2, rt, 6)
    for r in range(3, rt + 1):
        ws.cell(r, 3).number_format = INT
        ws.cell(r, 4).number_format = ws.cell(r, 5).number_format = MOEDA
    wb.defined_names["FOLHA_REF"] = DefinedName("FOLHA_REF", attr_text=f"pessoal!$E${rt}")
    nota(ws, rt + 2, "Fonte: Novo CAGED (MTE), salário de admissão, mediana, ago/2025 a jul/2026 (B4). Encargos e fator do quadro na aba premissas (C06 e C07).", 6)
    for col, w in zip("ABCDEF", (30, 9, 11, 16, 16, 60)):
        ws.column_dimensions[col].width = w

    # ---- simples
    ws = wb.create_sheet("simples")
    titulo(ws, "Simples Nacional: faixas dos anexos I e III e alíquota efetiva por ano", 6)
    blocos = {"I": 4, "III": 13}
    for anexo, r0 in blocos.items():
        ws.cell(r0 - 1, 1, f"Anexo {anexo}")
        for j, h in enumerate(["Faixa", "RBT12 mínima (R$)", "RBT12 máxima (R$)", "Alíquota nominal (%)", "Parcela a deduzir (R$)"], 1):
            ws.cell(r0, j, h)
        for i, row in f["simples"][anexo].iterrows():
            r = r0 + 1 + i
            for j, k in enumerate(("faixa", "rbt12_min", "rbt12_max", "aliquota_nominal_pct", "parcela_deduzir"), 1):
                ws.cell(r, j, float(row[k]) if j > 1 else int(row[k]))
    ws.cell(22, 1, "Ano")
    for a in range(1, ANOS + 1):
        ws.cell(22, 1 + a, a)
    ws.cell(23, 1, "Receita bruta do ano (R$)")
    ws.cell(24, 1, "Alíquota efetiva, anexo I (A&B)")
    ws.cell(25, 1, "Alíquota efetiva, anexo III (taxa de entrada e aluguel de jogos)")
    for a in range(1, ANOS + 1):
        c = get_column_letter(1 + a)
        ws[f"{c}23"] = f"=SUMIF(receitas!$C$3:$BJ$3,{c}22,receitas!$C$12:$BJ$12)"
        for r, r0 in ((24, 4), (25, 13)):
            mins, noms, pds = f"$B${r0 + 1}:$B${r0 + 6}", f"$D${r0 + 1}:$D${r0 + 6}", f"$E${r0 + 1}:$E${r0 + 6}"
            ws[f"{c}{r}"] = f"=IF({c}23<=0,0,MAX(0,({c}23*LOOKUP({c}23,{mins},{noms})/100-LOOKUP({c}23,{mins},{pds}))/{c}23))"
    for row in ws.iter_rows(min_row=1, max_row=25, max_col=6):
        for c in row:
            c.font, c.fill = FONTE, BRANCO
            if isinstance(c.value, (int, float)) and c.column > 1:
                c.alignment = Alignment(horizontal="right")
    ws["A1"].font = NEGRITO
    for r in (3, 4, 12, 13, 22):
        for c in range(1, 7):
            ws.cell(r, c).font = NEGRITO
    for r in (4, 13, 22):
        for c in range(1, 7 if r != 22 else 2 + ANOS):
            ws.cell(r, c).border = Border(top=FINA, bottom=FINA)
            ws.cell(r, c).alignment = Alignment(horizontal="center", wrap_text=True)
    for r in list(range(5, 11)) + list(range(14, 20)):
        for c in (2, 3, 5):
            ws.cell(r, c).number_format = MOEDA
        ws.cell(r, 4).number_format = DEC2
    for a in range(2, 2 + ANOS):
        ws.cell(23, a).number_format = MOEDA
        ws.cell(24, a).number_format = ws.cell(25, a).number_format = PCT
        ws.cell(25, a).border = Border(bottom=FINA)
    ws.cell(25, 1).border = Border(bottom=FINA)
    nota(ws, 27, "Fonte: LC 123/2006, anexos I e III, valores de 2026 (B4). Enquadramento a confirmar no cap. 8: A&B no anexo I (CNAE 5611-2/01); taxa de entrada e aluguel de jogos no anexo III. RBT12 aproximada pela receita do próprio ano.", 6)
    ws.column_dimensions["A"].width = 44
    for col in "BCDEF":
        ws.column_dimensions[col].width = 18

    # ---- receitas (mensal)
    ws = wb.create_sheet("receitas")
    titulo(ws, "Receitas mensais projetadas (preços de out/2026)", 2 + MESES)
    ws["A2"], ws["A3"] = "Mês", "Ano"
    linhas = {4: "Fator de rampa (% da ocupação madura)", 5: "Pessoas, fim de semana", 6: "Pessoas, dia útil",
              7: "Pessoas atendidas", 8: "Taxa de entrada, fim de semana (R$)", 9: "Taxa de entrada, dia útil (R$)",
              10: "Alimentos e bebidas (R$)", 11: "Aluguel de jogos (R$)", 12: "Receita bruta (R$)"}
    for r, t in linhas.items():
        ws.cell(r, 1, t)
    for i, c in enumerate(COL_MES):
        ws[f"{c}2"] = i + 1
        ws[f"{c}3"] = f"=ROUNDUP({c}2/12,0)"
        ws[f"{c}4"] = f"=MIN(1,P_R11+(1-P_R11)*({c}2-1)/(P_R12-1))"
        ws[f"{c}5"] = f"=P_R01*P_R02*(P_R13*60/P_R03)*P_R04*{c}4*P_M02"
        ws[f"{c}6"] = f"=P_R01*P_R02*(P_R14*60/P_R03)*P_R05*{c}4*P_M02"
        ws[f"{c}7"] = f"={c}5+{c}6"
        ws[f"{c}8"] = f"={c}5*P_R06"
        ws[f"{c}9"] = f"={c}6*P_R07"
        ws[f"{c}10"] = f"={c}7*P_R08"
        ws[f"{c}11"] = f"=P_R09*P_R10*{c}4"
        ws[f"{c}12"] = f"=SUM({c}8:{c}11)"
    estilo(ws, 2, 12, 2 + MESES)
    for c in COL_MES:
        ws[f"{c}4"].number_format = PCT
        for r in (5, 6, 7):
            ws[f"{c}{r}"].number_format = DEC1
        for r in range(8, 13):
            ws[f"{c}{r}"].number_format = MOEDA
    ws.column_dimensions["A"].width = 40
    ws.freeze_panes = "C4"

    # ---- custos (mensal)
    ws = wb.create_sheet("custos")
    titulo(ws, "Custos e despesas mensais projetados (preços de out/2026)", 2 + MESES)
    ws["A2"], ws["A3"] = "Mês", "Ano"
    linhas = {4: "Simples Nacional (R$)", 5: "CMV de alimentos e bebidas (R$)", 6: "Taxa de cartão (R$)",
              7: "Custos variáveis (R$)", 8: "Aluguel (R$)", 9: "Condomínio e IPTU (R$)",
              10: "Pessoal com encargos (R$)", 11: "Pró-labore (R$)", 12: "Utilidades (R$)", 13: "Marketing (R$)",
              14: "Contabilidade (R$)", 15: "Sistemas (R$)", 16: "Reposição do acervo (R$)",
              17: "Custos fixos (R$)", 18: "Custos totais (R$)"}
    for r, t in linhas.items():
        ws.cell(r, 1, t)
    for i, c in enumerate(COL_MES):
        ws[f"{c}2"] = i + 1
        ws[f"{c}3"] = f"=receitas!{c}3"
        ws[f"{c}4"] = (f"=receitas!{c}10*INDEX(simples!$B$24:$F$24,{c}3)"
                       f"+(receitas!{c}8+receitas!{c}9+receitas!{c}11)*INDEX(simples!$B$25:$F$25,{c}3)")
        ws[f"{c}5"] = f"=receitas!{c}10*P_C01"
        ws[f"{c}6"] = f"=receitas!{c}12*P_C02"
        ws[f"{c}7"] = f"=SUM({c}4:{c}6)"
        ws[f"{c}8"] = "=P_C03*P_C04"
        ws[f"{c}9"] = f"={c}8*P_C05"
        ws[f"{c}10"] = "=FOLHA_REF*P_C06*(1+P_C07)"
        ws[f"{c}11"] = "=P_C08"
        ws[f"{c}12"] = "=P_C09"
        ws[f"{c}13"] = "=P_C10"
        ws[f"{c}14"] = "=P_C11"
        ws[f"{c}15"] = "=P_C12"
        ws[f"{c}16"] = "=P_I05*P_C13/12"
        ws[f"{c}17"] = f"=SUM({c}8:{c}16)"
        ws[f"{c}18"] = f"={c}7+{c}17"
    estilo(ws, 2, 18, 2 + MESES, MOEDA)
    for c in COL_MES:
        ws[f"{c}2"].number_format = ws[f"{c}3"].number_format = "0"
    ws.column_dimensions["A"].width = 36
    ws.freeze_panes = "C4"

    # ---- investimento
    ws = wb.create_sheet("investimento")
    titulo(ws, "Investimento inicial e depreciação", 4)
    for j, h in enumerate(["Item", "Valor (R$)", "Vida útil (anos)", "Depreciação anual (R$)"], 1):
        ws.cell(2, j, h)
    itens = [("Reforma e instalações", "=P_I01*P_C03", 5), ("Mobiliário", "=P_I02", 10),
             ("Cozinha e bar", "=P_I03", 10), ("Caixa e informática", "=P_I04", 5),
             ("Acervo inicial de jogos", "=P_I05", 5), ("Gastos pré-operacionais", "=P_I06", 5),
             ("Caução do aluguel", "=P_I07*custos!$C$8", 0), ("Capital de giro", "=P_I08*custos!$C$17", 0)]
    for i, (n, fo, vida) in enumerate(itens):
        r = 3 + i
        ws.cell(r, 1, n); ws.cell(r, 2, fo); ws.cell(r, 3, vida)
        ws.cell(r, 4, f"=IF(C{r}>0,B{r}/C{r},0)")
    rt = 3 + len(itens)
    ws.cell(rt, 1, "Total"); ws.cell(rt, 2, f"=SUM(B3:B{rt - 1})"); ws.cell(rt, 4, f"=SUM(D3:D{rt - 1})")
    estilo(ws, 2, rt, 4, MOEDA)
    for r in range(3, rt):
        ws.cell(r, 3).number_format = "0"
    nota(ws, rt + 2, "Fonte: aba premissas (I01 a I08, C03). Reforma depreciada em 5 anos (prazo típico de locação); caução e capital de giro voltam no ano 5. Sem valor residual dos demais itens.", 4)
    ws.column_dimensions["A"].width = 30
    for col in "BCD":
        ws.column_dimensions[col].width = 18

    # ---- dre (anual)
    ws = wb.create_sheet("dre")
    titulo(ws, "Demonstração do resultado projetada (R$, preços de out/2026)", 1 + ANOS)
    ws["A2"] = "Conta"
    for a in range(1, ANOS + 1):
        ws.cell(2, 1 + a, f"Ano {a}")
        ws.cell(3, 1 + a, a)
    ws["A3"] = "Ano (número)"
    soma = lambda aba, lin, c: f"SUMIF({aba}!$C$3:$BJ$3,{c}$3,{aba}!$C${lin}:$BJ${lin})"
    contas = [(4, "Receita bruta", lambda c: "=" + soma("receitas", 12, c)),
              (5, "(-) Simples Nacional", lambda c: "=-" + soma("custos", 4, c)),
              (6, "Receita líquida", lambda c: f"={c}4+{c}5"),
              (7, "(-) CMV de alimentos e bebidas", lambda c: "=-" + soma("custos", 5, c)),
              (8, "(-) Taxa de cartão", lambda c: "=-" + soma("custos", 6, c)),
              (9, "Margem de contribuição", lambda c: f"={c}6+{c}7+{c}8"),
              (10, "(-) Custos fixos", lambda c: "=-" + soma("custos", 17, c)),
              (11, "EBITDA", lambda c: f"={c}9+{c}10"),
              (12, "(-) Depreciação e amortização", lambda c: "=-investimento!$D$11"),
              (13, "Lucro líquido", lambda c: f"={c}11+{c}12"),
              (14, "Margem líquida", lambda c: f"=IF({c}4=0,0,{c}13/{c}4)")]
    for r, n, fo in contas:
        ws.cell(r, 1, n)
        for a in range(1, ANOS + 1):
            c = get_column_letter(1 + a)
            ws[f"{c}{r}"] = fo(c)
    estilo(ws, 2, 14, 1 + ANOS, MOEDA)
    for a in range(2, 2 + ANOS):
        ws.cell(14, a).number_format = PCT
        ws.cell(3, a).number_format = "0"
    nota(ws, 16, "Fonte: abas receitas, custos e investimento. No Simples Nacional, IRPJ e CSLL estão no DAS; não há imposto sobre o lucro à parte.", 1 + ANOS)
    ws.column_dimensions["A"].width = 34
    for a in range(2, 2 + ANOS):
        ws.column_dimensions[get_column_letter(a)].width = 16

    # ---- fluxo de caixa (anual)
    ws = wb.create_sheet("fluxo_caixa")
    titulo(ws, "Fluxo de caixa projetado (R$, preços de out/2026)", 2 + ANOS)
    ws["A2"] = "Linha"
    for t in range(ANOS + 1):
        ws.cell(2, 2 + t, f"Ano {t}")
        ws.cell(3, 2 + t, t)
    ws["A3"] = "Ano (número)"
    rot = {4: "Investimento", 5: "EBITDA", 6: "Retorno da caução e do capital de giro", 7: "Fluxo líquido",
           8: "Fator de desconto", 9: "Fluxo descontado", 10: "Fluxo descontado acumulado",
           11: "Payback descontado (auxiliar)"}
    for r, n in rot.items():
        ws.cell(r, 1, n)
    for t in range(ANOS + 1):
        c = get_column_letter(2 + t)
        ws[f"{c}4"] = "=-investimento!$B$11" if t == 0 else 0
        ws[f"{c}5"] = 0 if t == 0 else f"=dre!{get_column_letter(1 + t)}11"
        ws[f"{c}6"] = "=investimento!$B$9+investimento!$B$10" if t == ANOS else 0
        ws[f"{c}7"] = f"=SUM({c}4:{c}6)"
        ws[f"{c}8"] = f"=1/(1+P_M01)^{c}3"
        ws[f"{c}9"] = f"={c}7*{c}8"
        ws[f"{c}10"] = f"={c}9" if t == 0 else f"={get_column_letter(1 + t)}10+{c}9"
        if t:
            p = get_column_letter(1 + t)
            ws[f"{c}11"] = f'=IF(AND({p}10<0,{c}10>=0),{p}3+(-{p}10)/{c}9,"")'
    estilo(ws, 2, 11, 2 + ANOS, MOEDA)
    for t in range(ANOS + 1):
        c = get_column_letter(2 + t)
        ws[f"{c}3"].number_format = "0"
        ws[f"{c}8"].number_format = "0.0000"
        ws[f"{c}11"].number_format = DEC2
    ws.column_dimensions["A"].width = 38
    for t in range(ANOS + 1):
        ws.column_dimensions[get_column_letter(2 + t)].width = 16

    # ---- indicadores
    ws = wb.create_sheet("indicadores")
    titulo(ws, "Indicadores de viabilidade do cenário ativo", 3)
    for j, h in enumerate(["Indicador", "Valor", "Unidade"], 1):
        ws.cell(2, j, h)
    ind = [("Cenário ativo", '=CHOOSE(P_CEN,"pessimista","provável","otimista")', ""),
           ("Investimento inicial", "=investimento!B11", "R$"),
           ("VPL a TMA real", "=SUM(fluxo_caixa!B9:G9)", "R$"),
           ("TIR real", '=IFERROR(IRR(fluxo_caixa!B7:G7),"não há")', "% ao ano"),
           ("Payback descontado", '=IF(COUNT(fluxo_caixa!C11:G11)=0,"mais de 5 anos",MIN(fluxo_caixa!C11:G11))', "anos"),
           ("Receita bruta mensal madura (mês 60)", "=receitas!BJ12", "R$"),
           ("Custos fixos mensais", "=custos!BJ17", "R$"),
           ("Margem de contribuição (ano 5)", "=dre!F9/dre!F4", "% da receita bruta"),
           ("Receita bruta mensal de equilíbrio", "=B9/B10", "R$"),
           ("Receita madura ÷ receita de equilíbrio", "=B8/B11", "vezes"),
           ("Ocupação de equilíbrio, fim de semana (dia útil na mesma proporção)", "=P_R04*(B11-receitas!BJ11)/(B8-receitas!BJ11)", "% das mesas"),
           ("Ocupação de equilíbrio, dia útil (mesma proporção)", "=P_R05*(B11-receitas!BJ11)/(B8-receitas!BJ11)", "% das mesas")]
    for i, (n, fo, u) in enumerate(ind):
        ws.cell(3 + i, 1, n); ws.cell(3 + i, 2, fo); ws.cell(3 + i, 3, u)
    rt = 2 + len(ind)
    wb.defined_names["P_CEN"] = DefinedName("P_CEN", attr_text="premissas!$D$2")
    estilo(ws, 2, rt, 3, MOEDA)
    ws["B6"].number_format = PCT
    ws["B7"].number_format = DEC2
    ws["B11"].number_format = MOEDA
    ws["B10"].number_format = ws["B13"].number_format = ws["B14"].number_format = PCT
    ws["B12"].number_format = DEC2
    ws["B3"].alignment = Alignment(horizontal="right")
    nota(ws, rt + 2, "Fonte: abas fluxo_caixa, dre, receitas e custos. Ocupação de equilíbrio por aproximação linear: a receita sem o aluguel de jogos cresce na proporção da ocupação.", 3)
    ws.column_dimensions["A"].width = 58
    ws.column_dimensions["B"].width = 18
    ws.column_dimensions["C"].width = 20

    # ---- sensibilidade (valores do Python)
    ws = wb.create_sheet("sensibilidade")
    titulo(ws, "Cenários, tornado e pontos de ruptura (calculados pelo script com as mesmas fórmulas)", 7)
    for j, h in enumerate(["Cenário", "Investimento (R$)", "Receita mensal madura (R$)", "EBITDA do ano 5 (R$)", "VPL (R$)", "TIR real", "Payback descontado (anos)"], 1):
        ws.cell(2, j, h)
    for i, cen in enumerate(CENARIOS):
        res = resultados[cen]
        r = 3 + i
        vals = [("pessimista", "provável", "otimista")[i], res["investimento"], res["receita_madura"], res["ebitda5"],
                res["vpl"], res["tir"] if res["tir"] == res["tir"] else "não há",
                res["payback"] if res["payback"] == res["payback"] else "mais de 5"]
        for j, x in enumerate(vals, 1):
            ws.cell(r, j, x)
    estilo(ws, 2, 5, 7, MOEDA)
    for r in range(3, 6):
        ws.cell(r, 6).number_format = PCT
        ws.cell(r, 7).number_format = DEC2
    r0 = 8
    ws.cell(r0 - 1, 1, "Tornado: VPL com a premissa no extremo e as demais no provável")
    for j, h in enumerate(["ID", "Premissa", "VPL no pessimista (R$)", "VPL no otimista (R$)", "Amplitude (R$)", "Ordem"], 1):
        ws.cell(r0, j, h)
    for i, t in enumerate(tornado):
        r = r0 + 1 + i
        for j, x in enumerate((t["id"], t["nome"], t["vpl_pess"], t["vpl_otim"], t["amplitude"], i + 1), 1):
            ws.cell(r, j, x)
    ultt = r0 + len(tornado)
    r1 = ultt + 3
    ws.cell(r1 - 1, 1, "Ponto de ruptura das premissas críticas (VPL = 0, demais no provável)")
    for j, h in enumerate(["ID", "Premissa", "Pessimista", "Provável", "Otimista", "Valor de ruptura", "Unidade"], 1):
        ws.cell(r1, j, h)
    for i, rp in enumerate(rupt):
        r = r1 + 1 + i
        for j, x in enumerate((rp["id"], rp["nome"], rp["pess"], rp["prov"], rp["otim"],
                               rp["ruptura"] if rp["ruptura"] == rp["ruptura"] else "fora do intervalo testado", rp["unidade"]), 1):
            ws.cell(r, j, x)
    ultr = r1 + len(rupt)
    # estilo dos três blocos
    for row in ws.iter_rows(min_row=1, max_row=ultr + 3, max_col=7):
        for c in row:
            c.font, c.fill = FONTE, BRANCO
    for rc, ru in ((2, 5), (r0, ultt), (r1, ultr)):
        for c in range(1, 8):
            ws.cell(rc, c).font = NEGRITO
            ws.cell(rc, c).alignment = Alignment(horizontal="center", wrap_text=True, vertical="center")
            ws.cell(rc, c).border = Border(top=FINA, bottom=FINA)
            ws.cell(ru, c).border = Border(bottom=FINA)
        for r in range(rc + 1, ru + 1):
            for c in range(1, 8):
                v = ws.cell(r, c).value
                ws.cell(r, c).alignment = Alignment(horizontal="left" if isinstance(v, str) else "right", wrap_text=True)
    ws["A1"].font = NEGRITO
    ws["A1"].alignment = Alignment(horizontal="center")
    for r in (r0 - 1, r1 - 1):
        ws.cell(r, 1).font = NEGRITO
    for r in range(r0 + 1, ultt + 1):
        for c in (3, 4, 5):
            ws.cell(r, c).number_format = MOEDA
    for i, rp in enumerate(rupt):
        r = r1 + 1 + i
        fmt = PCT if rp["unidade"].startswith("%") else (MOEDA if rp["unidade"].startswith("R$") else DEC2)
        for c in (3, 4, 5, 6):
            ws.cell(r, c).number_format = fmt
    nota(ws, ultr + 2, "Fonte: elaboração própria. Valores gravados pelo script _src/extras/gerar_modelo_financeiro.py, que recalcula o modelo em Python e confere o VPL com esta planilha nos três cenários; mudar a aba premissas não atualiza esta aba (rodar o script de novo).", 7)
    for col, w in zip("ABCDEFG", (12, 44, 18, 18, 18, 18, 18)):
        ws.column_dimensions[col].width = w

    # ---- checagens de coerência (§5) e triangulação (§4.1)
    ws = wb.create_sheet("checagens")
    titulo(ws, "Triangulação da receita e checagens de coerência (cenário ativo, mês 60)", 4)
    for j, h in enumerate(["Checagem", "Modelo", "Referência", "Fonte da referência"], 1):
        ws.cell(2, j, h)
    notas = f["notas"].set_index("unidade")
    ch = [
        ("Pessoas atendidas por dia", "=receitas!BJ7/(P_M02*7)", "=P_R01*P_R02*(P_R13+P_R14)*60/P_R03/7",
         "Capacidade: todas as mesas ocupadas todas as horas, com a permanência média"),
        ("Uso da capacidade", "=B3/C3", "", "Tem de ficar abaixo de 100%"),
        ("Grupos atendidos por dia", "=B3/P_R02", float(notas.loc["BV", "notas_por_dia"]),
         "NFC-e emitidas por dia, São Jogue Bela Vista, 10/10 a 16/10/2026 (B3); piso, pois a nota sai só a pedido"),
        ("Grupos atendidos por dia (Paralela)", "=B3/P_R02", float(notas.loc["PA", "notas_por_dia"]),
         "NFC-e por dia, São Jogue Paralela, 17/10 a 30/10/2026 (B3), com o fim de semana da eleição"),
        ("Receita mensal madura × piso de vendas da Bela Vista (R$)", "=receitas!BJ12",
         float(notas.loc["BV", "notas_por_dia"]) * f["valor_por_nota"] * 365 / 12,
         f"NFC-e por dia × R$ {f['valor_por_nota']:.2f} por nota (média das 4 notas das mesas dos observadores, sem jogos e serviço) × 30,4 dias"),
        ("Receita mensal madura × piso de vendas da Paralela (R$)", "=receitas!BJ12",
         float(notas.loc["PA", "notas_por_dia"]) * f["valor_por_nota"] * 365 / 12, "Idem, Paralela"),
        ("Receita por pessoa (R$)", "=receitas!BJ12/receitas!BJ7", f["valor_por_nota"] / 2,
         "Notas das mesas dos observadores: taxa e consumo por pessoa"),
        ("Taxa de entrada de fim de semana (R$)", "=P_R06", float(f["taxa_b1_mediana"]),
         f"Mediana do couvert ou taxa de jogo nos cardápios de {f['taxa_b1_n']} concorrentes (B1); acima dela, o posicionamento tem de justificar"),
        ("Pessoas atendidas por dia por funcionário", "=B3/(12*P_C06)", "",
         "Quadro de referência de 12 em dois turnos (aba pessoal)"),
        ("Visitas por ano ÷ população de Salvador", "=receitas!BJ7*12/C12", float(f["pop_salvador"]),
         "População estimada de Salvador em 2026 (IBGE, B4); visitas, não pessoas distintas"),
        ("Meses até a maturidade", "=P_R12", "", "Entre 6 e 12 meses (_instrucoes, §5)"),
    ]
    for i, (n, mo, ref, fonte) in enumerate(ch):
        r = 3 + i
        ws.cell(r, 1, n); ws.cell(r, 2, mo); ws.cell(r, 3, ref if ref != "" else None); ws.cell(r, 4, fonte)
    rt = 2 + len(ch)
    estilo(ws, 2, rt, 4, DEC2)
    ws["B4"].number_format = ws["B12"].number_format = "0.000%"
    for r in (7, 8, 9, 10):
        ws[f"B{r}"].number_format = ws[f"C{r}"].number_format = MOEDA
    ws["C12"].number_format = INT
    nota(ws, rt + 2, "Fonte: abas do modelo e Blocos B1, B3 e B4. A comparação com as notas é de ordem de grandeza: a NFC-e só é emitida quando o cliente pede.", 4)
    for col, w in zip("ABCD", (48, 18, 18, 70)):
        ws.column_dimensions[col].width = w
    return wb


# --------------------------------------------------------------- conferência
def recalcula(xlsx: Path, cenario: int) -> float:
    """Abre uma cópia no LibreOffice com o cenário dado, recalcula e devolve o VPL."""
    with tempfile.TemporaryDirectory() as tmp:
        tmp = Path(tmp)
        wb = openpyxl.load_workbook(xlsx)
        wb["premissas"]["D2"] = cenario
        src = tmp / "modelo.xlsx"
        wb.save(src)
        out = tmp / "out"
        subprocess.run([SOFFICE, "--headless", "--calc", "--convert-to", "xlsx", "--outdir", str(out), str(src)],
                       check=True, capture_output=True, timeout=180)
        v = openpyxl.load_workbook(out / "modelo.xlsx", data_only=True)
        erros = [f"{ws.title}!{c.coordinate}" for ws in v for row in ws.iter_rows() for c in row
                 if isinstance(c.value, str) and c.value.startswith(("#", "Err:"))]
        if erros:
            raise SystemExit(f"Erros de fórmula no LibreOffice: {erros[:10]}")
        return v["indicadores"]["B5"].value


def figura_tornado(tornado, base_vpl):
    top = tornado[:10][::-1]
    fig, ax = plt.subplots(figsize=(6.3, 4.2), dpi=300)
    y = range(len(top))
    mil = lambda x: x / 1000
    for i, t in enumerate(top):
        ax.barh(i, mil(t["vpl_pess"] - base_vpl), left=mil(base_vpl), height=0.6, color="#4d4d4d", label="Premissa no pessimista" if i == 0 else None)
        ax.barh(i, mil(t["vpl_otim"] - base_vpl), left=mil(base_vpl), height=0.6, color="#bdbdbd", label="Premissa no otimista" if i == 0 else None)
    ax.axvline(mil(base_vpl), color="black", linewidth=0.8)
    ax.axvline(0, color="black", linewidth=0.8, linestyle=":")
    ax.set_yticks(list(y))
    ax.set_yticklabels([f"{t['id']} {t['nome']}" for t in top], fontsize=7, fontname="Arial")
    ax.set_xlabel("VPL (R$ mil)", fontsize=8, fontname="Arial")
    ax.tick_params(axis="x", labelsize=7)
    for s in ("top", "right"):
        ax.spines[s].set_visible(False)
    ax.legend(prop={"family": "Arial", "size": 7}, frameon=False, loc="lower right")
    for t in ax.get_xticklabels():
        t.set_fontname("Arial")
    ax.xaxis.set_major_formatter(matplotlib.ticker.FuncFormatter(lambda x, _: f"{x:,.0f}".replace(",", ".")))
    fig.tight_layout()
    FIGURA.parent.mkdir(parents=True, exist_ok=True)
    fig.savefig(FIGURA, facecolor="white")
    plt.close(fig)


def main():
    f = ler_fontes()
    prem = premissas(f)
    quadro = quadro_pessoal(f)
    folha = sum(q * s for _, _, q, s, _ in quadro)
    resultados = {c: calcula(f, valores(prem, c), folha) for c in CENARIOS}
    base = resultados["provavel"]
    tornado = []
    for p in prem:
        if p["pess"] == p["otim"]:
            continue
        vp = calcula(f, dict(valores(prem, "provavel"), **{p["id"]: p["pess"]}), folha)["vpl"]
        vo = calcula(f, dict(valores(prem, "provavel"), **{p["id"]: p["otim"]}), folha)["vpl"]
        tornado.append(dict(id=p["id"], nome=p["nome"], vpl_pess=vp, vpl_otim=vo, amplitude=abs(vo - vp)))
    tornado.sort(key=lambda t: -t["amplitude"])
    por_id = {p["id"]: p for p in prem}
    rupt = [dict(por_id[t["id"]], ruptura=ruptura(f, prem, folha, t["id"])) for t in tornado[:5]]

    wb = planilha(f, prem, quadro, resultados, tornado, rupt)
    SAIDA.parent.mkdir(parents=True, exist_ok=True)
    wb.save(SAIDA)
    for i, cen in enumerate(CENARIOS, 1):
        vpl_lo = recalcula(SAIDA, i)
        dif = abs(vpl_lo - resultados[cen]["vpl"])
        print(f"Conferência {cen}: VPL Python {resultados[cen]['vpl']:,.2f} × LibreOffice {vpl_lo:,.2f} (diferença {dif:.4f})")
        if dif >= 1:
            raise SystemExit("VPL da planilha diverge do Python")
    figura_tornado(tornado, base["vpl"])

    # números citados no capítulo
    print(f"\nA&B por pessoa nas notas: R$ {f['ab_por_pessoa']:.2f}; valor médio por nota: R$ {f['valor_por_nota']:.2f}")
    print(f"Folha de referência: R$ {folha:,.2f} ({sum(q for _, _, q, _, _ in quadro)} pessoas)")
    print(f"Juro real: {f['juro_real']:.2f}%; PAS mercadorias: {f['pas_mercadorias_pct']:.1f}%; taxa B1 mediana {f['taxa_b1_mediana']} (n={f['taxa_b1_n']})")
    for cen in CENARIOS:
        r = resultados[cen]
        print(f"\n[{cen}] invest {r['investimento']:,.0f} | receita madura {r['receita_madura']:,.0f}/mês | pessoas {r['pessoas_madura']:,.0f}/mês"
              f" | fixos {r['fixos']:,.0f} | MC {r['mc_pct']:.1%} | equilíbrio {r['receita_equilibrio']:,.0f} | EBITDA5 {r['ebitda5']:,.0f}"
              f" | VPL {r['vpl']:,.0f} | TIR {r['tir']:.1%} | payback {r['payback']:.2f}")
        print("   inv:", {k: round(x) for k, x in r["inv"].items()})
        print("   anos:", r["ano"].round(0).to_dict("index"))
    print("\nTornado:")
    for t in tornado:
        print(f"  {t['id']} {t['nome']}: pess {t['vpl_pess']:,.0f} | otim {t['vpl_otim']:,.0f} | amplitude {t['amplitude']:,.0f}")
    print("\nRuptura:")
    for rp in rupt:
        print(f"  {rp['id']} {rp['nome']}: provável {rp['prov']} → ruptura {rp['ruptura']}")
    print(f"\nOK: {SAIDA.relative_to(ROOT)} e {FIGURA.relative_to(ROOT)}")


if __name__ == "__main__":
    main()
