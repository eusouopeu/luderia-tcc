"""
Estilo comum das planilhas de dados agregados (tabelas e gráficos).

Tabelas no padrão do TCC (Arial 10, fundo branco, bordas só no cabeçalho e no
fim) e gráficos no desenho do trabalho de marketing do autor (rótulo em negrito
na ponta da barra, legenda acima), com cores só das paletas slate e purple do
Tailwind CSS (pedido do autor, 30/10/2026).
Usado por `gerar_planilha_agregada.py` e `gerar_planilha_comentarios.py`;
cada script define `planilha_estilo.GRAF` (pasta dos PNGs) antes de desenhar.
"""

from __future__ import annotations

from pathlib import Path

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
from matplotlib.patches import Patch
from matplotlib.ticker import FuncFormatter, PercentFormatter
from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.formula import ArrayFormula

GRAF = Path(".")

# ---------------------------------------------------------------- gráficos
# Tailwind CSS v3: purple-800, purple-500, purple-300, slate-300 e slate-600 (pedido do autor, 30/10/2026)
ROXO_800, ROXO_500, ROXO_300, SLATE_300, SLATE_600 = "#6B21A8", "#A855F7", "#D8B4FE", "#CBD5E1", "#475569"
SLATE_400, ROXO_700, SLATE_700 = "#94A3B8", "#7E22CE", "#334155"
PALETA = [ROXO_800, ROXO_500, ROXO_300, SLATE_300, SLATE_600]
CLARAS = {ROXO_300, SLATE_300}  # rótulo em slate-700 para ser legível sobre o branco


def cor_rotulo(cor):
    return SLATE_700 if cor in CLARAS else cor


CINZA_GRADE = "#E2E8F0"  # slate-200

plt.rcParams.update({
    "font.family": "Liberation Sans",  # métrica igual à da Arial
    "font.size": 11,
    "axes.edgecolor": "#334155",
    "axes.linewidth": 0.8,
    "figure.dpi": 300,
    "savefig.dpi": 300,
    "figure.facecolor": "white",
    "axes.facecolor": "white",
})


def br(v, casas=2, pre="", suf=""):
    """Número no formato brasileiro (vírgula decimal, ponto de milhar)."""
    if v is None or (isinstance(v, float) and np.isnan(v)):
        return ""
    t = f"{v:,.{casas}f}".replace(",", "X").replace(".", ",").replace("X", ".")
    return f"{pre}{t}{suf}"


def _quebra(t, larg=22):
    palavras, linhas, atual = str(t).split(), [], ""
    for p in palavras:
        if len(atual) + len(p) + 1 > larg and atual:
            linhas.append(atual)
            atual = p
        else:
            atual = f"{atual} {p}".strip()
    linhas.append(atual)
    return "\n".join(linhas)


def _estilo_eixo(ax, eixo_valor="x"):
    for s in ("top", "right"):
        ax.spines[s].set_visible(False)
    if eixo_valor == "x":
        ax.spines["bottom"].set_visible(False)
        ax.xaxis.grid(True, color=CINZA_GRADE, linewidth=0.8)
        ax.tick_params(axis="y", length=0)
        ax.tick_params(axis="x", length=0, colors="#334155")
    else:
        ax.spines["left"].set_visible(False)
        ax.yaxis.grid(True, color=CINZA_GRADE, linewidth=0.8)
        ax.tick_params(axis="x", length=0)
        ax.tick_params(axis="y", length=0, colors="#334155")
    ax.set_axisbelow(True)


def _legenda_topo(fig, ax, nomes, cores):
    alcas = [Patch(facecolor=c, label=n) for n, c in zip(nomes, cores)]
    ax.legend(handles=alcas, loc="lower center", bbox_to_anchor=(0.5, 1.01), ncol=len(nomes),
              frameon=False, handlelength=0.9, handleheight=0.9, prop={"weight": "bold", "size": 10})


def barras_h(nome, rotulos, series, fmt, cores=None, destaque=None, xmax=None, eixo_pct=False, larg_rotulo=22):
    """Barras horizontais (uma ou mais séries), rótulo em negrito na ponta, na cor da série.

    series: dict nome -> lista de valores (mesma ordem de rotulos, de cima para baixo).
    destaque: rótulo cuja barra fica em purple-700, com as demais em slate-400 (só com uma série; ex.: Salvador).
    """
    nomes = list(series)
    cores = cores or PALETA[: len(nomes)]
    n, k = len(rotulos), len(nomes)
    alt_total = 0.72
    h = alt_total / k
    fig, ax = plt.subplots(figsize=(8, max(2.6, 0.32 * n * max(1, k * 0.8) + 1.0)))
    y = np.arange(n)[::-1]
    vmax = np.nanmax([np.nanmax(np.array(v, dtype=float)) for v in series.values()])
    for i, (s, cor) in enumerate(zip(nomes, cores)):
        vals = np.array(series[s], dtype=float)
        pos = y + alt_total / 2 - h * (i + 0.5)
        cor_barras = [cor] * n
        if destaque is not None and k == 1:
            cor_barras = [ROXO_700 if r == destaque else SLATE_400 for r in rotulos]
        ax.barh(pos, np.nan_to_num(vals), height=h * 0.92, color=cor_barras)
        for p, v, cb in zip(pos, vals, cor_barras):
            if np.isnan(v):
                continue
            ax.text(v + vmax * 0.015, p, fmt(v), va="center", ha="left", fontsize=9.5,
                    fontweight="bold", color=SLATE_700 if cb == SLATE_400 else cor_rotulo(cb))
    ax.set_yticks(y)
    ax.set_yticklabels([_quebra(r, larg_rotulo) for r in rotulos], fontsize=10, fontweight="bold", color="#1E293B")
    ax.set_xlim(0, xmax or vmax * 1.18)
    if eixo_pct:
        ax.xaxis.set_major_formatter(PercentFormatter(1.0, decimals=0))
    else:
        ax.xaxis.set_major_formatter(FuncFormatter(lambda v, _: br(v, 1 if vmax < 10 else 0)))
    _estilo_eixo(ax, "x")
    if k > 1:
        _legenda_topo(fig, ax, nomes, cores)
    return _salvar(fig, nome)


def barras_v(nome, rotulos, series, fmt, cores=None, eixo_pct=False, ymax=None):
    """Barras verticais agrupadas, rótulo em negrito acima da barra."""
    nomes = list(series)
    cores = cores or PALETA[: len(nomes)]
    n, k = len(rotulos), len(nomes)
    larg = 0.75 / k
    fig, ax = plt.subplots(figsize=(8, 4.2))
    x = np.arange(n)
    vmax = np.nanmax([np.nanmax(np.array(v, dtype=float)) for v in series.values()])
    for i, (s, cor) in enumerate(zip(nomes, cores)):
        vals = np.array(series[s], dtype=float)
        pos = x - 0.375 + larg * (i + 0.5)
        ax.bar(pos, np.nan_to_num(vals), width=larg * 0.92, color=cor)
        for p, v in zip(pos, vals):
            if not np.isnan(v):
                ax.text(p, v + vmax * 0.015, fmt(v), ha="center", va="bottom", fontsize=9.5,
                        fontweight="bold", color=cor_rotulo(cor))
    ax.set_xticks(x)
    ax.set_xticklabels([_quebra(r, 14) for r in rotulos], fontsize=10, fontweight="bold", color="#1E293B")
    ax.set_ylim(0, ymax or vmax * 1.2)
    if eixo_pct:
        ax.yaxis.set_major_formatter(PercentFormatter(1.0, decimals=0))
    else:
        ax.yaxis.set_major_formatter(FuncFormatter(lambda v, _: br(v, 1 if vmax < 10 else 0)))
    _estilo_eixo(ax, "y")
    if k > 1:
        _legenda_topo(fig, ax, nomes, cores)
    return _salvar(fig, nome)


def empilhado_100(nome, rotulos, partes, cores=None):
    """Barras horizontais empilhadas em 100%, rótulo branco dentro de cada parte (como no documento)."""
    nomes = list(partes)
    cores = cores or PALETA[: len(nomes)]
    n = len(rotulos)
    fig, ax = plt.subplots(figsize=(8, max(2.4, 0.9 * n + 1.0)))
    y = np.arange(n)[::-1]
    tot = np.sum([np.array(partes[p], dtype=float) for p in nomes], axis=0)
    esq = np.zeros(n)
    for p, cor in zip(nomes, cores):
        frac = np.array(partes[p], dtype=float) / tot
        ax.barh(y, frac, left=esq, height=0.62, color=cor)
        for yi, e, f in zip(y, esq, frac):
            if f >= 0.04:
                cor_txt = SLATE_700 if cor in CLARAS else "white"
                ax.text(e + f - 0.01, yi, br(f * 100, 0, suf="%"), ha="right", va="center",
                        fontsize=9.5, fontweight="bold", color=cor_txt)
        esq += frac
    ax.set_yticks(y)
    ax.set_yticklabels([_quebra(r, 18) for r in rotulos], fontsize=10, fontweight="bold", color="#1E293B")
    ax.set_xlim(0, 1)
    ax.xaxis.set_major_formatter(PercentFormatter(1.0, decimals=0))
    ax.set_xticks([0, 0.25, 0.5, 0.75, 1])
    _estilo_eixo(ax, "x")
    _legenda_topo(fig, ax, nomes, cores)
    return _salvar(fig, nome)


def linhas(nome, x, series, fmt, cores=None, tracejar_depois=None):
    nomes = list(series)
    cores = cores or PALETA[: len(nomes)]
    fig, ax = plt.subplots(figsize=(8, 4.2))
    for s, cor in zip(nomes, cores):
        v = np.array(series[s], dtype=float)
        ax.plot(x, v, color=cor, linewidth=2.4, marker="o", markersize=4)
        ult = np.where(~np.isnan(v))[0][-1]
        ax.text(x[ult], v[ult], "  " + fmt(v[ult]), color=cor_rotulo(cor), fontweight="bold", fontsize=9.5, va="center")
    if tracejar_depois is not None:
        ax.axvspan(tracejar_depois, x[-1] + 0.5, color="#F8FAFC", zorder=0)  # slate-50
        ax.text(tracejar_depois + 0.15, ax.get_ylim()[1] * 0.96, "projeção (Focus)", fontsize=9,
                color="#64748B", va="top")
    ax.yaxis.set_major_formatter(FuncFormatter(lambda v, _: br(v, 0, suf="%")))
    ax.tick_params(axis="x", labelsize=10)
    for t in ax.get_xticklabels():
        t.set_fontweight("bold")
    _estilo_eixo(ax, "y")
    _legenda_topo(fig, ax, nomes, cores)
    return _salvar(fig, nome)


def _salvar(fig, nome):
    pasta = globals()["GRAF"]
    pasta.mkdir(parents=True, exist_ok=True)
    caminho = pasta / f"{nome}.png"
    fig.savefig(caminho, dpi=300, bbox_inches="tight", facecolor="white")
    plt.close(fig)
    return caminho


# ---------------------------------------------------------------- planilha
ARIAL = Font(name="Arial", size=10, color="000000")
ARIAL_N = Font(name="Arial", size=10, color="000000", bold=True)
BRANCO = PatternFill("solid", start_color="FFFFFF", end_color="FFFFFF")
FINA = Side(style="thin", color="000000")
SEM = Side(style="thin", color="FFFFFF")
FMT = {"int": "#,##0", "dec": "#,##0.00", "pct": "0.00%", "txt": "@"}

wb = Workbook()
wb.remove(wb.active)
INDICE: list[tuple[str, str, str]] = []  # (aba, título, fonte)


def nova_aba(aba, titulo, cabecalho, linhas_, tipos, fonte, larguras=None, grafico=None, notas=()):
    """Escreve uma tabela: linha 1 título, linha 2 cabeçalho, dados, e rodapé com fonte.

    tipos: um por coluna ("txt", "int", "dec", "pct"); célula que começa com "=" vira fórmula.
    """
    titulo = titulo.split(":")[0].rstrip()  # só o nome da tabela (pedido do autor, 30/10/2026)
    ws = wb.create_sheet(aba)
    ws.sheet_view.showGridLines = False
    nc = len(cabecalho)
    ws.cell(1, 1, titulo)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=nc)
    for j, c in enumerate(cabecalho, 1):
        ws.cell(2, j, c)
    for i, lin in enumerate(linhas_, 3):
        for j, v in enumerate(lin, 1):
            if isinstance(v, ArrayFormula):
                v.ref = f"{get_column_letter(j)}{i}"
            if isinstance(v, float) and np.isnan(v):
                v = None
            ws.cell(i, j, v)
    ult = 2 + len(linhas_)
    for r in range(1, ult + 1):
        for j in range(1, nc + 1):
            cel = ws.cell(r, j)
            cel.fill = BRANCO
            if r == 1:
                cel.font, cel.alignment = ARIAL_N, Alignment(horizontal="center", vertical="center", wrap_text=True)
                continue
            if r == 2:
                cel.font = ARIAL_N
                cel.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
                cel.border = Border(top=FINA, bottom=FINA, left=SEM, right=SEM)
                continue
            t = tipos[j - 1]
            cel.font = ARIAL
            cel.alignment = Alignment(horizontal="left" if t == "txt" else "right", vertical="center", wrap_text=t == "txt")
            if t != "txt":
                cel.number_format = FMT[t]
            cel.border = Border(top=SEM, left=SEM, right=SEM, bottom=FINA if r == ult else SEM)
    ws.row_dimensions[1].height = 30
    ws.row_dimensions[2].height = 42
    rod = ult + 1
    for t in (f"Fonte: {fonte}", *notas):
        ws.cell(rod, 1, t).font = ARIAL
        ws.cell(rod, 1).alignment = Alignment(horizontal="left", vertical="top", wrap_text=False)
        rod += 1
    larguras = larguras or [28] + [14] * (nc - 1)
    for j, w in enumerate(larguras, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    if grafico is not None:
        img = XLImage(str(grafico))
        escala = 620 / img.width
        img.width, img.height = img.width * escala, img.height * escala
        ws.add_image(img, f"{get_column_letter(nc + 2)}2")
    INDICE.append((aba, titulo, fonte))
    return ws




def escrever_indice(wb_, titulo, itens, notas=()):
    """Aba "Índice" na frente: aba (com link), tabela e fonte; notas abaixo."""
    if "Índice" in wb_.sheetnames:
        del wb_["Índice"]
    ind = wb_.create_sheet("Índice", 0)
    ind.sheet_view.showGridLines = False
    ind.cell(1, 1, titulo)
    ind.merge_cells("A1:C1")
    for j, c in enumerate(["Aba", "Tabela", "Fonte"], 1):
        ind.cell(2, j, c)
    for i, (a, t, f) in enumerate(itens, 3):
        ind.cell(i, 1, a).hyperlink = f"#'{a}'!A1"
        ind.cell(i, 2, t)
        ind.cell(i, 3, f)
    ult = 2 + len(itens)
    for rr in range(1, ult + 1):
        for j in range(1, 4):
            cel = ind.cell(rr, j)
            cel.fill = BRANCO
            cel.font = ARIAL_N if rr <= 2 else ARIAL
            cel.alignment = Alignment(horizontal="center" if rr <= 2 else "left", vertical="center", wrap_text=rr > 2)
            if rr == 2:
                cel.border = Border(top=FINA, bottom=FINA, left=SEM, right=SEM)
            elif rr > 2:
                cel.border = Border(top=SEM, left=SEM, right=SEM, bottom=FINA if rr == ult else SEM)
    for j, w in enumerate([10, 70, 70], 1):
        ind.column_dimensions[get_column_letter(j)].width = w
    for k, t in enumerate(notas, ult + 2):
        ind.cell(k, 1, t).font = ARIAL


def dividir(arq_total, indice, saidas):
    """Separa a pasta completa em arquivos (um por bloco), renumerando tabelas e abas a partir de 1.

    indice: lista (aba, título, fonte) da pasta completa.
    saidas: lista (caminho, título do índice, abas de tabela, abas extras, notas do índice).
    """
    import re
    from openpyxl import load_workbook

    por_aba = {a: (t, f) for a, t, f in indice}
    feitos = []
    for caminho, titulo, abas, extras, notas in saidas:
        wb_ = load_workbook(arq_total)
        for nome in list(wb_.sheetnames):
            if nome not in abas and nome not in [e for e, _, _ in extras]:
                del wb_[nome]
        itens = []
        for k, aba in enumerate(abas, 1):
            ws = wb_[aba]
            novo_t = re.sub(r"^Tabela \d+", f"Tabela {k}", por_aba[aba][0])
            ws.cell(1, 1).value = novo_t
            ws.title = f"Tab {k}"
            itens.append((ws.title, novo_t, por_aba[aba][1]))
        for e, t, f in extras:
            itens.append((e, t, f))
        escrever_indice(wb_, titulo, itens, notas)
        wb_.save(caminho)
        feitos.append((caminho, len(abas)))
    return feitos
