"""
Planilha separada com os dados agregados dos comentários do B2, com gráficos.

Lê a cópia exportada e atualizada do banco "2 - Comentários" do Notion
(`2026-10-30_2_comentarios_all_atualizado.csv`, temas codificados em 30/10) e
grava `dados/Dados agregados/dados_agregados_comentarios.xlsx`, uma aba por
tabela com o gráfico ao lado, e os PNGs em `graficos_comentarios/`.

A aba "Dados B2" traz uma linha por comentário, sem o texto (texto integral não
é publicado; `_instrucoes/4_metodologia.md` §3.4): casa, metrópole, notas,
campos extras do Maps e uma coluna 0/1 por tema positivo, tema negativo,
contexto e tamanho de grupo. As tabelas calculam por fórmula sobre ela.

Uso: python3 _src/extras/gerar_planilha_comentarios.py [caminho do CSV]
Sem argumento, procura o CSV em `dados/[B2]…/_dados-brutos/notion_comentarios/`
e, se não achar, nos arquivos do projeto (`/mnt/project-files/b2/…`).
"""

from __future__ import annotations

import re
import sys
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl.styles import Alignment, Border
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / "_src" / "extras"))
from log_execucao import registrar_execucao  # noqa: E402
import planilha_estilo as est  # noqa: E402
from planilha_estilo import *  # noqa: E402,F403

D = ROOT / "dados"
B1 = D / "[B1] Estabelecimentos e cardapios"
NOME_CSV = "2026-10-30_2_comentarios_all_atualizado.csv"
CANDIDATOS = [D / "[B2] Avaliacoes recentes" / "_dados-brutos" / "notion_comentarios" / NOME_CSV,
              Path("/mnt/project-files/b2/notion_comentarios") / NOME_CSV]
SAIDA = D / "Dados agregados"
GRAF = SAIDA / "graficos_comentarios"

registrar_execucao(__file__)
est.GRAF = GRAF
arq = Path(sys.argv[1]) if len(sys.argv) > 1 else next((c for c in CANDIDATOS if c.exists()), None)
if arq is None or not arq.exists():
    sys.exit(f"CSV do '2 - Comentários' não encontrado; passe o caminho ou copie para {CANDIDATOS[0]}")
print(f"Lendo {arq}")
d = pd.read_csv(arq)

# ---------------------------------------------------------------- base
res = pd.read_excel(B1 / "estabelecimentos" / "B1_resumo_estabelecimentos.xlsx", sheet_name="estabelecimentos")
cap = d["Place ID"].map(dict(zip(res.place_id, res.capital_endereco)))
cap = cap.fillna(d.Estabelecimento.map(dict(zip(res.nome, res.capital_endereco))))
d["metropole"] = cap.fillna(d["Capital do endereço"])
REGIAO = {"Belém": "Norte", "Fortaleza": "Nordeste", "Recife": "Nordeste", "Salvador": "Nordeste",
          "Brasília": "Centro-Oeste", "Goiânia": "Centro-Oeste", "São Paulo": "Sudeste", "Rio de Janeiro": "Sudeste",
          "Belo Horizonte": "Sudeste", "Vitória": "Sudeste", "Curitiba": "Sul", "Florianópolis": "Sul",
          "Porto Alegre": "Sul"}
assert d.metropole.notna().all(), d[d.metropole.isna()].Estabelecimento.unique()

FAIXAS_PRECO = [(1, 20), (20, 40), (40, 60), (60, 80), (80, 100), (100, 120), (120, 140), (140, 160), (160, 180),
                (180, 200), (200, None)]
ROT_PRECO = [f"R$ {a} a {b}" if b else "R$ 200 ou mais" for a, b in FAIXAS_PRECO]


def faixa_preco(v):
    """Faixa do Maps; valor único digitado à mão vai para a faixa que o contém (50 → 40 a 60)."""
    if pd.isna(v):
        return None
    t = str(v).replace("–", "-").strip()
    if "200" in t and ("+" in t):
        return ROT_PRECO[-1]
    m = re.fullmatch(r"(\d+)-(\d+)", t)
    if m:
        a = int(m.group(1))
        return next(r for (x, _), r in zip(FAIXAS_PRECO, ROT_PRECO) if x == a)
    x = int(t)
    for (a, b), r in zip(FAIXAS_PRECO, ROT_PRECO):
        if b is None or a <= x < b:
            return r
    return None


d["preco_faixa"] = d["Preço por pessoa (R$)"].map(faixa_preco)
ESPERA = ["Sem espera", "Até 10 minutos", "De 10 a 30 minutos", "De 30 a 60 minutos", "Mais de 1 hora"]
d["espera"] = d["Tempo de espera"].map({"Sem espera": ESPERA[0], "até 10 minutos": ESPERA[1],
                                        "10 a 30 minutos": ESPERA[2], "30 a 60 minutos": ESPERA[3],
                                        "mais de 1 hora": ESPERA[4], "1 hora": ESPERA[4]})
BARULHO = ["Silencioso", "Moderado", "Alto; conversa possível", "Muito alto; conversa difícil"]
d["barulho"] = d["Nível de barulho"]

GRUPOS = [("g_1", "1 pessoa"), ("g_2", "2 pessoas"), ("g_3_4", "3 a 4 pessoas"), ("g_5_8", "5 a 8 pessoas"),
          ("g_9", "9 pessoas ou mais"), ("g_todos", "Todos os tamanhos")]


def grupos(v):
    """Tamanho do grupo (múltipla escolha no Maps, com variações digitadas à mão) em marcas 0/1."""
    out = dict.fromkeys([g for g, _ in GRUPOS], 0)
    if pd.isna(v):
        return out
    for p in [x.strip().lower() for x in str(v).split(",")]:
        if "todos" in p or "adequado" in p:
            out["g_todos"] = 1
        elif p.startswith("1 pessoa"):
            out["g_1"] = 1
        elif re.match(r"(de )?2 pessoas", p):
            out["g_2"] = 1
        elif re.match(r"(de )?3 a 4", p):
            out["g_3_4"] = 1
        elif re.match(r"(de )?5 a 8", p):
            out["g_5_8"] = 1
        elif p.startswith("mais de 9"):
            out["g_9"] = 1
        elif m := re.match(r"(\d) a (\d)", p):  # faixas largas: "3 a 8", "2 a 4", "2 a 9+", "5 a 9"
            a, b = int(m.group(1)), int(m.group(2))
            for g, lo, hi in (("g_2", 2, 2), ("g_3_4", 3, 4), ("g_5_8", 5, 8), ("g_9", 9, 99)):
                if a <= hi and b >= lo:
                    out[g] = 1
    return out


gr = pd.DataFrame([grupos(v) for v in d["Tamanho do grupo"]])
d["grupo_informado"] = (gr.sum(axis=1) > 0).astype(int)


def itens(v):
    return [t.strip() for t in str(v).split(",") if t.strip()] if pd.notna(v) else []


TEMAS = sorted({t for c in ("Itens positivos", "Itens negativos") for v in d[c] for t in itens(v)})
CONTEXTOS = ["amigos", "família", "casal", "aniversário", "hobbyista", "iniciante", "loja", "aluguel"]
pos = pd.DataFrame({t: d["Itens positivos"].map(lambda v, t=t: int(t in itens(v))) for t in TEMAS})
neg = pd.DataFrame({t: d["Itens negativos"].map(lambda v, t=t: int(t in itens(v))) for t in TEMAS})
ctx = pd.DataFrame({c: d.Contexto.map(lambda v, c=c: int(c in itens(v))) for c in CONTEXTOS})
d["com_tema"] = ((pos.sum(axis=1) + neg.sum(axis=1)) > 0).astype(int)

base = pd.DataFrame({
    "estabelecimento": d.Estabelecimento, "metropole": d.metropole, "regiao": d.metropole.map(REGIAO),
    "nota": d["Nota geral"], "comida": d.Comida, "servico": d["Serviço"], "ambiente": d.Ambiente,
    "preco_faixa": d.preco_faixa, "espera": d.espera, "barulho": d.barulho, "grupo_informado": d.grupo_informado,
    "analise_temas": d["Análise de temas"], "com_tema": d.com_tema,
})
for g, _ in GRUPOS:
    base[g] = gr[g]
for c in CONTEXTOS:
    base[f"ctx_{c}"] = ctx[c]
for i, t in enumerate(TEMAS):
    base[f"pos_{i:02d}"] = pos[t]
for i, t in enumerate(TEMAS):
    base[f"neg_{i:02d}"] = neg[t]
base = base.sort_values(["regiao", "metropole", "estabelecimento"], key=lambda s: s.map(
    {"Norte": 0, "Nordeste": 1, "Centro-Oeste": 2, "Sudeste": 3, "Sul": 4}) if s.name == "regiao" else s).reset_index(drop=True)
N = len(base)
ROTULO = {"estabelecimento": "casa", "metropole": "metrópole", "regiao": "região", "nota": "nota geral",
          "comida": "nota de comida", "servico": "nota de serviço", "ambiente": "nota de ambiente",
          "preco_faixa": "preço por pessoa", "espera": "tempo de espera", "barulho": "nível de barulho",
          "grupo_informado": "tamanho do grupo informado", "analise_temas": "análise de temas", "com_tema": "com tema"}
ROTULO.update({g: f"grupo: {r}" for g, r in GRUPOS})
ROTULO.update({f"ctx_{c}": f"contexto: {c}" for c in CONTEXTOS})
ROTULO.update({f"pos_{i:02d}": f"+ {t}" for i, t in enumerate(TEMAS)})
ROTULO.update({f"neg_{i:02d}": f"− {t}" for i, t in enumerate(TEMAS)})
TIPO_TXT = {"estabelecimento", "metropole", "regiao", "preco_faixa", "espera", "barulho", "analise_temas"}
LETRA = {c: get_column_letter(i) for i, c in enumerate(base.columns, 1)}
ABA_BASE = "Dados B2"


def R(col):
    return f"'{ABA_BASE}'!${LETRA[col]}$3:${LETRA[col]}${N + 2}"


NT = f'COUNTA({R("estabelecimento")})'
FONTE = ("B2, avaliações do Google Maps até \"3 meses atrás\" (coletas de 09 a 30/10/2026), banco \"2 - Comentários\" "
         "do Notion, cópia atualizada de 30/10/2026")
FONTE_TEMAS = FONTE + "; temas pelo dicionário de 30/10/2026 (textos/instrumentos/dicionario_temas_B2.md)"
N_CASAS = base.estabelecimento.nunique()

# Tab 1: comentários por metrópole
ORD = ["Norte", "Nordeste", "Centro-Oeste", "Sudeste", "Sul"]
mt = base.groupby(["regiao", "metropole"]).agg(n=("nota", "size")).reset_index()
mt["o"] = mt.regiao.map(ORD.index)
mt = mt.sort_values(["o", "n"], ascending=[True, False])
lin1 = []
for k, r in enumerate(mt.itertuples(), 3):
    ref = f"$B{k}"
    lin1.append([r.regiao, r.metropole,
                 f'=SUMPRODUCT(({R("metropole")}={ref})/COUNTIFS({R("estabelecimento")},{R("estabelecimento")}))',
                 f'=COUNTIF({R("metropole")},{ref})', f"=D{k}/D${3 + len(mt)}",
                 f'=AVERAGEIFS({R("nota")},{R("metropole")},{ref})',
                 f'=COUNTIFS({R("metropole")},{ref},{R("nota")},"<=3")/D{k}'])
t1 = 3 + len(mt)
lin1.append(["Total", "", f"=SUM(C3:C{t1 - 1})", f"=SUM(D3:D{t1 - 1})", f"=SUM(E3:E{t1 - 1})",
             f'=AVERAGE({R("nota")})', f'=COUNTIF({R("nota")},"<=3")/D{t1}'])
lin1.append(["Conferência (total − comentários da base; deve ser 0)", "", "", f"=D{t1}-{NT}", "", "", ""])
m1 = mt.sort_values("n", ascending=False)
g1 = barras_h("01_comentarios_por_metropole", m1.metropole.tolist(), {"comentários": m1.n.tolist()},
              lambda v: br(v, 0), destaque="Salvador")
nova_aba("Tab 1", f"Tabela 1 – Comentários por metrópole: {br(N, 0)} comentários de {N_CASAS} casas, mais da metade de São Paulo",
         ["Região", "Metrópole", "Casas", "Comentários", "% dos comentários", "Nota média", "% com nota 1 a 3"],
         lin1, ["txt", "txt", "int", "int", "pct", "dec", "pct"], FONTE, [16, 18, 9, 12, 12, 11, 12], g1,
         ["Nota: cada comentário conta uma vez, na metrópole do endereço da casa."])

# Tab 2: nota geral
lin2 = [[f"{n} estrela" + ("s" if n > 1 else ""), f'=COUNTIF({R("nota")},{n})', f"=B{k}/B$8"]
        for k, n in enumerate([5, 4, 3, 2, 1], 3)]
lin2 += [["Total", "=SUM(B3:B7)", "=SUM(C3:C7)"], ["Conferência (deve ser 0)", f"=B8-{NT}", ""]]
vc = base.nota.value_counts()
g2 = barras_v("02_nota_geral", ["5 estrelas", "4 estrelas", "3 estrelas", "2 estrelas", "1 estrela"],
              {"% dos comentários": [vc.get(n, 0) / N for n in (5, 4, 3, 2, 1)]}, lambda v: br(v * 100, 1, suf="%"),
              eixo_pct=True)
nova_aba("Tab 2", f"Tabela 2 – Nota geral: {vc.get(5, 0) / N:.0%} dos comentários dão 5 estrelas".replace(".", ","),
         ["Nota geral", "Comentários", "% dos comentários"], lin2, ["txt", "int", "pct"], FONTE, [26, 12, 14], g2)

# Tab 3: notas por aspecto
ASP = [("comida", "Comida"), ("servico", "Serviço"), ("ambiente", "Ambiente")]
lin3 = []
for k, (c, rot) in enumerate(ASP, 3):
    lin3.append([rot, f"=COUNT({R(c)})", f"=AVERAGE({R(c)})"] +
                [f"=COUNTIF({R(c)},{n})/B{k}" for n in (5, 4, 3, 2, 1)])
g3 = empilhado_100("03_notas_por_aspecto", [r for _, r in ASP],
                   {f"{n} estrela" + ("s" if n > 1 else ""): [(base[c] == n).sum() for c, _ in ASP] for n in (5, 4, 3, 2, 1)})
nova_aba("Tab 3", "Tabela 3 – Notas por aspecto: serviço e ambiente com mais 5 estrelas que a comida",
         ["Aspecto", "Comentários com nota", "Nota média", "5 estrelas", "4 estrelas", "3 estrelas", "2 estrelas",
          "1 estrela"], lin3, ["txt", "int", "dec", "pct", "pct", "pct", "pct", "pct"], FONTE, [14, 13, 11, 11, 11, 11, 11, 11],
         g3, ["Nota: só comentários com nota do aspecto; inclui as notas digitadas à mão pela convenção \"5 (comida 3)\" (aspecto não citado = nota geral)."])

# Tab 4: temas positivos e negativos
LT = f'SUM({R("com_tema")})'
ordem = sorted(range(len(TEMAS)), key=lambda i: -(pos.iloc[:, i].sum() + neg.iloc[:, i].sum()))
lin4 = []
for k, i in enumerate(ordem, 3):
    lin4.append([TEMAS[i], f'=SUM({R(f"pos_{i:02d}")})', f'=SUM({R(f"neg_{i:02d}")})', f"=B{k}+C{k}",
                 f"=B{k}/{LT}", f"=C{k}/{LT}", f'=IFERROR(C{k}/D{k},"")'])
nt = int(d.com_tema.sum())
g4 = barras_h("04_temas_positivos_negativos", [TEMAS[i] for i in ordem],
              {"Positivo": [pos.iloc[:, i].sum() / nt for i in ordem], "Negativo": [neg.iloc[:, i].sum() / nt for i in ordem]},
              lambda v: br(v * 100, 1, suf="%"), cores=[ROXO_800, SLATE_600], eixo_pct=True, larg_rotulo=24)
nova_aba("Tab 4", "Tabela 4 – Temas dos comentários: atendimento, comida, ambiente e acervo são os mais citados",
         ["Tema", "Comentários que elogiam", "Comentários que criticam", "Menções", "% dos comentários com tema (elogio)",
          "% dos comentários com tema (crítica)", "% de críticas nas menções"], lin4,
         ["txt", "int", "int", "int", "pct", "pct", "pct"], FONTE_TEMAS, [26, 12, 12, 10, 14, 14, 13], g4,
         [f"Nota: temas não exclusivos; base = {nt} comentários com pelo menos um tema (de {N}).",
          "Codificação: 355 linhas pelo autor (antes de 30/10) e 1.273 por Claude em 30/10/2026; recodificação cega do autor pendente."])

# Tab 5: críticas nos comentários de nota 1 a 3
ruins = d["Nota geral"] <= 3  # mesma ordem de pos/neg (a base é reordenada)
NR = f'COUNTIF({R("nota")},"<=3")'
ordem5 = sorted(range(len(TEMAS)), key=lambda i: -neg[ruins.values].iloc[:, i].sum())
ordem5 = [i for i in ordem5 if neg[ruins.values].iloc[:, i].sum() > 0]
lin5 = [[TEMAS[i], f'=SUMIFS({R(f"neg_{i:02d}")},{R("nota")},"<=3")', f"=B{k}/{NR}",
         f'=SUMIFS({R(f"neg_{i:02d}")},{R("nota")},">=4")'] for k, i in enumerate(ordem5, 3)]
g5 = barras_h("05_criticas_notas_baixas", [TEMAS[i] for i in ordem5],
              {"% dos comentários de nota 1 a 3": [neg[ruins.values].iloc[:, i].sum() / ruins.sum() for i in ordem5]},
              lambda v: br(v * 100, 1, suf="%"), cores=[SLATE_600], eixo_pct=True, larg_rotulo=24)
nova_aba("Tab 5", f"Tabela 5 – O que derruba a nota: atendimento, comida, preço e rapidez lideram as críticas nos {int(ruins.sum())} comentários de nota 1 a 3",
         ["Tema criticado", "Comentários de nota 1 a 3", "% dos comentários de nota 1 a 3", "Críticas em comentários de nota 4 ou 5"],
         lin5, ["txt", "int", "pct", "int"], FONTE_TEMAS, [26, 14, 16, 16], g5,
         ["Nota: temas não exclusivos; um comentário pode criticar vários temas."])

# Tab 6: contexto
lin6 = [[c.capitalize(), f'=SUM({R(f"ctx_{c}")})', f"=B{k}/{NT}"] for k, c in enumerate(CONTEXTOS, 3)]
cs = ctx.sum().sort_values(ascending=False)
g6 = barras_h("06_contexto", [c.capitalize() for c in cs.index], {"comentários": cs.tolist()}, lambda v: br(v, 0))
nova_aba("Tab 6", "Tabela 6 – Contexto citado: amigos e loja aparecem mais que família, aniversário ou casal",
         ["Contexto", "Comentários", "% dos comentários"], lin6, ["txt", "int", "pct"], FONTE_TEMAS, [18, 12, 14], g6,
         ["Nota: contextos não exclusivos e sem polaridade (ocasião, público, aluguel e loja)."])

# Tab 7: preço por pessoa
lin7 = [[r, f'=COUNTIF({R("preco_faixa")},A{k})', f"=B{k}/B${3 + len(ROT_PRECO)}"] for k, r in enumerate(ROT_PRECO, 3)]
t7 = 3 + len(ROT_PRECO)
lin7 += [["Total com preço informado", f"=SUM(B3:B{t7 - 1})", f"=SUM(C3:C{t7 - 1})"],
         ["Sem preço informado", f'={NT}-B{t7}', ""]]
pc = base.preco_faixa.value_counts()
g7 = barras_h("07_preco_por_pessoa", ROT_PRECO,
              {"% dos comentários com preço": [pc.get(r, 0) / pc.sum() for r in ROT_PRECO]},
              lambda v: br(v * 100, 1, suf="%"), eixo_pct=True)
acum = (pc.reindex(ROT_PRECO).fillna(0).cumsum() / pc.sum())
faixa_med = acum[acum >= 0.5].index[0]
nova_aba("Tab 7", f"Tabela 7 – Gasto por pessoa informado no Maps: a faixa mediana é {faixa_med}",
         ["Faixa de gasto por pessoa", "Comentários", "% dos com preço"], lin7, ["txt", "int", "pct"], FONTE,
         [26, 12, 14], g7, ["Nota: valores únicos digitados à mão entram na faixa que os contém (ex.: 50 em R$ 40 a 60)."])

# Tab 8 a 10: espera, barulho e grupo
lin8 = [[r, f'=COUNTIF({R("espera")},A{k})', f"=B{k}/B$8"] for k, r in enumerate(ESPERA, 3)]
lin8 += [["Total com espera informada", "=SUM(B3:B7)", "=SUM(C3:C7)"], ["Sem espera informada", f"={NT}-B8", ""]]
ec = base.espera.value_counts()
lin9 = [[r, f'=COUNTIF({R("barulho")},A{k})', f"=B{k}/B$7"] for k, r in enumerate(BARULHO, 3)]
lin9 += [["Total com barulho informado", "=SUM(B3:B6)", "=SUM(C3:C6)"], ["Sem barulho informado", f"={NT}-B7", ""]]
bc = base.barulho.value_counts()
g8 = barras_v("08_tempo_de_espera", ESPERA, {"% dos que informam": [ec.get(e, 0) / ec.sum() for e in ESPERA]},
              lambda v: br(v * 100, 1, suf="%"), eixo_pct=True)
nova_aba("Tab 8", "Tabela 8 – Tempo de espera: a maioria dos que informam não esperou",
         ["Tempo de espera", "Comentários", "% dos que informam"], lin8, ["txt", "int", "pct"], FONTE, [28, 12, 14], g8,
         ["Nota: campo opcional do Maps."])
g9 = barras_h("09_barulho", BARULHO, {"% dos que informam": [bc.get(b, 0) / bc.sum() for b in BARULHO]},
              lambda v: br(v * 100, 1, suf="%"), eixo_pct=True, larg_rotulo=20)
nova_aba("Tab 9", "Tabela 9 – Nível de barulho: moderado na maioria das menções; muito alto é raro",
         ["Nível de barulho", "Comentários", "% dos que informam"], lin9, ["txt", "int", "pct"], FONTE, [28, 12, 14], g9)
GI = f'SUM({R("grupo_informado")})'
lin10 = [[r, f'=SUM({R(g)})', f"=B{k}/{GI}"] for k, (g, r) in enumerate(GRUPOS, 3)]
lin10.append(["Comentários com tamanho de grupo informado", f"={GI}", ""])
ng = int(d.grupo_informado.sum())
g10 = barras_h("10_tamanho_grupo", [r for _, r in GRUPOS], {"% dos que informam": [gr[g].sum() / ng for g, _ in GRUPOS]},
               lambda v: br(v * 100, 1, suf="%"), eixo_pct=True)
nova_aba("Tab 10", "Tabela 10 – Tamanho do grupo: as casas são indicadas para grupos de 3 a 8 pessoas ou de todos os tamanhos",
         ["Tamanho do grupo", "Comentários", "% dos que informam"], lin10, ["txt", "int", "pct"], FONTE, [36, 12, 14], g10,
         ["Nota: múltipla escolha no Maps (não exclusiva); faixas digitadas à mão (ex.: \"3 a 8\") marcam as faixas que cobrem."])

# Tab 11: situação da codificação
AN = ["anterior a 30/10/2026", "Claude, 30/10/2026"]
lin11 = [[a, f'=COUNTIFS({R("analise_temas")},A{k},{R("com_tema")},1)',
          f'=COUNTIFS({R("analise_temas")},A{k},{R("com_tema")},0)', f"=B{k}+C{k}"] for k, a in enumerate(AN, 3)]
lin11 += [["Total", "=SUM(B3:B4)", "=SUM(C3:C4)", "=SUM(D3:D4)"], ["Conferência (deve ser 0)", "", "", f"=D5-{NT}"]]
g11 = empilhado_100("11_codificacao", ["Autor (antes de 30/10)", "Claude (30/10)"],
                    {"Com tema": [int(((d["Análise de temas"] == a) & (d.com_tema == 1)).sum()) for a in AN],
                     "Sem tema": [int(((d["Análise de temas"] == a) & (d.com_tema == 0)).sum()) for a in AN]})
nova_aba("Tab 11", "Tabela 11 – Codificação dos temas: quem codificou e quantos comentários ficaram sem tema",
         ["Análise de temas", "Com tema", "Sem tema", "Comentários"], lin11, ["txt", "int", "int", "int"], FONTE_TEMAS,
         [30, 12, 12, 12], g11, ["Nota: sem tema = comentário que não cita atributo (ex.: \"Excelente\"); não é o mesmo que não analisado."])

# ---------------------------------------------------------------- base e índice
ws = est.wb.create_sheet(ABA_BASE)
ws.sheet_view.showGridLines = False
ws.cell(1, 1, f"Base – os {N} comentários do B2, sem o texto (as tabelas calculam por fórmula sobre esta aba)")
for j, c in enumerate(base.columns, 1):
    ws.cell(2, j, ROTULO[c])
for i, r in enumerate(base.itertuples(index=False), 3):
    for j, v in enumerate(r, 1):
        ws.cell(i, j, None if (isinstance(v, float) and np.isnan(v)) or v is None else (v.item() if hasattr(v, "item") else v))
ult = N + 2
for rr in range(1, ult + 1):
    for j, c in enumerate(base.columns, 1):
        cel = ws.cell(rr, j)
        cel.fill = BRANCO
        if rr == 1:
            cel.font, cel.alignment = ARIAL_N, Alignment(horizontal="left")
        elif rr == 2:
            cel.font = ARIAL_N
            cel.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cel.border = Border(top=FINA, bottom=FINA, left=SEM, right=SEM)
        else:
            cel.font = ARIAL
            cel.alignment = Alignment(horizontal="left" if c in TIPO_TXT else "right")
            if c not in TIPO_TXT:
                cel.number_format = FMT["int"]
            if rr == ult:
                cel.border = Border(bottom=FINA, top=SEM, left=SEM, right=SEM)
ws.row_dimensions[2].height = 54
ws.freeze_panes = "B3"
for j, c in enumerate(base.columns, 1):
    ws.column_dimensions[get_column_letter(j)].width = 30 if c == "estabelecimento" else 12
ws.cell(ult + 1, 1, f"Fonte: {FONTE_TEMAS}.").font = ARIAL
ws.cell(ult + 2, 1, "Nota: colunas + e − marcam 1 quando o comentário elogia (+) ou critica (−) o tema; vazio nas notas = não informado.").font = ARIAL

ind = est.wb.create_sheet("Índice", 0)
ind.sheet_view.showGridLines = False
ind.cell(1, 1, "Dados agregados dos comentários (B2) – TCC Luderia")
ind.merge_cells("A1:B1")
for j, c in enumerate(["Aba", "Tabela"], 1):
    ind.cell(2, j, c)
for i, (a, t, _) in enumerate(INDICE + [(ABA_BASE, "Base dos comentários, sem o texto", "")], 3):
    ind.cell(i, 1, a).hyperlink = f"#'{a}'!A1"
    ind.cell(i, 2, t)
ult = 2 + len(INDICE) + 1
for rr in range(1, ult + 1):
    for j in (1, 2):
        cel = ind.cell(rr, j)
        cel.fill = BRANCO
        cel.font = ARIAL_N if rr <= 2 else ARIAL
        cel.alignment = Alignment(horizontal="center" if rr <= 2 else "left", vertical="center", wrap_text=rr > 2)
        if rr == 2:
            cel.border = Border(top=FINA, bottom=FINA, left=SEM, right=SEM)
        elif rr == ult:
            cel.border = Border(bottom=FINA, top=SEM, left=SEM, right=SEM)
ind.column_dimensions["A"].width = 10
ind.column_dimensions["B"].width = 110
ind.cell(ult + 2, 1, f"Fonte: {FONTE_TEMAS}.").font = ARIAL
ind.cell(ult + 3, 1, "Gráficos em PNG de 300 dpi em dados/Dados agregados/graficos_comentarios/. Gerado por _src/extras/gerar_planilha_comentarios.py.").font = ARIAL

SAIDA.mkdir(parents=True, exist_ok=True)
saida = SAIDA / "dados_agregados_comentarios.xlsx"
est.wb.save(saida)
print(f"{saida.relative_to(ROOT)}: {len(INDICE)} tabelas + base; {len(list(GRAF.glob('*.png')))} gráficos")
print(f"Conferência: {N} comentários; soma por metrópole = {int(mt.n.sum())} (diferença {N - int(mt.n.sum())}); "
      f"{nt} com tema; {N_CASAS} casas")
